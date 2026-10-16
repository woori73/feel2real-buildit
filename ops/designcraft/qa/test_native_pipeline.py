import subprocess,json,pathlib,shutil,sys
from docx import Document
import argparse
p=argparse.ArgumentParser();p.add_argument('--cli',required=True);p.add_argument('--out',required=True);a=p.parse_args()
root=pathlib.Path(__file__).resolve().parents[3];out=pathlib.Path(a.out).resolve();out.mkdir(exist_ok=True)
results=[]
cli=pathlib.Path(a.cli).resolve()
version=subprocess.check_output([str(cli),'--version'],text=True).strip().replace('designcraft-cli ', '')
for version,cli in [(version,cli)]:
 for ext in ['txt','md','rtf','docx']:
  d=out/(version+'-'+ext);d.mkdir(exist_ok=True)
  shutil.copytree(root/'ops/designcraft/templates/ko-textbook-a4',d,dirs_exist_ok=True)
  text='한글 제목\n첫 번째 문장입니다.\n두 번째 문장입니다.\n'
  if ext=='docx':
   doc=Document();doc.add_heading('한글 제목',0);doc.add_paragraph('첫 번째 문장입니다.');doc.add_paragraph('두 번째 문장입니다.');t=doc.add_table(rows=2,cols=2);t.cell(0,0).text='항목';t.cell(0,1).text='금액';t.cell(1,0).text='교재';t.cell(1,1).text='20000';doc.save(d/'manuscript.docx')
  elif ext=='rtf':
   def esc(x):return ''.join(('\\u%d?'%(ord(c) if ord(c)<32768 else ord(c)-65536)) if ord(c)>127 else ('\\par ' if c=='\n' else c) for c in x)
   (d/'manuscript.rtf').write_text('{\\rtf1\\ansi '+esc(text)+'}')
  else:(d/('manuscript.'+ext)).write_text(('# ' if ext=='md' else '')+text)
  def run(args,stdin=None):
   p=subprocess.run([str(cli)]+args,cwd=d,input=stdin,text=True,capture_output=True,timeout=45)
   if p.returncode:raise RuntimeError(p.stderr+p.stdout)
   return json.loads(p.stdout) if p.stdout.strip() else None
  try:
   imp=run(['script','-','--in','template.designcraft','--save','job.designcraft'],f'file.place={{"path":"manuscript.{ext}","autoflow":true}}\nedit.selectAll={{}}\nstyle.paragraph.apply={{"name":"ko-body"}}\n')
   info=run(['run','--in','job.designcraft','--cmd','document.inspect']);story=imp['results'][0]['story'];before=run(['run','--in','job.designcraft','--cmd',f'story.get={{"story":{story}}}'])['text']
   assert '첫 번째 문장입니다.' in before
   s=before.index('첫 번째 문장입니다.');off=len(before[:s].encode());end=off+len('첫 번째 문장입니다.'.encode())
   run(['script','-','--in','job.designcraft','--save','edited.designcraft'],'story.replaceRange='+json.dumps({'story':story,'start':off,'end':end,'text':'AI 승인 후 수정한 문장입니다.'},ensure_ascii=False)+'\n')
   after=run(['run','--in','edited.designcraft','--cmd',f'story.get={{"story":{story}}}'])['text'];assert 'AI 승인 후 수정한 문장입니다.' in after
   pdf=run(['run','--in','edited.designcraft','--cmd','file.exportPdf={"path":"result.pdf","standard":"none","tagged":true}'])
   run(['run','--in','edited.designcraft','--page','0','--scale','1','--export','page.png'])
   import fitz
   doc=fitz.open(d/'result.pdf');pdftext=''.join(p.get_text() for p in doc);assert ''.join('AI 승인 후 수정한 문장입니다.'.split()) in ''.join(pdftext.split())
   if ext=='docx':assert '교재' in pdftext and '20000' in pdftext
   results.append({'version':version,'format':ext,'result':'PASS','pages':len(doc),'tableTextPreserved':ext=='docx','pdf':str(d/'result.pdf')})
  except Exception as e:results.append({'version':version,'format':ext,'result':'FAIL','error':repr(e)})
print(json.dumps(results,ensure_ascii=False,indent=2));(out/'results.json').write_text(json.dumps(results,ensure_ascii=False,indent=2));sys.exit(any(r['result']!='PASS' for r in results))
