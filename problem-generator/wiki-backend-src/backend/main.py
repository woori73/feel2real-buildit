import json
import contextlib
import hashlib
import os
import re
import sqlite3
import struct
import zipfile
import zlib
import uuid
from datetime import datetime
from pathlib import Path
from typing import Literal, Optional

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, Response
from pydantic import BaseModel, Field
from pydantic import ValidationError
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

import prompt_logic as pl
import ai_router
import question_contract as qc

DB_PATH = os.getenv('BUILDIT_WIKI_DB_PATH', '/workspace/data/app.db')
WORKSPACE = Path(os.getenv('BUILDIT_WIKI_WORKSPACE', '/workspace'))
EXPORT_PATH = WORKSPACE / 'feel-to-real-project-export.zip'


def get_db():
    os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    with get_db() as conn:
        conn.execute('''CREATE TABLE IF NOT EXISTS studio_runs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            role TEXT NOT NULL,
            material_type TEXT NOT NULL,
            title TEXT NOT NULL,
            status TEXT NOT NULL DEFAULT '완료',
            created_at TEXT NOT NULL
        )''')
        conn.execute('''CREATE TABLE IF NOT EXISTS digital_twin_profiles (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            school TEXT NOT NULL,
            grade TEXT NOT NULL DEFAULT '',
            exam TEXT NOT NULL DEFAULT '',
            profile_json TEXT NOT NULL,
            version INTEGER NOT NULL DEFAULT 1,
            is_active INTEGER NOT NULL DEFAULT 1,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        )''')
        conn.execute('''CREATE TABLE IF NOT EXISTS generated_questions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            twin_id INTEGER,
            school TEXT NOT NULL DEFAULT '',
            grade TEXT NOT NULL DEFAULT '',
            exam TEXT NOT NULL DEFAULT '',
            mode TEXT NOT NULL DEFAULT 'new',
            q_kind TEXT NOT NULL DEFAULT 'mc',
            question_type TEXT NOT NULL DEFAULT '',
            stem TEXT NOT NULL DEFAULT '',
            source_text TEXT NOT NULL DEFAULT '',
            options_json TEXT NOT NULL DEFAULT '[]',
            answer_json TEXT NOT NULL DEFAULT '[]',
            multi INTEGER NOT NULL DEFAULT 0,
            score INTEGER NOT NULL DEFAULT 0,
            conditions_json TEXT NOT NULL DEFAULT '[]',
            answer_format TEXT NOT NULL DEFAULT '',
            accepted_answers_json TEXT NOT NULL DEFAULT '[]',
            sample_answer TEXT NOT NULL DEFAULT '',
            grading_rule TEXT NOT NULL DEFAULT '',
            exam_title TEXT NOT NULL DEFAULT '',
            exam_group_id TEXT NOT NULL DEFAULT '',
            favorite INTEGER NOT NULL DEFAULT 0,
            raw_json TEXT NOT NULL DEFAULT '{}',
            created_at TEXT NOT NULL,
            test_id INTEGER
        )''')
        # 내신 실전 제작: 강사 문제 DB에서 선택한 문항으로 구성한 시험지를 실제로 저장한다.
        conn.execute('''CREATE TABLE IF NOT EXISTS exam_papers (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT NOT NULL DEFAULT '',
            academy_name TEXT NOT NULL DEFAULT '',
            teacher_name TEXT NOT NULL DEFAULT '',
            question_ids_json TEXT NOT NULL DEFAULT '[]',
            created_at TEXT NOT NULL
        )''')
        # 내신 실전 제작 작업 중인 시험지의 "현재 선택 상태"를 지속시키는 단일 draft 레코드.
        # React state에만 의존하지 않고, 새로고침/페이지 이동 후에도 선택 문항과 순서를 복원하기 위해 사용한다.
        conn.execute('''CREATE TABLE IF NOT EXISTS exam_draft (
            id INTEGER PRIMARY KEY CHECK (id = 1),
            title TEXT NOT NULL DEFAULT '',
            academy_name TEXT NOT NULL DEFAULT '',
            teacher_name TEXT NOT NULL DEFAULT '',
            question_ids_json TEXT NOT NULL DEFAULT '[]',
            updated_at TEXT NOT NULL
        )''')
        # 학생 프로필: 이름/학교/학년/학기/시험(중간·기말)만 저장한다 (단순 유지 요구사항).
        conn.execute('''CREATE TABLE IF NOT EXISTS students (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            school TEXT NOT NULL DEFAULT '',
            grade TEXT NOT NULL DEFAULT '',
            semester TEXT NOT NULL DEFAULT '',
            exam_type TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL
        )''')
        # 학생 오답 관리: 학생별·시험별로 오답 문제 원문(=사용자가 최종 확인/수정한 추출 텍스트)과
        # 정답을 저장한다. 각 저장은 독립된 record이며(덮어쓰지 않고 계속 누적), student_id로
        # 학생과 연결한다.
        # source_text: 오답 사진/PDF에서 추출·정리된 원본 텍스트를 함께 보관한다 (마이그레이션으로 추가).
        # title/question_type: 강사 문제 DB와 유사한 목록/간략보기를 위한 표시용 라벨 (마이그레이션으로 추가).
        conn.execute('''CREATE TABLE IF NOT EXISTS student_wrong_notes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            student_id INTEGER NOT NULL,
            exam_label TEXT NOT NULL DEFAULT '',
            title TEXT NOT NULL DEFAULT '',
            question_type TEXT NOT NULL DEFAULT '',
            wrong_text TEXT NOT NULL DEFAULT '',
            answer_text TEXT NOT NULL DEFAULT '',
            source_text TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL,
            FOREIGN KEY (student_id) REFERENCES students(id)
        )''')
        conn.execute('''CREATE TABLE IF NOT EXISTS educator_insights (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            content TEXT NOT NULL,
            school TEXT NOT NULL DEFAULT '',
            grade TEXT NOT NULL DEFAULT '',
            student_context TEXT NOT NULL DEFAULT '',
            lesson_context TEXT NOT NULL DEFAULT '',
            year INTEGER,
            semester TEXT NOT NULL DEFAULT '',
            exam TEXT NOT NULL DEFAULT '',
            source TEXT NOT NULL DEFAULT 'director',
            created_at TEXT NOT NULL
        )''')
        conn.execute('''CREATE TABLE IF NOT EXISTS exam_material_sets (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            core_type TEXT NOT NULL CHECK (core_type IN ('school_print','student_wrong_note')),
            school TEXT NOT NULL DEFAULT '', grade TEXT NOT NULL DEFAULT '',
            year INTEGER NOT NULL, semester TEXT NOT NULL DEFAULT '', exam TEXT NOT NULL DEFAULT '',
            student_id INTEGER, title TEXT NOT NULL DEFAULT '',
            original_name TEXT NOT NULL, original_type TEXT NOT NULL DEFAULT 'application/octet-stream',
            original_blob BLOB NOT NULL, json_text TEXT NOT NULL DEFAULT '{}',
            html_text TEXT NOT NULL DEFAULT '', created_at TEXT NOT NULL,
            FOREIGN KEY (student_id) REFERENCES students(id)
        )''')
        conn.execute('CREATE INDEX IF NOT EXISTS idx_exam_material_context ON exam_material_sets(school, grade, year, semester, exam, core_type)')
        # FINAL/직보 결과물: 시험 Context(학교/학년/연도/학기/시험)별로 외부 LLM(Work/Claude/ChatGPT 등)이나
        # 강사가 완성해 저장하는 최종 산출물이다. exam_material_sets(자료)와 개념이 달라 별도 테이블로 둔다.
        # kind로 FINAL 완성본과 직보 완성본을 구분한다 — 이번 작업은 kind='final' 저장 UI만 연결하고
        # 'jikbo'는 향후 MCP save_jikbo 도구가 같은 API를 그대로 쓸 수 있도록 스키마만 열어 둔다.
        conn.execute('''CREATE TABLE IF NOT EXISTS final_deliverables (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            kind TEXT NOT NULL DEFAULT 'final' CHECK (kind IN ('final','jikbo')),
            school TEXT NOT NULL DEFAULT '', grade TEXT NOT NULL DEFAULT '',
            year INTEGER NOT NULL, semester TEXT NOT NULL DEFAULT '', exam TEXT NOT NULL DEFAULT '',
            title TEXT NOT NULL DEFAULT '',
            original_name TEXT NOT NULL DEFAULT '', original_type TEXT NOT NULL DEFAULT 'application/octet-stream',
            original_blob BLOB NOT NULL, html_text TEXT NOT NULL DEFAULT '', json_text TEXT NOT NULL DEFAULT '{}',
            created_at TEXT NOT NULL
        )''')
        conn.execute('CREATE INDEX IF NOT EXISTS idx_final_deliverable_context ON final_deliverables(school, grade, year, semester, exam, kind)')
        conn.execute('''CREATE TABLE IF NOT EXISTS csat_question_sources (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            exam_kind TEXT NOT NULL CHECK (exam_kind IN ('half','full')),
            source_filename TEXT NOT NULL,
            exam_metadata_json TEXT NOT NULL DEFAULT '{}',
            parser TEXT NOT NULL,
            raw_text TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL
        )''')
        conn.execute('''CREATE TABLE IF NOT EXISTS csat_questions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            source_id INTEGER NOT NULL,
            exam_kind TEXT NOT NULL,
            question_no INTEGER NOT NULL,
            csat_type_id TEXT NOT NULL,
            csat_type_name TEXT NOT NULL,
            passage TEXT NOT NULL DEFAULT '',
            prompt TEXT NOT NULL DEFAULT '',
            choices_json TEXT NOT NULL DEFAULT '[]',
            answer TEXT NOT NULL DEFAULT '',
            solution TEXT NOT NULL DEFAULT '',
            footnote TEXT NOT NULL DEFAULT '',
            group_label TEXT NOT NULL DEFAULT '',
            source_reference TEXT NOT NULL DEFAULT '',
            type_mismatch INTEGER NOT NULL DEFAULT 0,
            mismatch_reason TEXT NOT NULL DEFAULT '',
            raw_json TEXT NOT NULL DEFAULT '{}',
            created_at TEXT NOT NULL,
            FOREIGN KEY (source_id) REFERENCES csat_question_sources(id)
        )''')
        conn.execute('CREATE INDEX IF NOT EXISTS idx_csat_questions_kind_no ON csat_questions(exam_kind, question_no)')
        conn.execute('CREATE INDEX IF NOT EXISTS idx_csat_questions_type ON csat_questions(csat_type_id)')
        conn.execute('''CREATE TABLE IF NOT EXISTS csat_wrong_answers (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            question_id INTEGER,
            student_name TEXT NOT NULL,
            exam_name TEXT NOT NULL DEFAULT '',
            question_no INTEGER NOT NULL,
            csat_type_id TEXT NOT NULL,
            csat_type_name TEXT NOT NULL,
            wrong_answer TEXT NOT NULL DEFAULT '',
            correct_answer TEXT NOT NULL DEFAULT '',
            memo TEXT NOT NULL DEFAULT '',
            source_filename TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL,
            FOREIGN KEY (question_id) REFERENCES csat_questions(id)
        )''')
        conn.execute('CREATE INDEX IF NOT EXISTS idx_csat_wrong_lookup ON csat_wrong_answers(student_name, exam_name, question_no, csat_type_id)')
        # 수능 제작물의 공통 저장 기반. Question 원본/Assessment 구성/Template 표현을 분리한다.
        conn.execute('''CREATE TABLE IF NOT EXISTS csat_assessments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            assessment_type TEXT NOT NULL CHECK (assessment_type IN ('half','full','wrong','daily')),
            title TEXT NOT NULL DEFAULT '',
            template_id TEXT NOT NULL DEFAULT '',
            request_json TEXT NOT NULL DEFAULT '{}',
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        )''')
        conn.execute('''CREATE TABLE IF NOT EXISTS csat_assessment_items (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            assessment_id INTEGER NOT NULL,
            question_id INTEGER NOT NULL,
            position INTEGER NOT NULL,
            day_no INTEGER,
            created_at TEXT NOT NULL,
            FOREIGN KEY (assessment_id) REFERENCES csat_assessments(id),
            FOREIGN KEY (question_id) REFERENCES csat_questions(id),
            UNIQUE (assessment_id, position),
            UNIQUE (assessment_id, question_id)
        )''')
        conn.execute('''CREATE TABLE IF NOT EXISTS csat_question_usage (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            question_id INTEGER NOT NULL,
            assessment_id INTEGER NOT NULL,
            assessment_type TEXT NOT NULL,
            used_at TEXT NOT NULL,
            FOREIGN KEY (question_id) REFERENCES csat_questions(id),
            FOREIGN KEY (assessment_id) REFERENCES csat_assessments(id),
            UNIQUE (assessment_id, question_id)
        )''')
        conn.execute('CREATE INDEX IF NOT EXISTS idx_csat_assessment_items_order ON csat_assessment_items(assessment_id, day_no, position)')
        conn.execute('CREATE INDEX IF NOT EXISTS idx_csat_usage_lookup ON csat_question_usage(question_id, assessment_type, used_at)')
        conn.execute('''CREATE TABLE IF NOT EXISTS ai_ab_evaluations (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            school TEXT NOT NULL DEFAULT '', grade TEXT NOT NULL DEFAULT '', exam TEXT NOT NULL DEFAULT '',
            twin_id INTEGER, question_type TEXT NOT NULL DEFAULT '',
            gemini_provider TEXT NOT NULL, gemini_model TEXT NOT NULL, gemini_result_id TEXT NOT NULL,
            ox_provider TEXT NOT NULL, ox_model TEXT NOT NULL, ox_result_id TEXT NOT NULL,
            choice TEXT NOT NULL CHECK (choice IN ('gemini','ox-alpha','both-revise')),
            created_at TEXT NOT NULL
        )''')
        # Generated Question Store — MCP가 생성한 시험문항을 시험(Test) 단위로 영구 저장한다.
        # source_materials(학습자료) → generated_tests(생성 시험) → generated_questions(문항, test_id로 연결).
        # 기존 generated_questions/exam_material_sets/digital_twin_profiles는 건드리지 않고 위에 얹는다.
        conn.execute('''CREATE TABLE IF NOT EXISTS source_materials (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            school TEXT NOT NULL DEFAULT '',
            grade TEXT NOT NULL DEFAULT '',
            title TEXT NOT NULL DEFAULT '',
            publisher TEXT NOT NULL DEFAULT '',
            lesson TEXT NOT NULL DEFAULT '',
            material_type TEXT NOT NULL DEFAULT '',
            school_print_material_id INTEGER,
            created_at TEXT NOT NULL
        )''')
        conn.execute('CREATE UNIQUE INDEX IF NOT EXISTS uk_source_material_ident ON source_materials(school, grade, title)')
        conn.execute('''CREATE TABLE IF NOT EXISTS generated_tests (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            external_id TEXT NOT NULL DEFAULT '',
            source_material_id INTEGER NOT NULL,
            source_page TEXT NOT NULL DEFAULT '',
            school TEXT NOT NULL DEFAULT '',
            grade TEXT NOT NULL DEFAULT '',
            year INTEGER,
            semester TEXT NOT NULL DEFAULT '',
            exam TEXT NOT NULL DEFAULT '',
            test_type TEXT NOT NULL DEFAULT '',
            title TEXT NOT NULL DEFAULT '',
            twin_id INTEGER,
            twin_label TEXT NOT NULL DEFAULT '',
            twin_pattern TEXT NOT NULL DEFAULT '',
            difficulty TEXT NOT NULL DEFAULT '',
            target_vocabulary TEXT NOT NULL DEFAULT '[]',
            target_grammar TEXT NOT NULL DEFAULT '[]',
            target_passage TEXT NOT NULL DEFAULT '',
            generation_source TEXT NOT NULL DEFAULT 'mcp',
            question_count INTEGER NOT NULL DEFAULT 0,
            raw_request_json TEXT NOT NULL DEFAULT '{}',
            created_at TEXT NOT NULL,
            FOREIGN KEY (source_material_id) REFERENCES source_materials(id)
        )''')
        conn.execute("CREATE UNIQUE INDEX IF NOT EXISTS uk_generated_test_external ON generated_tests(external_id) WHERE external_id <> ''")
        conn.execute('CREATE INDEX IF NOT EXISTS idx_generated_test_material ON generated_tests(source_material_id)')
        conn.execute('CREATE INDEX IF NOT EXISTS idx_generated_test_context ON generated_tests(school, grade, exam)')


def migrate_schema():
    """기존에 만들어져 있던 DB에는 없는 컬럼을 안전하게 추가한다.
    기존 데이터/테이블은 그대로 두고 누락된 컬럼만 ALTER TABLE로 보강한다 (데이터 삭제·재생성 없음)."""
    with get_db() as conn:
        cols = {row['name'] for row in conn.execute("PRAGMA table_info(student_wrong_notes)").fetchall()}
        if 'source_text' not in cols:
            conn.execute("ALTER TABLE student_wrong_notes ADD COLUMN source_text TEXT NOT NULL DEFAULT ''")
        if 'title' not in cols:
            conn.execute("ALTER TABLE student_wrong_notes ADD COLUMN title TEXT NOT NULL DEFAULT ''")
        if 'question_type' not in cols:
            conn.execute("ALTER TABLE student_wrong_notes ADD COLUMN question_type TEXT NOT NULL DEFAULT ''")
        insight_cols = {row['name'] for row in conn.execute("PRAGMA table_info(educator_insights)").fetchall()}
        for name, definition in (
            ('year', 'INTEGER'), ('semester', "TEXT NOT NULL DEFAULT ''"),
            ('exam', "TEXT NOT NULL DEFAULT ''"), ('source', "TEXT NOT NULL DEFAULT 'director'"),
        ):
            if name not in insight_cols:
                conn.execute(f'ALTER TABLE educator_insights ADD COLUMN {name} {definition}')
        material_cols = {row['name'] for row in conn.execute("PRAGMA table_info(exam_material_sets)").fetchall()}
        if 'html_text' not in material_cols:
            conn.execute("ALTER TABLE exam_material_sets ADD COLUMN html_text TEXT NOT NULL DEFAULT ''")
        source_cols = {row['name'] for row in conn.execute("PRAGMA table_info(csat_question_sources)").fetchall()}
        for name, definition in (
            ('source_hash', "TEXT NOT NULL DEFAULT ''"),
            ('original_stored', 'INTEGER NOT NULL DEFAULT 0'),
            ('owner_scope', "TEXT NOT NULL DEFAULT 'shared'"),
        ):
            if name not in source_cols:
                conn.execute(f'ALTER TABLE csat_question_sources ADD COLUMN {name} {definition}')
        conn.execute("CREATE UNIQUE INDEX IF NOT EXISTS uk_csat_source_hash_kind ON csat_question_sources(source_hash, exam_kind) WHERE source_hash <> ''")
        wrong_cols = {row['name'] for row in conn.execute("PRAGMA table_info(csat_wrong_answers)").fetchall()}
        if 'question_id' not in wrong_cols:
            conn.execute('ALTER TABLE csat_wrong_answers ADD COLUMN question_id INTEGER')
        conn.execute('CREATE INDEX IF NOT EXISTS idx_csat_wrong_question ON csat_wrong_answers(question_id)')
        gq_cols = {row['name'] for row in conn.execute("PRAGMA table_info(generated_questions)").fetchall()}
        if 'test_id' not in gq_cols:
            conn.execute("ALTER TABLE generated_questions ADD COLUMN test_id INTEGER")


def seed_default_twins():
    """기존 HTML에 내장되어 있던 6개 학교(신사중3, 신구중2, 신구중3, 압구정중2, 압구정중3, 언북중3)
    Digital Twin을 기본 데이터로 자동 Insert한다. 이미 같은 학교·학년의 Digital Twin이 있으면
    건드리지 않고, 없는 학교만 채워 넣는다 (기존 학교/업로드 자료를 삭제하거나 덮어쓰지 않음).
    이후 HTML 업로드는 이 6개를 지우지 않고 새 학교만 추가한다."""
    with get_db() as conn:
        now = datetime.now().strftime('%Y.%m.%d %H:%M')
        for profile in pl.DEFAULT_SCHOOL_PROFILES:
            existing = conn.execute(
                'SELECT id FROM digital_twin_profiles WHERE is_active = 1 AND school = ? AND grade = ?',
                (profile['school'], profile['grade']),
            ).fetchone()
            if existing:
                continue
            types = [
                {'name': pl.TYPE_DEF_BY_ID[type_id]['label'], 'description': pl.TYPE_DEF_BY_ID[type_id]['prompt'], 'ratio': '', 'difficulty': ''}
                for type_id in profile['types'] if type_id in pl.TYPE_DEF_BY_ID
            ]
            data = {
                'school': profile['school'], 'grade': profile['grade'], 'exam': '전체',
                'dna': profile['desc'],
                'description': f"{profile['school']} {profile['grade']} 기본 Digital Twin (기존 HTML 기준 데이터)",
                'types': types,
                'rules': ['기존 HTML의 학교 출제 DNA와 유형 순서를 그대로 따른다.', '같은 소재·같은 문장으로 문항을 반복하지 않는다.'],
                'difficulty': pl.DIFFICULTY_LABELS['actual'],
                'questionRatio': {}, 'trapPatterns': [], 'choiceStyle': '', 'writingStyle': '',
                'updatedAt': now, 'rawSource': '',
            }
            conn.execute(
                'INSERT INTO digital_twin_profiles (school, grade, exam, profile_json, version, is_active, created_at, updated_at) VALUES (?,?,?,?,?,1,?,?)',
                (data['school'], data['grade'], data['exam'], json.dumps(data, ensure_ascii=False), 1, now, now),
            )


# ==================================================================================
# Legacy 마이그레이션 — PR #63 이전에 save_generated_questions(구형 tool)로 저장된 특정 MCP
# 문항 묶음을 Generated Test 구조에 1회 연결한다. 문항 내용은 복제·재생성하지 않고 test_id만
# 채운다. exam_group_id별 external_id guard로 완전 idempotent(재배포해도 중복 안 생김).
# 향후 신규 저장은 save_generated_test가 정식 schema로 하므로 여기에 새 항목이 추가될 일은 없다.
# ==================================================================================
_LEGACY_TEST_MIGRATIONS = [
    {
        'exam_group_id': 'mcp_f5061b3f4b5b43778291b2c1e25c37fc',
        'school': '신사중', 'grade': '3학년', 'exam': '전체', 'year': 2026, 'semester': '2학기',
        'source_material': 'The Boiling Frog Myth', 'source_page': '2',
        'material_type': 'ted_ed', 'test_type': 'vocabulary',
        'title': 'The Boiling Frog Myth — Page 2 Vocabulary Test',
        'twin_label': '신사중 3학년 출제자 디지털트윈',
        'twin_pattern': '복수정답·문학개념 전이·심리/담화 추론·조건형 영작',
        'difficulty': '실제 시험 수준의 중상 난이도',
        'delete_question_ids': [83, 84],  # MCP 진단 중 잘못 저장된 dummy — 내용이 test/test3 짧은 것만 삭제
    },
]

# sourceText에 우회 저장된 Digital Twin 태그. 실제 형식은 key:value 쌍이며 대소문자가 섞여 있다:
#   [문항ID:Q30 | 대상어휘:be oblivious to | 어휘등급:CHALLENGE | 난이도:상 | Page:4 | DigitalTwin근거:... | DigitalTwin패턴:...]
# (초기 저장분은 bare 6-value: [대상어휘 | 어휘등급 | 난이도 | Page | 근거 | 패턴] 도 지원)
_LEGACY_TAG_RE = re.compile(r'\[([^\[\]]+)\]')
_LEGACY_TAG_KEYS = {
    '대상어휘': 'target_vocab', '어휘등급': 'vocab_level', '난이도': 'difficulty',
    'page': 'page', 'digitaltwin근거': 'twin_basis', 'digitaltwin패턴': 'twin_pattern',
}


def _parse_legacy_twin_tag(source_text: str) -> Optional[dict]:
    if not source_text:
        return None
    for group in _LEGACY_TAG_RE.findall(source_text):
        if '|' not in group:
            continue
        fields = [f.strip() for f in group.split('|')]
        if all(':' in f for f in fields):
            out: dict = {}
            for field in fields:
                key, _, val = field.partition(':')
                canon = _LEGACY_TAG_KEYS.get(key.strip().lower())
                if canon and val.strip():
                    out.setdefault(canon, val.strip())
            if out.get('target_vocab'):
                return out
        elif len(fields) == 6:
            return {'target_vocab': fields[0], 'vocab_level': fields[1], 'difficulty': fields[2],
                    'page': fields[3], 'twin_basis': fields[4], 'twin_pattern': fields[5]}
    return None


def _aggregate_legacy_vocab(rows) -> list:
    vocab: list = []
    for r in rows:
        tag = _parse_legacy_twin_tag(r['source_text'])
        if tag and tag.get('target_vocab') and tag['target_vocab'] not in vocab:
            vocab.append(tag['target_vocab'])
    return vocab


def _audit_note(conn, school: str, grade: str, year: int, semester: str, exam: str, content: str) -> None:
    conn.execute(
        'INSERT INTO educator_insights (content, school, grade, student_context, lesson_context, '
        'year, semester, exam, source, created_at) VALUES (?,?,?,?,?,?,?,?,?,?)',
        (content, school, grade, '', 'Generated Question Store 마이그레이션', year, semester, exam,
         'claude', datetime.now().isoformat()),
    )


def backfill_legacy_generated_tests():
    with get_db() as conn:
        for m in _LEGACY_TEST_MIGRATIONS:
            ext = 'legacy-migrate-' + m['exam_group_id']
            audit = []
            did_work = False

            # 1) 잘못 저장된 진단 row 삭제 — 내용이 test/test3 짧은 dummy이고 이 묶음 소속이 아닐 때만.
            for qid in m.get('delete_question_ids', []):
                r = conn.execute(
                    'SELECT id, stem, source_text, exam_group_id FROM generated_questions WHERE id = ?', (qid,),
                ).fetchone()
                if not r:
                    continue
                blob = ((r['stem'] or '') + ' ' + (r['source_text'] or '')).strip().lower()
                if r['exam_group_id'] != m['exam_group_id'] and len(blob) <= 12 and 'test' in blob:
                    conn.execute('DELETE FROM generated_questions WHERE id = ?', (qid,))
                    audit.append(f'delete id={qid}: 삭제함 (content={blob!r})')
                    did_work = True
                else:
                    audit.append(f'delete id={qid}: SKIP — 진단 데이터로 확인 안 됨 (content={blob!r}, group={r["exam_group_id"]})')

            group_rows = conn.execute(
                'SELECT * FROM generated_questions WHERE exam_group_id = ? ORDER BY id', (m['exam_group_id'],),
            ).fetchall()
            twin = _match_digital_twin(conn, m['school'], m['grade'], m['exam'])
            twin_id = twin.get('id') if twin else None
            twin_label = m.get('twin_label') or (pl._twin_name(twin) if twin else '') or f"{m['school']} {m['grade']} 출제자 디지털트윈"
            twin_pattern = m.get('twin_pattern') or (twin.get('dna') if twin else '') or ''
            difficulty = m.get('difficulty') or '실제 시험 수준의 중상 난이도'

            existing = conn.execute('SELECT * FROM generated_tests WHERE external_id = ?', (ext,)).fetchone()

            # 2a) 이미 연결됨 — metadata(target_vocabulary/twin_pattern)가 비어 있으면 태그 재파싱으로 보강.
            if existing:
                needs_repair = (not (json.loads(existing['target_vocabulary'] or '[]'))
                                or not existing['twin_pattern'] or not existing['twin_label'])
                if needs_repair and group_rows:
                    vocab = _aggregate_legacy_vocab(group_rows)
                    conn.execute(
                        'UPDATE generated_tests SET target_vocabulary=?, twin_label=?, twin_pattern=?, '
                        'difficulty=?, source_page=?, question_count=? WHERE id=?',
                        (json.dumps(vocab, ensure_ascii=False), twin_label, twin_pattern, difficulty,
                         m['source_page'], len(group_rows), existing['id']),
                    )
                    audit.append(f'repaired generated_test id={existing["id"]}: target_vocabulary={len(vocab)}, '
                                 f'twin_label={twin_label!r}, twin_pattern set')
                    did_work = True
                if did_work:
                    _audit_note(conn, m['school'], m['grade'], m['year'], m['semester'], m['exam'],
                                f'[backfill {m["exam_group_id"]}] ' + ' | '.join(audit))
                continue

            # 2b) 최초 연결 — 문항 내용 복제 없이 test_id만 채운다.
            if not group_rows:
                if did_work:
                    _audit_note(conn, m['school'], m['grade'], m['year'], m['semester'], m['exam'],
                                f'[backfill {m["exam_group_id"]}] 묶음 없음. ' + ' | '.join(audit))
                continue

            vocab = _aggregate_legacy_vocab(group_rows)
            sm_id = _upsert_source_material(conn, m['school'], m['grade'], m['source_material'],
                                           material_type=m['material_type'])
            now = datetime.now().isoformat()
            cur = conn.execute(
                '''INSERT INTO generated_tests
                (external_id, source_material_id, source_page, school, grade, year, semester, exam,
                 test_type, title, twin_id, twin_label, twin_pattern, difficulty, target_vocabulary,
                 target_grammar, target_passage, generation_source, question_count, raw_request_json, created_at)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',
                (ext, sm_id, m['source_page'], m['school'], m['grade'], m['year'], m['semester'],
                 m['exam'], m['test_type'], m['title'], twin_id, twin_label, twin_pattern, difficulty,
                 json.dumps(vocab, ensure_ascii=False), '[]', '',
                 'mcp', len(group_rows), json.dumps({'migrated_from_exam_group_id': m['exam_group_id']}, ensure_ascii=False), now),
            )
            test_id = cur.lastrowid
            conn.execute(
                'UPDATE generated_questions SET test_id = ?, school = ?, grade = ?, exam = ? WHERE exam_group_id = ?',
                (test_id, m['school'], m['grade'], m['exam'], m['exam_group_id']),
            )
            audit.append(f'linked {len(group_rows)}문항 -> generated_test id={test_id}, source_material_id={sm_id}, '
                         f'target_vocabulary={len(vocab)}, twin_id={twin_id}, twin_label={twin_label!r}')
            _audit_note(conn, m['school'], m['grade'], m['year'], m['semester'], m['exam'],
                        f'[backfill {m["exam_group_id"]}] ' + ' | '.join(audit))


@contextlib.asynccontextmanager
async def _lifespan(_app: FastAPI):
    # mcp_server의 StreamableHTTP session manager는 자체 task group을 갖고 있어서
    # app.mount()만으로는 시작되지 않는다 — 부모 app의 lifespan 안에서 명시적으로
    # run()에 진입시켜야 한다. mcp_server는 main을 지연 import하므로 순환 import를
    # 피하기 위해 여기서도 지연 import한다(서버 startup 시점에만 실행됨).
    from mcp_server import mcp as _mcp
    # 모든 모듈 레벨 정의가 끝난 뒤(=서버 startup) 실행한다 — backfill이 파일 뒤쪽의
    # _upsert_source_material / _match_digital_twin 을 참조하므로 import 시점엔 부를 수 없다.
    backfill_legacy_generated_tests()
    async with _mcp.session_manager.run():
        yield


app = FastAPI(lifespan=_lifespan)
init_db()
migrate_schema()
seed_default_twins()


class StudioRun(BaseModel):
    role: str
    material_type: str
    title: str


@app.get('/api/health')
def health():
    return {'ok': True}


@app.get('/api/ai/providers')
def ai_providers():
    """Return capability flags only. Secret values never leave the server."""
    return ai_router.availability()


@app.get('/api/ai/openrouter/free-models')
def openrouter_models():
    """Public-safe model metadata only; the OpenRouter credential remains server-side."""
    return {'models': ai_router.openrouter_free_models(), 'autoModel': ai_router.FREE_ROUTER}


@app.get('/api/runs')
def list_runs():
    with get_db() as conn:
        rows = conn.execute('SELECT * FROM studio_runs ORDER BY id DESC LIMIT 20').fetchall()
        return [dict(row) for row in rows]


@app.post('/api/runs')
def create_run(run: StudioRun):
    if not run.title.strip():
        raise HTTPException(status_code=400, detail='자료 이름을 입력해 주세요.')
    created_at = datetime.now().strftime('%Y.%m.%d %H:%M')
    with get_db() as conn:
        cursor = conn.execute('INSERT INTO studio_runs (role, material_type, title, created_at) VALUES (?, ?, ?, ?)', (run.role, run.material_type, run.title.strip(), created_at))
        row = conn.execute('SELECT * FROM studio_runs WHERE id = ?', (cursor.lastrowid,)).fetchone()
        return dict(row)


# ==================================================================================
# Educator Brain — 수업 직후의 관찰과 판단을 다음 제작 작업에서 다시 쓸 수 있게 축적한다.
# 이번 범위는 기록 생성과 최근 기록 조회다. 기존 기록 수정/삭제는 별도 관리 UX가 생길 때 추가한다.
# ==================================================================================

class EducatorInsightCreate(BaseModel):
    content: str
    school: str = ''
    grade: str = ''
    student_context: str = ''
    lesson_context: str = ''
    year: Optional[int] = None
    semester: str = ''
    exam: str = ''
    # 'claude'/'chatgpt'는 MCP save_class_insights(mcp_server.py)가 쓰는 외부 LLM 출처 값이다.
    # 기존 'director'/'school_teacher' 값과 동작·우선순위 정렬(_insight_context_items)은 그대로 유지된다.
    source: Literal['director', 'school_teacher', 'claude', 'chatgpt'] = 'director'


@app.post('/api/educator-insights', status_code=201)
def create_educator_insight(payload: EducatorInsightCreate):
    content = payload.content.strip()
    if not content:
        raise HTTPException(status_code=400, detail='수업에서 발견한 내용을 입력해 주세요.')
    if len(content) > 10000:
        raise HTTPException(status_code=400, detail='수업 인사이트는 10,000자까지 기록할 수 있습니다.')
    now = datetime.now().isoformat()
    with get_db() as conn:
        cursor = conn.execute(
            '''INSERT INTO educator_insights
               (content, school, grade, student_context, lesson_context, year, semester, exam, source, created_at)
               VALUES (?,?,?,?,?,?,?,?,?,?)''',
            (
                content, payload.school.strip()[:100], payload.grade.strip()[:50],
                payload.student_context.strip()[:200], payload.lesson_context.strip()[:200], payload.year,
                payload.semester.strip()[:20], payload.exam.strip()[:80], payload.source, now,
            ),
        )
        row = conn.execute('SELECT * FROM educator_insights WHERE id = ?', (cursor.lastrowid,)).fetchone()
        return dict(row)


@app.get('/api/educator-insights')
def list_educator_insights(limit: int = 5, school: str = '', grade: str = '', year: Optional[int] = None,
                           semester: str = '', exam: str = ''):
    safe_limit = max(1, min(limit, 20))
    with get_db() as conn:
        clauses, values = ['1=1'], []
        for column, value in (('school', school), ('grade', grade), ('semester', semester), ('exam', exam)):
            if value:
                clauses.append(f'{column} = ?'); values.append(value)
        if year is not None:
            clauses.append('(year = ? OR year IS NULL)'); values.append(year)
        values.append(safe_limit)
        rows = conn.execute(f"SELECT * FROM educator_insights WHERE {' AND '.join(clauses)} ORDER BY id DESC LIMIT ?", values).fetchall()
        return [dict(row) for row in rows]


def _validate_json_upload(json_file: Optional[UploadFile], json_text: str, required: bool = True) -> str:
    """core_type='school_print'는 JSON을 나중에 추가할 수 있도록 선택 입력으로 둔다
    (NAESIN_WIKI_V1_1_TASK.md STEP1 fix). student_wrong_note는 기존과 동일하게 필수."""
    raw = json_text.strip()
    if json_file and json_file.filename:
        raw = json_file.file.read().decode('utf-8-sig').strip()
    if not raw:
        if required:
            raise HTTPException(status_code=400, detail='JSON 파일을 올리거나 JSON을 붙여넣어 주세요.')
        return '{}'
    try:
        parsed = json.loads(raw)
    except (ValueError, UnicodeDecodeError):
        raise HTTPException(status_code=400, detail='올바른 JSON 형식인지 확인해 주세요.')
    return json.dumps(parsed, ensure_ascii=False)


@app.post('/api/exam-materials', status_code=201)
def create_exam_material(core_type: Literal['school_print', 'student_wrong_note'] = Form(...),
                         school: str = Form(...), grade: str = Form(...), year: int = Form(...),
                         semester: str = Form(...), exam: str = Form(...), title: str = Form(''),
                         student_id: Optional[int] = Form(None), original: Optional[UploadFile] = File(None),
                         json_file: Optional[UploadFile] = File(None), json_text: str = Form(''),
                         html_text: str = Form('')):
    if core_type == 'student_wrong_note' and student_id is None:
        raise HTTPException(status_code=400, detail='학생을 선택해 주세요.')
    has_original = bool(original and original.filename)
    if core_type == 'student_wrong_note' and not has_original:
        raise HTTPException(status_code=400, detail='원본 파일을 올려 주세요.')
    canonical_html = html_text.strip() if core_type == 'school_print' and isinstance(html_text, str) else ''
    if core_type == 'school_print' and not has_original and not canonical_html:
        raise HTTPException(status_code=400, detail='HTML 자료를 올리거나 붙여넣어 주세요.')
    original_name, original_type, original_bytes = '', 'application/octet-stream', b''
    if has_original:
        original_bytes = original.file.read()
        if not original_bytes or len(original_bytes) > 30 * 1024 * 1024:
            raise HTTPException(status_code=400, detail='원본은 30MB 이하의 파일이어야 합니다.')
        original_name, original_type = original.filename[:255], original.content_type or 'application/octet-stream'
    canonical_json = _validate_json_upload(json_file, json_text, required=core_type == 'student_wrong_note')
    now = datetime.now().isoformat()
    with get_db() as conn:
        cursor = conn.execute('''INSERT INTO exam_material_sets
            (core_type, school, grade, year, semester, exam, student_id, title, original_name,
             original_type, original_blob, json_text, html_text, created_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',
            (core_type, school.strip(), grade.strip(), year, semester.strip(), exam.strip(), student_id,
             title.strip()[:150], original_name, original_type,
             original_bytes, canonical_json, canonical_html, now))
        row = conn.execute('''SELECT m.id,m.core_type,m.school,m.grade,m.year,m.semester,m.exam,m.student_id,
            s.name student_name,m.title,m.original_name,m.original_type,m.html_text,m.created_at
            FROM exam_material_sets m LEFT JOIN students s ON s.id=m.student_id WHERE m.id=?''', (cursor.lastrowid,)).fetchone()
        return dict(row)


@app.get('/api/exam-materials')
def list_exam_materials(school: str, grade: str, year: int, semester: str, exam: str):
    with get_db() as conn:
        rows = conn.execute('''SELECT m.id,m.core_type,m.school,m.grade,m.year,m.semester,m.exam,m.student_id,
            s.name student_name,m.title,m.original_name,m.original_type,m.created_at,
            (CASE WHEN m.html_text <> '' THEN 1 ELSE 0 END) AS has_html
            FROM exam_material_sets m LEFT JOIN students s ON s.id=m.student_id
            WHERE m.school=? AND m.grade=? AND m.year=? AND m.semester=? AND m.exam=? ORDER BY m.id DESC''',
            (school, grade, year, semester, exam)).fetchall()
        return [dict(row) for row in rows]


@app.get('/api/exam-materials/{material_id}/original')
def open_exam_material_original(material_id: int):
    with get_db() as conn:
        row = conn.execute('SELECT original_blob,original_type,original_name FROM exam_material_sets WHERE id=?', (material_id,)).fetchone()
        if not row: raise HTTPException(status_code=404, detail='자료를 찾지 못했습니다.')
        safe_name = re.sub(r'[^0-9A-Za-z._-]', '_', row['original_name'])
        return Response(row['original_blob'], media_type=row['original_type'], headers={'Content-Disposition': f'inline; filename="{safe_name}"'})


@app.get('/api/exam-materials/{material_id}/json')
def open_exam_material_json(material_id: int):
    with get_db() as conn:
        row = conn.execute('SELECT json_text FROM exam_material_sets WHERE id=?', (material_id,)).fetchone()
        if not row: raise HTTPException(status_code=404, detail='자료를 찾지 못했습니다.')
        return Response(row['json_text'], media_type='application/json; charset=utf-8')


@app.get('/api/exam-materials/{material_id}/html')
def open_exam_material_html(material_id: int):
    with get_db() as conn:
        row = conn.execute('SELECT html_text FROM exam_material_sets WHERE id=?', (material_id,)).fetchone()
        if not row: raise HTTPException(status_code=404, detail='자료를 찾지 못했습니다.')
        return Response(row['html_text'] or '', media_type='text/html; charset=utf-8')


# ==================================================================================
# FINAL/직보 결과물 — 시험 Context별로 완성된 산출물을 저장·조회한다. 학교 프린트(exam_material_sets)와
# 같은 원본+HTML+JSON 3종 저장 방식을 그대로 따라 사용자가 이미 익숙한 방식을 재사용한다.
# 이번 작업에서 LLM이 FINAL을 자동 생성하지는 않는다 — 이 API는 저장·조회 자리만 마련한다.
# ==================================================================================

@app.post('/api/final-deliverables', status_code=201)
def create_final_deliverable(kind: Literal['final', 'jikbo'] = Form('final'),
                             school: str = Form(...), grade: str = Form(...), year: int = Form(...),
                             semester: str = Form(...), exam: str = Form(...), title: str = Form(''),
                             original: Optional[UploadFile] = File(None),
                             json_file: Optional[UploadFile] = File(None), json_text: str = Form(''),
                             html_text: str = Form('')):
    has_original = bool(original and original.filename)
    canonical_html = html_text.strip()
    if not has_original and not canonical_html:
        raise HTTPException(status_code=400, detail='결과 파일을 올리거나 HTML을 붙여넣어 주세요.')
    original_name, original_type, original_bytes = '', 'application/octet-stream', b''
    if has_original:
        original_bytes = original.file.read()
        if not original_bytes or len(original_bytes) > 30 * 1024 * 1024:
            raise HTTPException(status_code=400, detail='결과 파일은 30MB 이하여야 합니다.')
        original_name, original_type = original.filename[:255], original.content_type or 'application/octet-stream'
    canonical_json = _validate_json_upload(json_file, json_text, required=False)
    now = datetime.now().isoformat()
    with get_db() as conn:
        cursor = conn.execute('''INSERT INTO final_deliverables
            (kind, school, grade, year, semester, exam, title, original_name, original_type,
             original_blob, html_text, json_text, created_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)''',
            (kind, school.strip(), grade.strip(), year, semester.strip(), exam.strip(),
             title.strip()[:150], original_name, original_type, original_bytes, canonical_html, canonical_json, now))
        row = conn.execute('''SELECT id,kind,school,grade,year,semester,exam,title,original_name,created_at,
            (CASE WHEN html_text <> '' THEN 1 ELSE 0 END) AS has_html
            FROM final_deliverables WHERE id=?''', (cursor.lastrowid,)).fetchone()
        return dict(row)


@app.get('/api/final-deliverables')
def list_final_deliverables(school: str, grade: str, year: int, semester: str, exam: str, kind: Optional[str] = None):
    with get_db() as conn:
        query = '''SELECT id,kind,school,grade,year,semester,exam,title,original_name,created_at,
            (CASE WHEN html_text <> '' THEN 1 ELSE 0 END) AS has_html
            FROM final_deliverables WHERE school=? AND grade=? AND year=? AND semester=? AND exam=?'''
        values: list = [school, grade, year, semester, exam]
        if kind:
            query += ' AND kind=?'
            values.append(kind)
        query += ' ORDER BY id DESC'
        rows = conn.execute(query, values).fetchall()
        return [dict(row) for row in rows]


@app.get('/api/final-deliverables/{deliverable_id}/original')
def open_final_deliverable_original(deliverable_id: int):
    with get_db() as conn:
        row = conn.execute('SELECT original_blob,original_type,original_name FROM final_deliverables WHERE id=?', (deliverable_id,)).fetchone()
        if not row or not row['original_name']: raise HTTPException(status_code=404, detail='자료를 찾지 못했습니다.')
        safe_name = re.sub(r'[^0-9A-Za-z._-]', '_', row['original_name'])
        return Response(row['original_blob'], media_type=row['original_type'], headers={'Content-Disposition': f'inline; filename="{safe_name}"'})


@app.get('/api/final-deliverables/{deliverable_id}/html')
def open_final_deliverable_html(deliverable_id: int):
    with get_db() as conn:
        row = conn.execute('SELECT html_text FROM final_deliverables WHERE id=?', (deliverable_id,)).fetchone()
        if not row: raise HTTPException(status_code=404, detail='자료를 찾지 못했습니다.')
        return Response(row['html_text'] or '', media_type='text/html; charset=utf-8')


@app.get('/api/final-deliverables/{deliverable_id}/json')
def open_final_deliverable_json(deliverable_id: int):
    with get_db() as conn:
        row = conn.execute('SELECT json_text FROM final_deliverables WHERE id=?', (deliverable_id,)).fetchone()
        if not row: raise HTTPException(status_code=404, detail='자료를 찾지 못했습니다.')
        return Response(row['json_text'], media_type='application/json; charset=utf-8')


# ==================================================================================
# Digital Twin — HTML 업로드 기반 (AI 분석/Digital Twin 자동 생성 기능은 완전히 삭제됨)
# ==================================================================================

def sanitize_profile_text(value: object) -> str:
    if not isinstance(value, str):
        return ''
    cleaned = re.sub(r'<style[\s\S]*?</style>|<script[\s\S]*?</script>', ' ', value, flags=re.IGNORECASE)
    cleaned = re.sub(r':root\s*\{[^}]*\}|--[\w-]+\s*:[^;{}]+;?', ' ', cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r'<[^>]*>', ' ', cleaned)
    cleaned = re.sub(r'\s+', ' ', cleaned).strip()
    if re.search(r'(^|\s)(:root|--[\w-]+|\.\w+\s*\{|@media|display\s*:|color\s*:|background\s*:|학교명 미확인)', cleaned, re.IGNORECASE):
        return ''
    return cleaned


def sanitize_profile_data(data: dict) -> dict:
    cleaned = dict(data)
    cleaned['school'] = sanitize_profile_text(cleaned.get('school')) or '업로드 자료'
    for key in ('grade', 'exam', 'dna', 'description', 'difficulty', 'updatedAt', 'choiceStyle', 'writingStyle'):
        cleaned[key] = sanitize_profile_text(cleaned.get(key))
    cleaned['rules'] = [sanitize_profile_text(rule) for rule in cleaned.get('rules', []) if sanitize_profile_text(rule)]
    cleaned['trapPatterns'] = [sanitize_profile_text(rule) for rule in cleaned.get('trapPatterns', []) if sanitize_profile_text(rule)]
    cleaned['types'] = [
        {**item, 'name': sanitize_profile_text(item.get('name')), 'description': sanitize_profile_text(item.get('description')), 'ratio': sanitize_profile_text(item.get('ratio')), 'difficulty': sanitize_profile_text(item.get('difficulty'))}
        for item in cleaned.get('types', []) if isinstance(item, dict) and sanitize_profile_text(item.get('name'))
    ]
    return cleaned


class ImportedTwinProfile(BaseModel):
    school: str
    grade: str = ''
    exam: str = ''
    dna: str = ''
    description: str = ''
    types: list[dict] = []
    rules: list[str] = []
    difficulty: str = ''
    questionRatio: dict = {}
    trapPatterns: list[str] = []
    choiceStyle: str = ''
    writingStyle: str = ''
    updatedAt: str = ''
    rawSource: str = ''


def digital_twin_response(row: sqlite3.Row) -> dict:
    """Return the persisted row contract, not potentially stale profile_json metadata."""
    profile = sanitize_profile_data(json.loads(row['profile_json']))
    return {
        **profile,
        'id': row['id'],
        'school': sanitize_profile_text(row['school']) or profile.get('school', ''),
        'grade': sanitize_profile_text(row['grade']),
        'exam': sanitize_profile_text(row['exam']),
        'version': row['version'],
        'updatedAt': row['updated_at'],
    }


@app.get('/api/digital-twins')
def list_digital_twins(school: Optional[str] = None, grade: Optional[str] = None, exam: Optional[str] = None):
    query = 'SELECT * FROM digital_twin_profiles WHERE is_active = 1'
    values: list[str] = []
    if school:
        query += ' AND school = ?'
        values.append(school)
    if grade:
        query += ' AND grade = ?'
        values.append(grade)
    if exam:
        query += ' AND exam = ?'
        values.append(exam)
    query += ' ORDER BY updated_at DESC'
    with get_db() as conn:
        rows = conn.execute(query, values).fetchall()
        return [digital_twin_response(row) for row in rows]


@app.get('/api/digital-twins/{twin_id}')
def get_digital_twin(twin_id: int):
    with get_db() as conn:
        row = conn.execute('SELECT * FROM digital_twin_profiles WHERE id = ?', (twin_id,)).fetchone()
    if not row:
        raise HTTPException(status_code=404, detail='Digital Twin을 찾지 못했습니다.')
    return digital_twin_response(row)


@app.post('/api/digital-twins/import')
def import_digital_twin(profile: ImportedTwinProfile):
    """강사가 만든 Digital Twin HTML을 업로드하면 파싱된 정보를 저장한다.
    이 저장을 기준으로 학교 → 학년 → 시험 계층이 자동으로 생성된다 (별도 관리자 입력 불필요)."""
    now = datetime.now().strftime('%Y.%m.%d %H:%M')
    data = sanitize_profile_data(profile.model_dump())
    data['school'] = data['school'] or '업로드 자료'
    data['types'] = data['types'] or [{'name': '기출 분석 대기', 'description': '업로드 자료의 문항 분석이 필요합니다.', 'ratio': '', 'difficulty': ''}]
    data['rules'] = data['rules'] or ['업로드 자료의 문항 구성과 난이도를 기준으로 생성']
    data['updatedAt'] = now
    data['rawSource'] = (data.get('rawSource') or '')[:50000]
    with get_db() as conn:
        latest = conn.execute('SELECT COALESCE(MAX(version), 0) AS version FROM digital_twin_profiles WHERE school = ? AND grade = ? AND exam = ?', (data['school'], data['grade'], data['exam'])).fetchone()
        version = int(latest['version']) + 1
        conn.execute('UPDATE digital_twin_profiles SET is_active = 0 WHERE school = ? AND grade = ? AND exam = ?', (data['school'], data['grade'], data['exam']))
        cursor = conn.execute('INSERT INTO digital_twin_profiles (school, grade, exam, profile_json, version, is_active, created_at, updated_at) VALUES (?, ?, ?, ?, ?, 1, ?, ?)', (data['school'], data['grade'], data['exam'], json.dumps(data, ensure_ascii=False), version, now, now))
        return {**data, 'id': cursor.lastrowid, 'version': version}


@app.delete('/api/digital-twins/{twin_id}')
def delete_digital_twin(twin_id: int):
    with get_db() as conn:
        conn.execute('UPDATE digital_twin_profiles SET is_active = 0 WHERE id = ?', (twin_id,))
    return {'ok': True}


@app.get('/api/schools')
def school_hierarchy():
    """기존 SQLite 전체 기록에서 학교 → 학년 → 시험 계층을 만든다.
    Digital Twin만 조회하던 이전 구현 때문에 학생/문제 기록의 학교가 새 UI에서 누락됐다."""
    with get_db() as conn:
        twin_rows = conn.execute('SELECT id, school, grade, exam, version, updated_at FROM digital_twin_profiles WHERE is_active = 1').fetchall()
        question_rows = conn.execute("SELECT DISTINCT school, grade, exam FROM generated_questions WHERE TRIM(school) <> ''").fetchall()
        student_rows = conn.execute("SELECT DISTINCT school, grade, exam_type AS exam FROM students WHERE TRIM(school) <> ''").fetchall()
    tree: dict[str, dict[str, dict]] = {}
    for row in twin_rows:
        school = sanitize_profile_text(row['school']) or '학교명 미확인'
        grade = sanitize_profile_text(row['grade']) or '학년 미지정'
        exam = sanitize_profile_text(row['exam']) or '시험 미지정'
        tree.setdefault(school, {}).setdefault(grade, {})
        tree[school][grade][exam] = {'twinId': row['id'], 'version': row['version'], 'updatedAt': row['updated_at'], 'hasDigitalTwin': True}
    for row in [*question_rows, *student_rows]:
        school = sanitize_profile_text(row['school']) or '학교명 미확인'
        grade = sanitize_profile_text(row['grade']) or '학년 미지정'
        exam = sanitize_profile_text(row['exam']) or '시험 미지정'
        tree.setdefault(school, {}).setdefault(grade, {}).setdefault(
            exam, {'twinId': None, 'version': None, 'updatedAt': '', 'hasDigitalTwin': False}
        )
    result = []
    for school, grades in tree.items():
        grade_list = []
        for grade, exams in grades.items():
            exam_list = [{'exam': exam, **info} for exam, info in exams.items()]
            grade_list.append({'grade': grade, 'exams': exam_list})
        result.append({'school': school, 'grades': grade_list})
    return result


# ==================================================================================
# BuildIt Assistant — 기존 Wiki 기억을 최소 범위로 조회하는 대화형 작업 안내.
# 세션 state/history는 브라우저가 보유하며 서버는 DB를 수정하지 않는다.
# ==================================================================================

class AssistantState(BaseModel):
    goal: str = ''
    school: str = ''
    grade: str = ''
    exam: str = ''
    selected_twin_id: Optional[int] = None
    needed_materials: list[str] = Field(default_factory=list)
    recent_request: str = ''
    current_page: str = 'state'


class AssistantMessage(BaseModel):
    role: Literal['user', 'assistant']
    content: str


class AssistantAttachment(BaseModel):
    name: str
    type: str = ''


class AssistantRequest(BaseModel):
    message: str
    api_key: str = ''
    openrouter_api_key: str = ''
    provider: str = 'auto'
    state: AssistantState = Field(default_factory=AssistantState)
    history: list[AssistantMessage] = Field(default_factory=list)
    attachments: list[AssistantAttachment] = Field(default_factory=list)


def _assistant_entities(message: str, state: AssistantState) -> dict:
    school_match = re.search(r'([가-힣A-Za-z0-9]+(?:중학교|중|고등학교|고))(?=\s|\d|$)', message)
    grade_match = re.search(r'(?:중\s*)?([1-3])\s*학년', message) or re.search(r'(?:중학교|중)\s*([1-3])(?=\s|학년|문제|시험|$)', message)
    exam_match = re.search(r'((?:1|2)학기\s*)?(중간고사|기말고사|중간|기말)', message)
    school = school_match.group(1) if school_match else state.school
    grade = f'중{grade_match.group(1)}' if grade_match else state.grade
    exam = (exam_match.group(0).replace(' ', '') if exam_match else state.exam)
    if exam == '중간': exam = '중간고사'
    if exam == '기말': exam = '기말고사'
    goal = state.goal
    if re.search(r'문제|시험지|내신', message): goal = '학교별 내신 문제 제작'
    elif re.search(r'오답|학생', message): goal = '학생 오답 정리'
    elif re.search(r'인사이트|수업\s*기록', message): goal = '수업 인사이트 기록'
    elif re.search(r'Digital\s*Twin|디지털\s*트윈|기출\s*분석', message, re.I): goal = 'Digital Twin 구축'
    return {'school': school, 'grade': grade, 'exam': exam, 'goal': goal or '요청한 작업 준비'}


def _assistant_grade_number(grade: str) -> str:
    """UI/대화의 '중3'과 Digital Twin 저장값 '3학년'을 같은 학년으로 조회한다."""
    match = re.fullmatch(r'(?:중)?\s*([1-3])\s*(?:학년)?', (grade or '').strip())
    return match.group(1) if match else ''


def _assistant_context(entities: dict) -> dict:
    school, grade, exam = entities['school'], entities['grade'], entities['exam']
    grade_number = _assistant_grade_number(grade)
    context = {'twin': None, 'availableExams': [], 'questionTypes': [], 'questionCount': 0, 'studentCount': 0, 'insightCount': 0}
    if not school:
        return context
    with get_db() as conn:
        twin_query = 'SELECT * FROM digital_twin_profiles WHERE is_active=1 AND school=?'
        twin_values: list = [school]
        if grade_number:
            twin_query += " AND REPLACE(REPLACE(grade, '중', ''), '학년', '')=?"; twin_values.append(grade_number)
        if exam:
            twin_query += " AND (exam=? OR exam='전체')"; twin_values.append(exam)
        twin_query += " ORDER BY CASE WHEN exam=? THEN 0 ELSE 1 END, updated_at DESC LIMIT 1"
        twin_values.append(exam)
        twin = conn.execute(twin_query, twin_values).fetchone()
        if twin:
            context['twin'] = digital_twin_response(twin)
        exam_rows = conn.execute(
            "SELECT DISTINCT exam FROM digital_twin_profiles WHERE is_active=1 AND school=? AND (? = '' OR REPLACE(REPLACE(grade, '중', ''), '학년', '')=?) ORDER BY exam",
            (school, grade_number, grade_number),
        ).fetchall()
        context['availableExams'] = [row['exam'] for row in exam_rows if row['exam']]
        if context['twin']:
            context['questionTypes'] = [item.get('name', '') if isinstance(item, dict) else str(item) for item in context['twin'].get('types', []) if item]
        filters, values = ['school=?'], [school]
        if grade_number: filters.append("REPLACE(REPLACE(grade, '중', ''), '학년', '')=?"); values.append(grade_number)
        if exam: filters.append('exam=?'); values.append(exam)
        context['questionCount'] = conn.execute(f"SELECT COUNT(*) n FROM generated_questions WHERE {' AND '.join(filters)}", values).fetchone()['n']
        student_filters, student_values = ['school=?'], [school]
        if grade_number: student_filters.append("REPLACE(REPLACE(grade, '중', ''), '학년', '')=?"); student_values.append(grade_number)
        context['studentCount'] = conn.execute(f"SELECT COUNT(*) n FROM students WHERE {' AND '.join(student_filters)}", student_values).fetchone()['n']
        insight_filters, insight_values = ['school=?'], [school]
        if grade_number: insight_filters.append("REPLACE(REPLACE(grade, '중', ''), '학년', '')=?"); insight_values.append(grade_number)
        context['insightCount'] = conn.execute(f"SELECT COUNT(*) n FROM educator_insights WHERE {' AND '.join(insight_filters)}", insight_values).fetchone()['n']
    return context


def _attachment_route(attachments: list[AssistantAttachment], message: str, goal: str) -> tuple[str, str]:
    names = ' '.join(item.name.lower() for item in attachments)
    if re.search(r'\.(html?|json|txt)$', names): return 'twin', 'Digital Twin 가져오기'
    if re.search(r'오답|학생', message + goal): return 'wrongnote', '학생 오답 OCR 열기'
    if attachments and re.search(r'\.(pdf|png|jpe?g|hwp|hwpx|docx)$', names): return 'school', '학교 프린트 OCR 열기'
    if goal == '학교별 내신 문제 제작': return 'home', '문제 제작 화면 열기'
    if goal == '수업 인사이트 기록': return 'brain', '수업 인사이트 열기'
    if goal == 'Digital Twin 구축': return 'twin', 'Digital Twin 열기'
    return 'state', '현재 상태 보기'


def _assistant_fallback(entities: dict, context: dict, attachments: list[AssistantAttachment]) -> tuple[str, list[str]]:
    school, grade, exam, goal = entities['school'], entities['grade'], entities['exam'], entities['goal']
    target = ' '.join(part for part in (school, grade, exam) if part) or '현재 요청'
    lines = [f'{target}의 목표를 “{goal}”으로 정리했습니다.']
    twin = context['twin']
    if school and twin:
        type_text = ', '.join(context['questionTypes'][:5]) or '저장된 학교별 출제 패턴'
        exam_text = ', '.join(context['availableExams']) or '시험 미지정'
        lines.append(f'{school} {grade} Digital Twin을 확인했습니다. 연결 가능한 시험은 {exam_text}, 출제 유형은 {type_text}입니다. 저장 문제 {context["questionCount"]}개와 학생·수업 기록도 연결할 수 있습니다.')
    elif school:
        lines.append(f'{school} {grade}에 연결된 Digital Twin을 찾지 못했습니다. 기존 학교 기록은 보존한 채 기출 또는 Digital Twin 자료를 먼저 추가할 수 있습니다.')
    else:
        lines.append('연결할 기억을 찾으려면 학교·학년·시험을 알려 주세요.')
    needed = []
    if goal == '학교별 내신 문제 제작':
        if not exam: needed.append('시험')
        if not attachments: needed.extend(['시험 범위 지문 또는 학교 프린트', '시험 범위 주요 문법 사항'])
    if attachments:
        lines.append(f'첨부한 {", ".join(item.name for item in attachments)}은 기존 업로드 도구에서 처리하도록 연결하겠습니다. 파일 내용은 이 대화 API에 저장하지 않습니다.')
    elif needed:
        if goal == '학교별 내신 문제 제작':
            lines.append('이번 시험 범위 지문이나 학교 프린트와 시험 범위 주요 문법 사항이 필요합니다.' + (' 시험도 선택해 주세요.' if not exam else '') + ' 알려 주시면 저장된 출제 유형을 바로 적용하겠습니다.')
        else:
            lines.append('아직 필요한 것은 ' + ', '.join(needed) + '입니다.')
    if goal == '학교별 내신 문제 제작':
        lines.append('준비가 되면 학교·학년·Twin이 채워진 문제 제작 화면을 열겠습니다.')
    return '\n\n'.join(lines), needed


@app.post('/api/assistant/chat')
def assistant_chat(payload: AssistantRequest):
    message = payload.message.strip()
    if not message and not payload.attachments:
        raise HTTPException(status_code=400, detail='요청이나 첨부 자료를 입력해 주세요.')
    entities = _assistant_entities(message, payload.state)
    context = _assistant_context(entities)
    fallback, needed = _assistant_fallback(entities, context, payload.attachments)
    route, action = _attachment_route(payload.attachments, message, entities['goal'])
    text, model = fallback, 'rules-fallback'
    if payload.api_key.strip() or payload.openrouter_api_key.strip() or os.getenv('GEMINI_API_KEY', '').strip() or os.getenv('OPENROUTER_API_KEY', '').strip():
        minimal_context = json.dumps({'goal': entities['goal'], 'school': entities['school'], 'grade': entities['grade'], 'exam': entities['exam'], **context, 'attachments': [a.name for a in payload.attachments]}, ensure_ascii=False)
        recent = payload.history[-6:]
        prompt = f'확인된 컨텍스트: {minimal_context}\n최근 대화: ' + json.dumps([m.model_dump() for m in recent], ensure_ascii=False) + f'\n사용자 요청: {message}\n부족한 항목: {json.dumps(needed, ensure_ascii=False)}. 확인된 사실만 사용해 자연스럽게 답하고, 정적 기능 목록을 반복하지 마라.'
        try:
            result = ai_router.generate(payload.provider, '너는 BuildIt Wiki의 한국어 작업 Assistant다. 링크만 답하지 말고 목표를 요약하고, 확인된 기억과 부족한 자료, 작업 순서를 짧고 친절하게 설명한다. 제공된 최소 컨텍스트 밖의 DB 사실은 추측하지 않는다. 학생 이름·연락처 등 개인정보를 요청하거나 추측하지 않는다. DB 수정이나 저장을 완료했다고 말하지 않는다.', prompt, temperature=0.25, allow_noncritical_fallback=True)
            text, model = result.text, result.provider
        except Exception as exc:
            if payload.provider in ('gemini', 'ox-alpha'):
                _handle_gemini_error(exc)
    next_state = {**payload.state.model_dump(), **entities, 'selected_twin_id': context['twin']['id'] if context['twin'] else None, 'needed_materials': needed, 'recent_request': message}
    return {'text': text, 'state': next_state, 'context': context, 'page': route, 'action': action, 'provider': result.provider if 'result' in locals() else 'rules-fallback', 'model': result.model if 'result' in locals() else model}


# ==================================================================================
# 문제 생성 / 기존 문제 변형 — 기존 HTML의 AI 프롬프트 로직을 그대로 사용 (prompt_logic.py)
# ==================================================================================

def _load_twin(conn, twin_id: Optional[int], school: Optional[str], grade: Optional[str], exam: Optional[str]):
    row = None
    if twin_id:
        row = conn.execute('SELECT * FROM digital_twin_profiles WHERE id = ? AND is_active = 1', (twin_id,)).fetchone()
    elif school:
        query = 'SELECT * FROM digital_twin_profiles WHERE is_active = 1 AND school = ?'
        values = [school]
        if grade:
            query += ' AND grade = ?'
            values.append(grade)
        if exam:
            query += ' AND exam = ?'
            values.append(exam)
        query += ' ORDER BY updated_at DESC LIMIT 1'
        row = conn.execute(query, values).fetchone()
    if not row:
        raise HTTPException(status_code=404, detail='선택한 학교·학년·시험의 Digital Twin을 찾지 못했습니다. Digital Twin 관리에서 HTML을 먼저 업로드해 주세요.')
    return digital_twin_response(row)


def _save_generated(conn, twin: Optional[dict], mode: str, questions: list[dict], exam_title: str, exam_group_id: str,
                    test_id: Optional[int] = None):
    """생성/변형된 문항을 DB에 저장하고, 프런트엔드(GET /api/questions와 동일한 snake_case 형태)로
    정규화한 결과를 반환한다. AI 원본 응답의 camelCase 키를 그대로 돌려주면 프런트 표시가 깨지므로 여기서 통일한다."""
    now = datetime.now().isoformat()
    twin_school = twin.get('school', '') if twin else ''
    twin_grade = twin.get('grade', '') if twin else ''
    twin_exam = twin.get('exam', '') if twin else ''
    saved = []
    for q in questions:
        q_kind = q.get('type') or ('sa' if 'conditions' in q else 'mc')
        options = q.get('options', []) or []
        answer = q.get('answer', []) or []
        conditions = q.get('conditions', []) or []
        accepted_answers = q.get('acceptedAnswers', []) or []
        question_type = q.get('questionType', '')
        stem = q.get('stem', '')
        source_text = q.get('sourceText', '')
        multi = bool(q.get('multi'))
        score = int(q.get('score') or 0)
        answer_format = q.get('answerFormat', '')
        sample_answer = q.get('sampleAnswer', '')
        grading_rule = q.get('gradingRule', '')
        row_data = (
            twin.get('id') if twin else None, twin_school, twin_grade, twin_exam, mode, q_kind,
            question_type, stem, source_text,
            json.dumps(options, ensure_ascii=False), json.dumps(answer, ensure_ascii=False),
            1 if multi else 0, score,
            json.dumps(conditions, ensure_ascii=False), answer_format,
            json.dumps(accepted_answers, ensure_ascii=False), sample_answer, grading_rule,
            exam_title, exam_group_id, json.dumps(q, ensure_ascii=False), now, test_id,
        )
        cursor = conn.execute(
            '''INSERT INTO generated_questions
            (twin_id, school, grade, exam, mode, q_kind, question_type, stem, source_text, options_json, answer_json,
             multi, score, conditions_json, answer_format, accepted_answers_json, sample_answer, grading_rule,
             exam_title, exam_group_id, raw_json, created_at, test_id)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',
            row_data,
        )
        saved.append({
            'id': cursor.lastrowid, 'twin_id': twin.get('id') if twin else None,
            'school': twin_school, 'grade': twin_grade, 'exam': twin_exam,
            'mode': mode, 'q_kind': q_kind, 'question_type': question_type,
            'stem': stem, 'source_text': source_text, 'options': options, 'answer': answer,
            'multi': 1 if multi else 0, 'score': score, 'conditions': conditions,
            'answer_format': answer_format, 'accepted_answers': accepted_answers,
            'sample_answer': sample_answer, 'grading_rule': grading_rule,
            'exam_title': exam_title, 'exam_group_id': exam_group_id, 'favorite': 0, 'created_at': now,
            'test_id': test_id,
        })
    return saved


def _handle_gemini_error(exc):
    if isinstance(exc, ai_router.ProviderError):
        status = exc.status if exc.status in (401, 403, 429, 502, 503, 504) else 502
        message = str(exc)
        if status == 429:
            message = 'AI 요청이 많아 잠시 사용할 수 없습니다. 잠시 후 다시 시도해 주세요.'
        elif status == 504:
            message = 'AI 응답 시간이 초과되었습니다. 잠시 후 다시 시도해 주세요.'
        elif status in (502, 503) and exc.provider == 'ox-alpha':
            message = 'Ox Alpha 연결에 실패했습니다. 잠시 후 다시 시도해 주세요.'
        raise HTTPException(status_code=status, detail={'message': message, 'provider': exc.provider, 'status': status})
    if isinstance(exc, pl.GeminiCallError):
        if exc.http_status in (401, 403):
            raise HTTPException(status_code=401, detail='잘못된 API Key입니다. 설정에서 키를 다시 확인해 주세요.')
        raise HTTPException(status_code=502, detail=f'AI 요청을 완료하지 못했습니다: {exc}')
    raise HTTPException(status_code=502, detail='AI 서비스 연결을 확인하지 못했습니다. 네트워크 상태를 확인해 주세요.')


def _parse_generated_json(text: str, repair):
    """Repair only malformed JSON, once. Schema/quality failures are never retried here."""
    try:
        return json.loads(re.sub(r'```json|```', '', text).strip())
    except (json.JSONDecodeError, TypeError):
        repaired = repair()
        try:
            return json.loads(re.sub(r'```json|```', '', repaired).strip())
        except (json.JSONDecodeError, TypeError):
            raise HTTPException(status_code=502, detail='AI 응답 JSON 파싱 실패(repair 1회 포함)')


def _contract_error(exc: ValidationError):
    raise HTTPException(status_code=502, detail='AI 응답 schema validation 실패: ' + qc.validation_fields(exc))


class NewQuestionRequest(BaseModel):
    api_key: str = ''
    openrouter_api_key: str = ''
    provider: str = 'auto'
    twin_id: Optional[int] = None
    school: Optional[str] = None
    grade: Optional[str] = None
    exam: Optional[str] = None
    passage: str
    grammar: str
    count: int = 10
    difficulty: Literal['actual', 'easy', 'hard', 'killer', 'writing'] = 'actual'
    composition: Literal['balanced', 'objective', 'writing'] = 'balanced'
    types: list[str] = []  # 사용자가 직접 유형을 지정하면 이 값을 사용, 비어 있으면 Digital Twin의 유형을 자동 적용


def _composition(total: int, mode: str):
    if mode == 'objective':
        sa = max(1, round(total * 0.15))
    elif mode == 'writing':
        sa = max(2, round(total * 0.35))
    else:
        sa = max(1, round(total * 0.20))
    return total - sa, sa


def _batches(total: int, mc: int, sa: int):
    arr = []
    left, start = total, 1
    while left > 0:
        size = min(5, left)
        bsa = min(sa, max(0, round(size * sa / left))) if left else 0
        if sa > 0 and left == size:
            bsa = sa
        bsa = min(bsa, size)
        bmc = size - bsa
        arr.append({'start': start, 'end': start + size - 1, 'mc': bmc, 'sa': bsa, 'size': size})
        start += size
        left -= size
        mc -= bmc
        sa -= bsa
    return arr


@app.post('/api/questions/generate')
def generate_questions(payload: NewQuestionRequest):
    if not payload.passage.strip() or not payload.grammar.strip():
        raise HTTPException(status_code=400, detail='새 영어 지문과 핵심 문법을 입력해 주세요.')
    with get_db() as conn:
        twin = _load_twin(conn, payload.twin_id, payload.school, payload.grade, payload.exam)
        type_labels = [
            f"{pl.TYPE_DEF_BY_ID[type_id]['label']} — {pl.TYPE_DEF_BY_ID[type_id]['prompt']}"
            for type_id in payload.types if type_id in pl.TYPE_DEF_BY_ID
        ] or None
        mc_total, sa_total = _composition(payload.count, payload.composition)
        batches = _batches(payload.count, mc_total, sa_total)
        all_questions: list[dict] = []
        previous = ''
        for batch in batches:
            rules = pl.build_new_rules(twin, {
                'type_labels': type_labels, 'difficulty': payload.difficulty,
                'start': batch['start'], 'end': batch['end'], 'mc': batch['mc'], 'sa': batch['sa'], 'previous': previous,
            })
            user_msg = f"[시험 범위 주요 문법—문법 문항과 서답형 문항에 최우선 반영]\n{payload.grammar}\n\n[새 지문—내용 일치 문제에서는 이 전체 지문을 제시할 것]\n{payload.passage}"
            batch_questions: list[dict] = []
            quality_issues: list[str] = []
            for quality_attempt in range(2):
                retry_rules = rules
                if quality_issues:
                    retry_rules += ' [직전 응답 품질 오류—모두 수정] ' + ' / '.join(quality_issues)
                try:
                    result = ai_router.generate(payload.provider, retry_rules, user_msg, schema=pl.schema_new(batch['mc'], batch['sa']), temperature=0.35, gemini_key=payload.api_key, openrouter_key=payload.openrouter_api_key)
                    text = result.text
                except Exception as exc:
                    _handle_gemini_error(exc)
                    return
                def repair_json():
                    repair_prompt = '다음 응답의 내용은 바꾸지 말고 유효한 JSON 객체로만 복구하라:\n' + text
                    return ai_router.generate(payload.provider, retry_rules, repair_prompt, schema=pl.schema_new(batch['mc'], batch['sa']), temperature=0, gemini_key=payload.api_key, openrouter_key=payload.openrouter_api_key).text
                parsed = _parse_generated_json(text, repair_json)
                try:
                    parsed = qc.validate_new(parsed)
                except ValidationError as exc:
                    _contract_error(exc)
                mc_qs = [pl.prepare_question({**q, 'type': 'mc'}) for q in parsed.get('mcQuestions', [])]
                sa_qs = []
                for q in parsed.get('saQuestions', []):
                    q['type'] = 'sa'
                    q['acceptedAnswers'] = q.get('acceptedAnswers') or [q.get('sampleAnswer', '')]
                    q['answerFormat'] = q.get('answerFormat') or '정답 한 개를 정확히 작성'
                    q['gradingRule'] = q.get('gradingRule') or '철자와 문법이 모범답안과 일치해야 함'
                    sa_qs.append(q)
                batch_questions = mc_qs + sa_qs
                for question in batch_questions:
                    question['grammarFocus'] = payload.grammar.strip()
                quality_issues = pl.batch_quality_issues(batch_questions)
                if not quality_issues:
                    break
            if quality_issues:
                raise HTTPException(status_code=502, detail='AI가 시험지 형식 검증을 통과하지 못했습니다: ' + '; '.join(quality_issues[:3]))
            all_questions.extend(batch_questions)
            previous = ' || '.join(f"{i+1}. {q.get('questionType','')} / {q.get('stem','')} / {(q.get('sourceText') or '')[:120]}" for i, q in enumerate(all_questions))[-7000:]
        exam_title = f"{twin.get('school','')} {twin.get('grade','')} 맞춤형 시험지".strip()
        exam_group_id = f"exam_{int(datetime.now().timestamp())}"
        saved = _save_generated(conn, twin, 'new', all_questions, exam_title, exam_group_id)
        return {'examTitle': exam_title, 'examGroupId': exam_group_id, 'questions': saved, 'provider': result.provider, 'model': result.model, 'twin': {'id': twin['id'], 'school': twin.get('school', ''), 'version': twin['version']}}


class AbComparisonRequest(NewQuestionRequest):
    provider: str = 'auto'


def _ab_generate(provider: str, payload: AbComparisonRequest, twin: dict) -> dict:
    mc, sa = _composition(payload.count, payload.composition)
    type_labels = [f"{pl.TYPE_DEF_BY_ID[t]['label']} — {pl.TYPE_DEF_BY_ID[t]['prompt']}" for t in payload.types if t in pl.TYPE_DEF_BY_ID] or None
    rules = pl.build_new_rules(twin, {'type_labels': type_labels, 'difficulty': payload.difficulty, 'start': 1, 'end': payload.count, 'mc': mc, 'sa': sa, 'previous': ''})
    prompt = f"[시험 범위 주요 문법]\n{payload.grammar}\n\n[새 지문]\n{payload.passage}"
    result = ai_router.generate(provider, rules, prompt, schema=pl.schema_new(mc, sa), temperature=0.35, gemini_key=payload.api_key, openrouter_key=payload.openrouter_api_key)
    parsed = json.loads(re.sub(r'```json|```', '', result.text).strip())
    questions = [pl.prepare_question({**q, 'type': 'mc'}) for q in parsed.get('mcQuestions', [])]
    questions += [{**q, 'type': 'sa'} for q in parsed.get('saQuestions', [])]
    issues = pl.batch_quality_issues(questions)
    if issues:
        raise ai_router.ProviderError('JSON 문제 구조 검증 실패: ' + '; '.join(issues[:3]), provider)
    return {'resultId': uuid.uuid4().hex, 'provider': result.provider, 'model': result.model, 'questions': questions}


@app.post('/api/questions/ab-compare')
def compare_questions(payload: AbComparisonRequest):
    if not payload.passage.strip() or not payload.grammar.strip():
        raise HTTPException(status_code=400, detail='새 영어 지문과 핵심 문법을 입력해 주세요.')
    with get_db() as conn:
        twin = _load_twin(conn, payload.twin_id, payload.school, payload.grade, payload.exam)
    results, errors = {}, {}
    for provider in ('gemini', 'ox-alpha'):
        try:
            results[provider] = _ab_generate(provider, payload, twin)
        except Exception as exc:
            errors[provider] = str(exc)
    if not results:
        raise HTTPException(status_code=502, detail='두 provider 모두 비교 결과를 생성하지 못했습니다.')
    return {'contextContract': {'school': twin.get('school', ''), 'grade': twin.get('grade', ''), 'exam': twin.get('exam', ''), 'twinId': twin['id'], 'types': payload.types, 'grammar': payload.grammar}, 'results': results, 'errors': errors}


class AbEvaluationCreate(BaseModel):
    school: str = ''
    grade: str = ''
    exam: str = ''
    twin_id: Optional[int] = None
    question_type: str = ''
    gemini_model: str
    gemini_result_id: str
    ox_model: str
    ox_result_id: str
    choice: Literal['gemini', 'ox-alpha', 'both-revise']


@app.post('/api/questions/ab-evaluations', status_code=201)
def save_ab_evaluation(payload: AbEvaluationCreate):
    now = datetime.now().isoformat()
    with get_db() as conn:
        cursor = conn.execute('''INSERT INTO ai_ab_evaluations
            (school,grade,exam,twin_id,question_type,gemini_provider,gemini_model,gemini_result_id,ox_provider,ox_model,ox_result_id,choice,created_at)
            VALUES (?,?,?,?,?,'gemini',?,?,'ox-alpha',?,?,?,?)''',
            (payload.school, payload.grade, payload.exam, payload.twin_id, payload.question_type, payload.gemini_model, payload.gemini_result_id, payload.ox_model, payload.ox_result_id, payload.choice, now))
        row = conn.execute('SELECT * FROM ai_ab_evaluations WHERE id=?', (cursor.lastrowid,)).fetchone()
        return dict(row)


@app.get('/api/questions/ab-evaluations')
def list_ab_evaluations(limit: int = 100):
    with get_db() as conn:
        return [dict(row) for row in conn.execute('SELECT * FROM ai_ab_evaluations ORDER BY id DESC LIMIT ?', (max(1, min(limit, 500)),)).fetchall()]


class TransformQuestionRequest(BaseModel):
    api_key: str = ''
    openrouter_api_key: str = ''
    provider: str = 'auto'
    twin_id: Optional[int] = None
    school: Optional[str] = None
    grade: Optional[str] = None
    exam: Optional[str] = None
    original: str
    passage: str = ''
    count: int = 3
    strength: Literal['light', 'normal', 'strong'] = 'normal'
    target: Literal['same', 'related', 'different'] = 'same'
    options: bool = True
    answerPos: bool = True
    harder: bool = False
    schoolStyle: bool = True
    desiredType: Optional[str] = None  # TYPE_DEFS id, 사용자가 "원하는 유형으로 변경"을 직접 선택한 경우


@app.get('/api/question-types')
def question_types():
    """자유 유형 선택 UI용 문제 유형 카탈로그 (기존 HTML TYPE_DEFS)."""
    return pl.TYPE_DEFS


def _recommend_type(twin: dict) -> tuple[Optional[str], Optional[str]]:
    """항목 10: 유형 미지정 시 Digital Twin에서 가장 비중이 높은 유형을 추천."""
    types = twin.get('types') or []
    if not types:
        return None, None

    def ratio_value(item):
        raw = str(item.get('ratio', '') or '')
        m = re.search(r'\d+', raw)
        return int(m.group()) if m else 0

    best = sorted(types, key=ratio_value, reverse=True)[0]
    return best.get('name'), best.get('description')


@app.post('/api/questions/transform')
def transform_questions(payload: TransformQuestionRequest):
    """항목 7~10: 기존 HTML '기존 문제 변형' 기능 이식.
    학교 선택 → Digital Twin DB 읽기 → 학교 출제 DNA 적용 → 문제 변형.
    유형을 지정하지 않으면 Digital Twin을 보고 자동 추천한다."""
    if not payload.original.strip():
        raise HTTPException(status_code=400, detail='기존 문제를 입력해 주세요.')
    with get_db() as conn:
        twin = _load_twin(conn, payload.twin_id, payload.school, payload.grade, payload.exam)
        cfg = {
            'strength': payload.strength, 'target': payload.target, 'count': payload.count,
            'options': payload.options, 'answerPos': payload.answerPos, 'harder': payload.harder, 'schoolStyle': payload.schoolStyle,
        }
        if payload.desiredType:
            type_def = pl.TYPE_DEF_BY_ID.get(payload.desiredType)
            if type_def:
                cfg['desiredTypeLabel'] = type_def['label']
                cfg['desiredTypePrompt'] = type_def['prompt']
            else:
                cfg['desiredTypeLabel'] = payload.desiredType
        elif payload.target != 'same':
            recommended_name, recommended_desc = _recommend_type(twin)
            if recommended_name:
                cfg['recommended'] = recommended_name + (f'({recommended_desc})' if recommended_desc else '')
        rules = pl.build_transform_rules(twin, cfg)
        user_msg = f"[원본 문제]\n{payload.original}\n\n[관련 전체 지문]\n{payload.passage or '제공되지 않음'}"
        try:
            provider_result = ai_router.generate(payload.provider, rules, user_msg, schema=pl.schema_transform(payload.count), temperature=0.35, gemini_key=payload.api_key, openrouter_key=payload.openrouter_api_key)
            text = provider_result.text
        except Exception as exc:
            _handle_gemini_error(exc)
            return
        def repair_json():
            repair_prompt = '다음 응답의 내용은 바꾸지 말고 유효한 JSON 객체로만 복구하라:\n' + text
            return ai_router.generate(payload.provider, rules, repair_prompt, schema=pl.schema_transform(payload.count), temperature=0, gemini_key=payload.api_key, openrouter_key=payload.openrouter_api_key).text
        data = _parse_generated_json(text, repair_json)
        try:
            data = qc.validate_transform(data)
        except ValidationError as exc:
            _contract_error(exc)
        questions = []
        for q in data.get('questions', []):
            q['type'] = 'sa' if q.get('type') == 'sa' else 'mc'
            if q['type'] == 'mc':
                q = pl.repair_transform_question(q)
            else:
                q['conditions'] = q.get('conditions') or []
                q['acceptedAnswers'] = q.get('acceptedAnswers') or []
            questions.append(q)
        quality_issues = pl.batch_quality_issues(questions)
        if quality_issues:
            raise HTTPException(status_code=502, detail='AI가 변형문제 조건 검증을 통과하지 못했습니다: ' + '; '.join(quality_issues[:3]))
        exam_title = data.get('examTitle') or f"{twin.get('school','')} 기존 문제 변형"
        exam_group_id = f"transform_{int(datetime.now().timestamp())}"
        saved = _save_generated(conn, twin, 'transform', questions, exam_title, exam_group_id)
        return {'examTitle': exam_title, 'examGroupId': exam_group_id, 'questions': saved, 'provider': provider_result.provider, 'model': provider_result.model, 'twin': {'id': twin['id'], 'school': twin.get('school', ''), 'version': twin['version']}}


# ==================================================================================
# 강사 문제 DB — SQLite 영구 저장 (React 상태/로컬스토리지 저장 금지)
# ==================================================================================

@app.get('/api/questions')
def list_questions(school: Optional[str] = None, grade: Optional[str] = None, exam: Optional[str] = None,
                    mode: Optional[str] = None, search: Optional[str] = None, limit: int = 200,
                    test_id: Optional[int] = None):
    query = 'SELECT * FROM generated_questions WHERE 1=1'
    values: list = []
    for field, value in (('school', school), ('grade', grade), ('exam', exam), ('mode', mode), ('test_id', test_id)):
        if value is not None and value != '':
            query += f' AND {field} = ?'
            values.append(value)
    if search:
        query += ' AND (stem LIKE ? OR source_text LIKE ? OR question_type LIKE ?)'
        needle = f'%{search}%'
        values.extend([needle, needle, needle])
    query += ' ORDER BY id DESC LIMIT ?'
    values.append(max(1, min(limit, 1000)))
    with get_db() as conn:
        rows = conn.execute(query, values).fetchall()
    result = []
    for row in rows:
        item = dict(row)
        for jkey in ('options_json', 'answer_json', 'conditions_json', 'accepted_answers_json'):
            try:
                item[jkey.replace('_json', '')] = json.loads(item.pop(jkey))
            except (TypeError, json.JSONDecodeError):
                item[jkey.replace('_json', '')] = []
        item.pop('raw_json', None)
        result.append(item)
    return result


class ManualQuestionsCreate(BaseModel):
    """외부(Claude/MCP 등)에서 이미 제작한 문항을 사용자가 강사 문제 DB에 직접 붙여넣어 저장한다.
    AI 호출·Digital Twin 분석 없이 기존 question_contract + _save_generated만 재사용한다.
    source_material이 있으면 기존 create_generated_test 경로로 Generated Test까지 연결한다."""
    school: str = ''
    grade: str = ''
    year: Optional[int] = None
    semester: str = ''
    exam: str = '전체'
    exam_title: str = ''
    source_material: str = ''
    source_page: str = ''
    digital_twin: str = ''
    questions: list[dict]


@app.post('/api/questions', status_code=201)
def create_manual_questions(payload: ManualQuestionsCreate):
    if not payload.questions:
        raise HTTPException(status_code=422, detail='저장할 문항이 없습니다.')
    if payload.source_material.strip():
        title = payload.exam_title.strip() or f'{payload.source_material.strip()} 수동 저장 문항'
        result = create_generated_test(GeneratedTestCreate(
            school=payload.school, grade=payload.grade, source_material=payload.source_material,
            questions=payload.questions, source_page=payload.source_page, title=title,
            year=payload.year, semester=payload.semester, exam=payload.exam or '전체',
            digital_twin=payload.digital_twin, test_type='manual', generation_source='manual',
        ))
        return {'saved': True, 'question_count': result['question_count'],
                'examGroupId': result.get('external_id', ''), 'test_id': result.get('test_id'),
                'questions': result.get('questions', [])}
    title = payload.exam_title.strip() or f'{payload.school} {payload.grade} 수동 저장 문항'.strip()
    try:
        validated = qc.validate_transform({'examTitle': title, 'questions': payload.questions})
    except ValidationError as exc:
        raise HTTPException(status_code=422, detail='문항 검증 실패(저장하지 않음): ' + qc.validation_fields(exc)) from exc
    twin_like = {'id': None, 'school': payload.school.strip(), 'grade': payload.grade.strip(),
                 'exam': (payload.exam or '전체').strip()}
    group_id = f'manual_{uuid.uuid4().hex}'
    try:
        with get_db() as conn:
            saved = _save_generated(conn, twin_like, 'manual', validated['questions'], title, group_id)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f'문항을 저장하지 못했습니다: {exc}') from exc
    return {'saved': True, 'question_count': len(saved), 'examGroupId': group_id, 'questions': saved}


@app.delete('/api/questions/{question_id}')
def delete_question(question_id: int):
    with get_db() as conn:
        conn.execute('DELETE FROM generated_questions WHERE id = ?', (question_id,))
    return {'ok': True}


class FavoriteUpdate(BaseModel):
    favorite: bool


@app.patch('/api/questions/{question_id}/favorite')
def update_question_favorite(question_id: int, payload: FavoriteUpdate):
    with get_db() as conn:
        conn.execute('UPDATE generated_questions SET favorite = ? WHERE id = ?', (1 if payload.favorite else 0, question_id))
    return {'ok': True}


# ==================================================================================
# Generated Question Store — MCP가 생성한 시험문항을 시험(Test) 단위로 영구 저장한다.
#   Source Material(학습자료) → Generated Test(생성 시험) → Questions → Answer/Explanation.
# 새 저장 로직을 늘리지 않고 기존 question_contract.validate_transform + _save_generated를 재사용한다.
# 저장은 all-or-nothing이며, 문항 검증은 어떤 것도 쓰기 전에 먼저 한다(요구사항 5: 생성 성공 ≠ 저장 성공).
# ==================================================================================

class GeneratedTestCreate(BaseModel):
    school: str
    grade: str
    source_material: str
    questions: list[dict]
    source_page: str = ''
    test_type: str = ''
    title: str = ''
    year: Optional[int] = None
    semester: str = ''
    exam: str = '전체'
    digital_twin: str = ''
    twin_pattern: str = ''
    difficulty: str = ''
    target_vocabulary: list[str] = []
    target_grammar: list[str] = []
    target_passage: str = ''
    publisher: str = ''
    lesson: str = ''
    material_type: str = ''
    external_id: str = ''
    generation_source: str = 'mcp'  # 'mcp'(MCP 자동저장) | 'manual'(사용자 직접 붙여넣기)


def _derive_external_id(school: str, grade: str, source_material: str, source_page: str,
                        title: str, questions: list[dict]) -> str:
    """external_id 미지정 시 요청 내용으로 안정적인 idempotency 키를 만든다 — 같은 요청 재시도가
    같은 시험을 중복 저장하지 않도록 한다(요구사항 5)."""
    stems = [str(q.get('stem', '')) for q in questions]
    seed = '|'.join([school, grade, source_material, source_page, title, *stems])
    return hashlib.sha256(seed.encode('utf-8')).hexdigest()[:16]


def _upsert_source_material(conn, school: str, grade: str, title: str, publisher: str = '',
                            lesson: str = '', material_type: str = '',
                            school_print_material_id: Optional[int] = None) -> int:
    """(school, grade, title) unique 기준 upsert. 이미 있으면 그 id를 재사용하고 빈 필드만 채운다 —
    같은 '학습자료'가 시험마다 중복 생성되지 않게 한다(요구사항 7: 관계 유지)."""
    school, grade, title = school.strip(), grade.strip(), title.strip()
    if not school or not grade or not title:
        raise HTTPException(status_code=422, detail='source material은 school/grade/title이 모두 필요합니다.')
    if school_print_material_id is None:
        needle = f"%{title.split(':')[0].split('(')[0].strip()[:20]}%"
        linked = conn.execute(
            "SELECT id FROM exam_material_sets WHERE core_type='school_print' AND school=? AND grade=? "
            "AND (title LIKE ? OR ? LIKE '%' || title || '%') ORDER BY id DESC LIMIT 1",
            (school, grade, needle, title),
        ).fetchone()
        if linked:
            school_print_material_id = linked['id']
    row = conn.execute(
        'SELECT id, publisher, lesson, material_type, school_print_material_id FROM source_materials '
        'WHERE school=? AND grade=? AND title=?',
        (school, grade, title),
    ).fetchone()
    if row:
        updates, params = [], []
        for field, value in (('publisher', publisher), ('lesson', lesson), ('material_type', material_type)):
            if value and not row[field]:
                updates.append(f'{field}=?')
                params.append(value)
        if school_print_material_id is not None and row['school_print_material_id'] is None:
            updates.append('school_print_material_id=?')
            params.append(school_print_material_id)
        if updates:
            params.append(row['id'])
            conn.execute(f"UPDATE source_materials SET {', '.join(updates)} WHERE id=?", params)
        return row['id']
    cursor = conn.execute(
        'INSERT INTO source_materials (school, grade, title, publisher, lesson, material_type, '
        'school_print_material_id, created_at) VALUES (?,?,?,?,?,?,?,?)',
        (school, grade, title, publisher.strip(), lesson.strip(), material_type.strip(),
         school_print_material_id, datetime.now().isoformat()),
    )
    return cursor.lastrowid


def _generated_test_row(row) -> dict:
    item = dict(row)
    for jkey in ('target_vocabulary', 'target_grammar'):
        try:
            item[jkey] = json.loads(item.get(jkey) or '[]')
        except (TypeError, json.JSONDecodeError):
            item[jkey] = []
    item.pop('raw_request_json', None)
    return item


def create_generated_test(payload: GeneratedTestCreate) -> dict:
    """생성 문항을 시험 단위로 저장하는 단일 write 경로. 저장이 성공했을 때만 saved=True를 준다."""
    school, grade = payload.school.strip(), payload.grade.strip()
    source_material = payload.source_material.strip()
    test_type = payload.test_type.strip()
    if not school or not grade or not source_material:
        raise HTTPException(status_code=422, detail='school/grade/source_material은 비어 있을 수 없습니다.')
    if not payload.questions:
        raise HTTPException(status_code=422, detail='questions는 최소 1개 이상이어야 합니다.')
    title = payload.title.strip() or f'{source_material} {test_type or "Test"}'.strip()
    # 1) 어떤 것도 쓰기 전에 문항 계약을 먼저 검증한다 (실패 시 아무 것도 저장하지 않음).
    try:
        validated = qc.validate_transform({'examTitle': title, 'questions': payload.questions})
    except ValidationError as exc:
        raise HTTPException(status_code=422, detail='문항 검증 실패(저장하지 않음): ' + qc.validation_fields(exc)) from exc
    external_id = payload.external_id.strip() or _derive_external_id(
        school, grade, source_material, payload.source_page.strip(), title, payload.questions)
    try:
        with get_db() as conn:
            existing = conn.execute(
                'SELECT t.*, sm.title source_material_title FROM generated_tests t '
                'JOIN source_materials sm ON sm.id = t.source_material_id WHERE t.external_id=?',
                (external_id,),
            ).fetchone()
            if existing:
                return {
                    'saved': True, 'duplicate': True, 'test_id': existing['id'],
                    'question_count': existing['question_count'],
                    'source_material': existing['source_material_title'], 'source_page': existing['source_page'],
                    'external_id': external_id,
                }
            source_material_id = _upsert_source_material(
                conn, school, grade, source_material, payload.publisher, payload.lesson, payload.material_type)
            twin = _match_digital_twin(conn, school, grade, payload.exam or '전체')
            twin_id = twin.get('id') if twin else None
            twin_label = payload.digital_twin.strip() or (pl._twin_name(twin) if twin else '')
            now = datetime.now().isoformat()
            group_id = f'gtest_{uuid.uuid4().hex}'
            cursor = conn.execute(
                '''INSERT INTO generated_tests
                (external_id, source_material_id, source_page, school, grade, year, semester, exam,
                 test_type, title, twin_id, twin_label, twin_pattern, difficulty, target_vocabulary,
                 target_grammar, target_passage, generation_source, question_count, raw_request_json, created_at)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',
                (external_id, source_material_id, payload.source_page.strip(), school, grade, payload.year,
                 payload.semester.strip(), payload.exam or '전체', test_type, title, twin_id, twin_label,
                 payload.twin_pattern.strip(), payload.difficulty.strip(),
                 json.dumps(payload.target_vocabulary, ensure_ascii=False),
                 json.dumps(payload.target_grammar, ensure_ascii=False), payload.target_passage.strip(),
                 payload.generation_source or 'mcp', 0,
                 json.dumps(payload.model_dump(exclude={'questions'}), ensure_ascii=False), now),
            )
            test_id = cursor.lastrowid
            twin_like = {'id': twin_id, 'school': school, 'grade': grade, 'exam': payload.exam or '전체'}
            saved = _save_generated(conn, twin_like, payload.generation_source or 'mcp',
                                    validated['questions'], title, group_id, test_id=test_id)
            conn.execute('UPDATE generated_tests SET question_count=? WHERE id=?', (len(saved), test_id))
    except HTTPException:
        raise
    except Exception as exc:  # 저장 도중 예기치 못한 실패 — 트랜잭션은 이미 롤백됨. 생성 성공과 구분한다.
        raise HTTPException(status_code=500, detail=f'생성 시험을 저장하지 못했습니다: {exc}') from exc
    return {
        'saved': True, 'duplicate': False, 'test_id': test_id, 'question_count': len(saved),
        'source_material': source_material, 'source_page': payload.source_page.strip(),
        'external_id': external_id, 'questions': saved,
    }


@app.post('/api/generated-tests')
def create_generated_test_endpoint(payload: GeneratedTestCreate):
    return create_generated_test(payload)


@app.get('/api/generated-tests/{test_id}')
def get_generated_test(test_id: int):
    with get_db() as conn:
        row = conn.execute(
            'SELECT t.*, sm.title source_material_title, sm.publisher source_material_publisher, '
            'sm.lesson source_material_lesson, sm.school_print_material_id '
            'FROM generated_tests t JOIN source_materials sm ON sm.id = t.source_material_id WHERE t.id=?',
            (test_id,),
        ).fetchone()
    if not row:
        raise HTTPException(status_code=404, detail='생성 시험을 찾지 못했습니다.')
    questions = sorted(list_questions(test_id=test_id, limit=1000), key=lambda q: q['id'])
    return {'test': _generated_test_row(row), 'questions': questions}


@app.get('/api/source-materials')
def list_source_materials(school: Optional[str] = None, grade: Optional[str] = None):
    query = ('SELECT sm.*, (SELECT COUNT(*) FROM generated_tests t WHERE t.source_material_id = sm.id) '
             'test_count FROM source_materials sm WHERE 1=1')
    values: list = []
    for field, value in (('school', school), ('grade', grade)):
        if value:
            query += f' AND sm.{field} = ?'
            values.append(value)
    query += ' ORDER BY sm.id DESC'
    with get_db() as conn:
        rows = conn.execute(query, values).fetchall()
    return [dict(row) for row in rows]


@app.get('/api/source-materials/{material_id}/generated-tests')
def list_generated_tests_for_material(material_id: int):
    with get_db() as conn:
        rows = conn.execute(
            'SELECT * FROM generated_tests WHERE source_material_id=? ORDER BY id DESC', (material_id,),
        ).fetchall()
    return [_generated_test_row(row) for row in rows]


# ==================================================================================
# FINAL/JIKBO Workspace Context — 향후 MCP/외부 LLM이 한 번의 요청으로 읽을 수 있도록
# 화면에 흩어진 자료를 DB 원본 그대로 구조화된 JSON으로 조립한다(화면 HTML을 긁지 않는다).
# 기존 저장/조회 API의 SQL 패턴을 그대로 재사용하고 별도 테이블에 중복 저장하지 않는다.
# read-only이며 이 섹션의 함수들은 아무 것도 쓰지 않는다.
# ==================================================================================

def _match_digital_twin(conn, school: str, grade: str, exam: str) -> Optional[dict]:
    """NaesinWikiV1.tsx FINAL Workspace가 쓰는 것과 같은 우선순위: school+grade가 정확히 같고,
    exam이 정확히 같은 Twin을 우선하되 없으면 '전체' Twin으로 대체한다."""
    row = conn.execute(
        "SELECT * FROM digital_twin_profiles WHERE is_active=1 AND school=? AND grade=? AND (exam=? OR exam='전체') "
        "ORDER BY CASE WHEN exam=? THEN 0 ELSE 1 END, updated_at DESC LIMIT 1",
        (school, grade, exam, exam),
    ).fetchone()
    if not row:
        return None
    twin = digital_twin_response(row)
    return {**twin, 'type': 'digital_twin', 'source': 'digital_twin_profiles', 'title': twin.get('description') or twin.get('dna') or 'Digital Twin'}


def _material_context_items(conn, core_type: str, school: str, grade: str, year: int, semester: str, exam: str) -> list:
    """exam_material_sets를 list_exam_materials와 같은 Context 조건으로 조회하되, 목록 화면과
    달리 AI가 실제로 읽을 수 있도록 html_text/json_text 원문까지 그대로 포함한다."""
    rows = conn.execute(
        '''SELECT m.id, m.core_type, m.school, m.grade, m.year, m.semester, m.exam, m.student_id,
            s.name student_name, m.title, m.original_name, m.html_text, m.json_text, m.created_at
            FROM exam_material_sets m LEFT JOIN students s ON s.id = m.student_id
            WHERE m.core_type=? AND m.school=? AND m.grade=? AND m.year=? AND m.semester=? AND m.exam=?
            ORDER BY m.id DESC''',
        (core_type, school, grade, year, semester, exam),
    ).fetchall()
    default_title = '학교 프린트' if core_type == 'school_print' else '학생 오답'
    items = []
    for row in rows:
        try:
            parsed_json = json.loads(row['json_text']) if row['json_text'] else {}
        except (TypeError, json.JSONDecodeError):
            parsed_json = {}
        items.append({
            'id': row['id'], 'type': core_type, 'title': row['title'] or row['original_name'] or default_title,
            'created_at': row['created_at'], 'source': 'exam_material_sets',
            'student_id': row['student_id'], 'student_name': row['student_name'],
            'has_original': bool(row['original_name']), 'original_name': row['original_name'],
            'has_html': bool(row['html_text']), 'html_text': row['html_text'] or None,
            'json_text': parsed_json,
        })
    return items


def _final_deliverable_context_items(conn, kind: str, school: str, grade: str, year: int, semester: str, exam: str) -> list:
    rows = conn.execute(
        '''SELECT id, kind, title, original_name, html_text, json_text, created_at
           FROM final_deliverables WHERE kind=? AND school=? AND grade=? AND year=? AND semester=? AND exam=?
           ORDER BY id DESC''',
        (kind, school, grade, year, semester, exam),
    ).fetchall()
    items = []
    for row in rows:
        try:
            parsed_json = json.loads(row['json_text']) if row['json_text'] else {}
        except (TypeError, json.JSONDecodeError):
            parsed_json = {}
        items.append({
            'id': row['id'], 'type': 'final_deliverable', 'title': row['title'] or 'FINAL 완성본',
            'created_at': row['created_at'], 'source': 'final_deliverables', 'kind': row['kind'],
            'has_original': bool(row['original_name']), 'original_name': row['original_name'],
            'has_html': bool(row['html_text']), 'html_text': row['html_text'] or None,
            'json_text': parsed_json,
        })
    return items


def _question_context_items(school: str, grade: str, exam: str) -> list:
    rows = list_questions(school=school, grade=grade, exam=exam, limit=200)
    return [{
        **row, 'type': 'generated_question', 'source': 'generated_questions',
        'title': f"{row.get('question_type') or '문항'} · {(row.get('stem') or '')[:60]}",
    } for row in rows]


def _generated_test_context_items(conn, school: str, grade: str, exam: str) -> list:
    """Generated Question Store에 저장된 생성 시험을 Context 항목으로 담는다(추적성). exam이 '전체'가
    아니면 그 시험과 '전체' 저장분을 함께 본다 — 기존 Twin 매칭 우선순위와 같은 관대함을 따른다."""
    rows = conn.execute(
        'SELECT t.*, sm.title source_material_title FROM generated_tests t '
        'JOIN source_materials sm ON sm.id = t.source_material_id '
        "WHERE t.school=? AND t.grade=? AND (t.exam=? OR t.exam='전체' OR ?='전체') ORDER BY t.id DESC",
        (school, grade, exam, exam),
    ).fetchall()
    items = []
    for row in rows:
        data = _generated_test_row(row)
        items.append({
            'id': data['id'], 'type': 'generated_test', 'source': 'generated_tests',
            'title': data['title'] or data['source_material_title'],
            'source_material': data['source_material_title'], 'source_page': data['source_page'],
            'test_type': data['test_type'], 'question_count': data['question_count'],
            'twin_label': data['twin_label'], 'generation_source': data['generation_source'],
            'created_at': data['created_at'],
        })
    return items


def _insight_context_items(school: str, grade: str, year: int, semester: str, exam: str, after: str = '') -> list:
    """수업 인사이트를 Context 항목 형태(id/type/title/created_at/source)로 변환한다.
    after가 주어지면 그 시각 이후 기록만 남긴다(직보의 post_final_insights용).
    학교 선생님 힌트를 박원장 인사이트보다 먼저 보여주는 기존 화면 우선순위(sortedInsights)를 그대로 따른다."""
    rows = list_educator_insights(200, school, grade, year, semester, exam)
    items = [{
        'id': row['id'], 'type': 'educator_insight', 'title': None, 'content': row['content'],
        'created_at': row['created_at'], 'source': row['source'],
    } for row in rows if not after or row['created_at'] > after]
    items.sort(key=lambda item: 0 if item['source'] == 'school_teacher' else 1)
    return items


@app.get('/api/workspaces/final/context')
def get_final_workspace_context(school: str, grade: str, year: int, semester: str, exam: str):
    with get_db() as conn:
        digital_twin = _match_digital_twin(conn, school, grade, exam)
        school_prints = _material_context_items(conn, 'school_print', school, grade, year, semester, exam)
        student_errors = _material_context_items(conn, 'student_wrong_note', school, grade, year, semester, exam)
        final_deliverables = _final_deliverable_context_items(conn, 'final', school, grade, year, semester, exam)
        generated_tests = _generated_test_context_items(conn, school, grade, exam)
    return {
        'exam_context': {'school': school, 'grade': grade, 'year': year, 'semester': semester, 'exam': exam},
        'digital_twin': digital_twin,
        'school_prints': school_prints,
        'student_errors': student_errors,
        'class_insights': _insight_context_items(school, grade, year, semester, exam),
        'generated_questions': _question_context_items(school, grade, exam),
        'generated_tests': generated_tests,
        'final_deliverables': final_deliverables,
    }


@app.get('/api/workspaces/jikbo/context')
def get_jikbo_workspace_context(school: str, grade: str, year: int, semester: str, exam: str):
    """직보는 FINAL과 별도 원천 저장소가 아니다 — FINAL 완성본 → FINAL 이후 최신 수업 인사이트 →
    학생 반복 오답 순으로만 압축한다. 학교 프린트/제작 문제/Digital Twin 전체는 FINAL 제작 과정에서
    이미 반영된 자료라 다시 담지 않는다."""
    with get_db() as conn:
        final_items = _final_deliverable_context_items(conn, 'final', school, grade, year, semester, exam)
        latest_final = final_items[0] if final_items else None
        student_weaknesses = _material_context_items(conn, 'student_wrong_note', school, grade, year, semester, exam)
        jikbo_deliverables = _final_deliverable_context_items(conn, 'jikbo', school, grade, year, semester, exam)
    ready = latest_final is not None
    post_final_insights = _insight_context_items(school, grade, year, semester, exam, after=latest_final['created_at']) if ready else []
    return {
        'exam_context': {'school': school, 'grade': grade, 'year': year, 'semester': semester, 'exam': exam},
        'ready': ready,
        'latest_final': latest_final,
        'post_final_insights': post_final_insights,
        'student_weaknesses': student_weaknesses,
        'jikbo_deliverables': jikbo_deliverables,
    }


# ==================================================================================
# 내신 실전 제작 — 강사 문제 DB에서 선택한 문항으로 구성한 시험지를 실제로 저장한다.
# ==================================================================================

class ExamPaperCreate(BaseModel):
    title: str
    academy_name: str = ''
    teacher_name: str = ''
    question_ids: list[int] = []


@app.post('/api/exam-papers')
def create_exam_paper(payload: ExamPaperCreate):
    if not payload.question_ids:
        raise HTTPException(status_code=400, detail='저장할 문항을 강사 문제 DB에서 먼저 선택해 주세요.')
    now = datetime.now().isoformat()
    with get_db() as conn:
        rows = conn.execute(
            f"SELECT id FROM generated_questions WHERE id IN ({','.join('?' * len(payload.question_ids))})",
            payload.question_ids,
        ).fetchall()
        valid_ids = [row['id'] for row in rows]
        if not valid_ids:
            raise HTTPException(status_code=400, detail='선택한 문항을 강사 문제 DB에서 찾지 못했습니다.')
        cursor = conn.execute(
            'INSERT INTO exam_papers (title, academy_name, teacher_name, question_ids_json, created_at) VALUES (?,?,?,?,?)',
            (payload.title.strip() or '내신 실전 시험지', payload.academy_name.strip(), payload.teacher_name.strip(), json.dumps(valid_ids), now),
        )
        return {'id': cursor.lastrowid, 'title': payload.title, 'academy_name': payload.academy_name, 'teacher_name': payload.teacher_name, 'question_ids': valid_ids, 'created_at': now}


@app.get('/api/exam-papers')
def list_exam_papers():
    with get_db() as conn:
        rows = conn.execute('SELECT * FROM exam_papers ORDER BY id DESC').fetchall()
    result = []
    for row in rows:
        item = dict(row)
        item['question_ids'] = json.loads(item.pop('question_ids_json'))
        result.append(item)
    return result


# ==================================================================================
# 내신 실전 제작 — 현재 작업 중인 시험지의 선택 문항/순서를 DB에 지속시키는 draft.
# React state에만 두지 않고 새로고침·페이지 이동 후에도 그대로 복원되도록 한다.
# ==================================================================================

class ExamDraftUpdate(BaseModel):
    title: str = ''
    academy_name: str = ''
    teacher_name: str = ''
    question_ids: list[int] = []


@app.get('/api/exam-draft')
def get_exam_draft():
    with get_db() as conn:
        row = conn.execute('SELECT * FROM exam_draft WHERE id = 1').fetchone()
    if not row:
        return {'title': '', 'academy_name': '', 'teacher_name': '', 'question_ids': []}
    return {
        'title': row['title'], 'academy_name': row['academy_name'], 'teacher_name': row['teacher_name'],
        'question_ids': json.loads(row['question_ids_json']),
    }


@app.put('/api/exam-draft')
def save_exam_draft(payload: ExamDraftUpdate):
    now = datetime.now().isoformat()
    with get_db() as conn:
        conn.execute(
            '''INSERT INTO exam_draft (id, title, academy_name, teacher_name, question_ids_json, updated_at)
               VALUES (1, ?, ?, ?, ?, ?)
               ON CONFLICT(id) DO UPDATE SET
                 title = excluded.title, academy_name = excluded.academy_name,
                 teacher_name = excluded.teacher_name, question_ids_json = excluded.question_ids_json,
                 updated_at = excluded.updated_at''',
            (payload.title, payload.academy_name, payload.teacher_name, json.dumps(payload.question_ids), now),
        )
    return {'ok': True}


# ==================================================================================
# 내신 실전 제작 — 오피스 2007(.docx) 문제지/정답지 내보내기
# ==================================================================================

_UNDERLINE_RE = re.compile(r'\[\[u\]\](.*?)\[\[/u\]\]', re.DOTALL)
_OPTION_MARKS = ['①', '②', '③', '④', '⑤']


def _render_markup_segments(text: str):
    """[[u]]...[[/u]] 형태로 저장된 밑줄 마크업을 (텍스트, 밑줄여부) 세그먼트 목록으로 변환한다."""
    text = text or ''
    segments: list[tuple[str, bool]] = []
    pos = 0
    for m in _UNDERLINE_RE.finditer(text):
        if m.start() > pos:
            segments.append((text[pos:m.start()], False))
        segments.append((m.group(1), True))
        pos = m.end()
    if pos < len(text):
        segments.append((text[pos:], False))
    if not segments:
        segments.append(('', False))
    return [(seg.replace('[[u]]', '').replace('[[/u]]', ''), underline) for seg, underline in segments]


def _set_korean_font(run, name: str = '맑은 고딕'):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), name)


def _add_markup_run(paragraph, text: str, bold: bool = False):
    for seg_text, underline in _render_markup_segments(text):
        if not seg_text:
            continue
        run = paragraph.add_run(seg_text)
        run.bold = bold
        run.underline = underline
        _set_korean_font(run)
    return paragraph


def _set_columns(section, num: int = 1):
    """섹션 본문을 num단으로 설정한다 (1=단일 칼럼, 2=신문형 2단)."""
    sect_pr = section._sectPr
    cols = sect_pr.find(qn('w:cols'))
    if cols is None:
        cols = OxmlElement('w:cols')
        sect_pr.append(cols)
    cols.set(qn('w:num'), str(num))
    if num > 1:
        cols.set(qn('w:space'), '425')  # 칼럼 사이 간격(약 0.3인치)
    elif qn('w:space') in cols.attrib:
        del cols.attrib[qn('w:space')]


def _configure_a4_section(section, columns: int = 1):
    """문제11(6차): A4 용지 크기와 여백을 명시적으로 지정한다 (python-docx 기본값은 Letter이므로
    A4로 맞춰야 한다). 문제지 본문은 2단, 표지/정답 섹션은 1단으로 재사용한다."""
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(15)
    section.bottom_margin = Mm(15)
    section.left_margin = Mm(14)
    section.right_margin = Mm(14)
    _set_columns(section, columns)


def _new_docx_with_korean_style() -> Document:
    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = '맑은 고딕'
    normal.font.size = Pt(10)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), '맑은 고딕')
    _configure_a4_section(doc.sections[0], columns=1)
    return doc


def _start_section(doc: Document, columns: int, new_page: bool):
    """문제11(6차): 표지(1단) → 본문(2단) → 정답(1단, 새 페이지)처럼 구간별로 칼럼 수가
    다른 섹션을 이어붙인다. new_page=True면 항상 새 페이지에서 시작한다(명시적 page_break 불필요)."""
    section = doc.add_section(WD_SECTION.NEW_PAGE if new_page else WD_SECTION.CONTINUOUS)
    _configure_a4_section(section, columns=columns)
    return section


def _add_doc_heading(doc: Document, title: str, subtitle: str, label: str):
    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading_run = heading.add_run(title or '문제지')
    heading_run.bold = True
    heading_run.font.size = Pt(18)
    _set_korean_font(heading_run)

    if subtitle:
        sub = doc.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub_run = sub.add_run(subtitle)
        sub_run.font.size = Pt(11)
        _set_korean_font(sub_run)

    label_p = doc.add_paragraph()
    label_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label_run = label_p.add_run(label)
    label_run.bold = True
    _set_korean_font(label_run)


def _build_question_section(doc: Document, questions: list[dict]):
    """문제지 섹션. 객관식 선택지는 ①~⑤ 각각 별도 문단(줄)으로 추가한다
    (한 문단에 이어 붙이지 않는다)."""
    for idx, q in enumerate(questions, start=1):
        head = doc.add_paragraph()
        num_run = head.add_run(f"{idx}. [{q.get('score', 0)}점]  ")
        num_run.bold = True
        _set_korean_font(num_run)
        _add_markup_run(head, q.get('stem', ''))

        if q.get('source_text'):
            src = doc.add_paragraph()
            _add_markup_run(src, q.get('source_text', ''))

        if q.get('q_kind') == 'sa':
            conditions = q.get('conditions') or []
            if conditions:
                cond_p = doc.add_paragraph()
                cond_run = cond_p.add_run('조건: ' + ' / '.join(conditions))
                _set_korean_font(cond_run)
            blank = doc.add_paragraph()
            blank_run = blank.add_run('답: ______________________________________________________')
            _set_korean_font(blank_run)
        else:
            options = q.get('options') or []
            for i, opt in enumerate(options):
                opt_p = doc.add_paragraph()
                mark = _OPTION_MARKS[i] if i < len(_OPTION_MARKS) else str(i + 1)
                mark_run = opt_p.add_run(f"{mark} ")
                _set_korean_font(mark_run)
                _add_markup_run(opt_p, opt)
        doc.add_paragraph('')


def _build_answer_section(doc: Document, questions: list[dict]):
    """정답 섹션. 요청하신 형태(예: "1. ③")대로 문항 번호당 한 줄씩 정답만 간단히 표시한다."""
    for idx, q in enumerate(questions, start=1):
        line = doc.add_paragraph()
        num_run = line.add_run(f"{idx}. ")
        num_run.bold = True
        _set_korean_font(num_run)
        if q.get('q_kind') == 'sa':
            ans_run = line.add_run(f"모범답안: {q.get('sample_answer', '') or '-'}")
        else:
            answer_idx = sorted(set(q.get('answer') or []))
            labels = ', '.join(_OPTION_MARKS[a - 1] if 1 <= a <= len(_OPTION_MARKS) else str(a) for a in answer_idx)
            ans_run = line.add_run(labels or '-')
        ans_run.bold = True
        _set_korean_font(ans_run)
        if q.get('q_kind') == 'sa' and q.get('grading_rule'):
            rule_p = doc.add_paragraph()
            rule_run = rule_p.add_run(f"    채점 기준: {q.get('grading_rule', '')}")
            rule_run.font.size = Pt(9.5)
            _set_korean_font(rule_run)


def _build_combined_exam_document(meta: dict, questions: list[dict]) -> Document:
    """ZIP으로 문제지·정답지 두 파일을 나눠 만들던 방식을 폐지하고, 문제지(A4, 2단) 마지막에
    새 섹션(1단)으로 정답을 이어붙인 단일 .docx(오피스 2007 호환) 하나만 생성한다."""
    doc = _new_docx_with_korean_style()  # 표지: 1단
    subtitle = ' · '.join(filter(None, [meta.get('academy_name', ''), meta.get('teacher_name', '')]))
    _add_doc_heading(doc, meta.get('title') or '내신 실전 시험지', subtitle, '문제지')

    _start_section(doc, columns=2, new_page=False)  # 본문: 2단 (표지 바로 아래에서 이어짐)
    _build_question_section(doc, questions)

    _start_section(doc, columns=1, new_page=True)  # 정답: 새 페이지, 1단
    answer_label = doc.add_paragraph()
    answer_label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    answer_label_run = answer_label.add_run('정답')
    answer_label_run.bold = True
    answer_label_run.font.size = Pt(15)
    _set_korean_font(answer_label_run)
    doc.add_paragraph('')
    _build_answer_section(doc, questions)

    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run(subtitle)
    footer_run.font.size = Pt(9)
    _set_korean_font(footer_run)
    return doc


def _build_wrong_note_combined_document(title: str, subtitle: str, notes: list[dict], document_label: str = '오답 복습 문제') -> Document:
    """문제9~11(6차): 학생이 선택한 오답 기록 여러 개를 하나의 통합 오답 복습 문제지로 만든다.
    내신 실전 제작(_build_combined_exam_document)과 동일한 A4/2단/맑은 고딕 10pt/페이지 나누기
    helper(_new_docx_with_korean_style, _start_section, _add_doc_heading, _set_korean_font,
    _add_markup_run)를 그대로 재사용한다 — 새로운 스타일 체계를 만들지 않는다."""
    doc = _new_docx_with_korean_style()
    _add_doc_heading(doc, title, subtitle, document_label)

    _start_section(doc, columns=2, new_page=False)
    for idx, note in enumerate(notes, start=1):
        head = doc.add_paragraph()
        num_run = head.add_run(f"{idx}. ")
        num_run.bold = True
        _set_korean_font(num_run)
        if note.get('question_type'):
            type_run = head.add_run(f"[{note['question_type']}]  ")
            type_run.bold = True
            _set_korean_font(type_run)
        if note.get('title'):
            _add_markup_run(head, note['title'])
        body = doc.add_paragraph()
        _add_markup_run(body, note.get('wrong_text', ''))
        doc.add_paragraph('')

    _start_section(doc, columns=1, new_page=True)
    answer_label = doc.add_paragraph()
    answer_label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    answer_label_run = answer_label.add_run('정답')
    answer_label_run.bold = True
    answer_label_run.font.size = Pt(15)
    _set_korean_font(answer_label_run)
    doc.add_paragraph('')
    for idx, note in enumerate(notes, start=1):
        line = doc.add_paragraph()
        num_run = line.add_run(f"{idx}. ")
        num_run.bold = True
        _set_korean_font(num_run)
        ans_run = line.add_run(note.get('answer_text') or '-')
        ans_run.bold = True
        _set_korean_font(ans_run)

    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run(subtitle)
    footer_run.font.size = Pt(9)
    _set_korean_font(footer_run)
    return doc


class ExamExportRequest(BaseModel):
    title: str = ''
    academy_name: str = ''
    teacher_name: str = ''
    question_ids: list[int] = []


@app.post('/api/exam-papers/export-docx')
def export_exam_docx(payload: ExamExportRequest):
    """선택된 전체 문항을 기준으로 문제지 + 정답을 하나의 오피스 2007(.docx) 파일로 내보낸다.
    각 문항은 강사 문제 DB(generated_questions)에 저장된 전체 데이터(지문/발문/선택지/조건/정답 등)를
    그대로 불러와 렌더링한다 — 실전 제작 화면에서 별도로 재구성하지 않는다."""
    if not payload.question_ids:
        raise HTTPException(status_code=400, detail='내보낼 문항을 먼저 선택해 주세요.')
    with get_db() as conn:
        rows = conn.execute(
            f"SELECT * FROM generated_questions WHERE id IN ({','.join('?' * len(payload.question_ids))})",
            payload.question_ids,
        ).fetchall()
    by_id = {row['id']: dict(row) for row in rows}
    ordered = [by_id[qid] for qid in payload.question_ids if qid in by_id]
    if not ordered:
        raise HTTPException(status_code=400, detail='선택한 문항을 강사 문제 DB에서 찾지 못했습니다.')

    questions = []
    for row in ordered:
        questions.append({
            'score': row.get('score', 0),
            'stem': row.get('stem', ''),
            'source_text': row.get('source_text', ''),
            'q_kind': row.get('q_kind', 'mc'),
            'options': json.loads(row.get('options_json') or '[]'),
            'answer': json.loads(row.get('answer_json') or '[]'),
            'conditions': json.loads(row.get('conditions_json') or '[]'),
            'sample_answer': row.get('sample_answer', ''),
            'grading_rule': row.get('grading_rule', ''),
        })
    meta = {
        'title': payload.title.strip() or '내신 실전 시험지',
        'academy_name': payload.academy_name.strip(),
        'teacher_name': payload.teacher_name.strip(),
    }
    doc = _build_combined_exam_document(meta, questions)

    export_dir = WORKSPACE / 'data' / 'exports'
    export_dir.mkdir(parents=True, exist_ok=True)
    safe_title = re.sub(r'[\\/:*?"<>|]', '_', meta['title']) or '내신_실전_시험지'
    doc_path = export_dir / 'exam-combined.docx'
    doc.save(doc_path)
    return FileResponse(
        doc_path,
        media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        filename=f'{safe_title}.docx',
    )


# ==================================================================================
# 학생 프로필 — 이름/학교/학년/학기/시험(중간·기말)만 저장한다.
# ==================================================================================

class StudentCreate(BaseModel):
    name: str
    school: str = ''
    grade: str = ''
    semester: str = ''
    exam_type: str = ''


@app.post('/api/students')
def create_student(payload: StudentCreate):
    if not payload.name.strip():
        raise HTTPException(status_code=400, detail='학생 이름을 입력해 주세요.')
    now = datetime.now().strftime('%Y.%m.%d %H:%M')
    with get_db() as conn:
        cursor = conn.execute(
            'INSERT INTO students (name, school, grade, semester, exam_type, created_at) VALUES (?,?,?,?,?,?)',
            (payload.name.strip(), payload.school.strip(), payload.grade.strip(), payload.semester.strip(), payload.exam_type.strip(), now),
        )
        row = conn.execute('SELECT * FROM students WHERE id = ?', (cursor.lastrowid,)).fetchone()
        return dict(row)


@app.get('/api/students')
def list_students():
    with get_db() as conn:
        rows = conn.execute('SELECT * FROM students ORDER BY id DESC').fetchall()
        return [dict(row) for row in rows]


@app.delete('/api/students/{student_id}')
def delete_student(student_id: int):
    with get_db() as conn:
        conn.execute('DELETE FROM student_wrong_notes WHERE student_id = ?', (student_id,))
        conn.execute('DELETE FROM students WHERE id = ?', (student_id,))
    return {'ok': True}


# ==================================================================================
# 오답 사진/PDF → 텍스트 추출 2차 정리 (Gemini).
# 1차 브라우저 추출(PDF 텍스트 레이어 + 필요 시 OCR)로 얻은 원문을, 요약·재작성 없이
# 원본 시험지 구조(문제 번호/지문/발문/선택지/밑줄/채점 표시 등) 그대로 정돈만 한다.
# 사용자에게 별도 API Key 입력 UI를 새로 만들지 않고, 이미 앱에 설정되어 있는 Gemini 키를 그대로 사용한다.
# ==================================================================================

class OcrCleanRequest(BaseModel):
    api_key: str = ''
    raw_text: str


class PassageRequest(BaseModel):
    api_key: str = ''
    source_text: str


# 2027학년도 6월 모의평가 18~45번을 기준으로 한 번호별 표준 유형이다.
# 번호는 1차 힌트이고, 저장 시 발문/문항 구조를 함께 검사해 불일치를 별도 기록한다.
CSAT_TYPE_RANGES = [
    ((18,), 'purpose', '목적', ('목적으로', 'purpose')),
    ((19,), 'emotion', '심경', ('심경', 'mood', 'feeling')),
    ((20,), 'claim', '주장', ('주장', 'claim')),
    ((21,), 'implication', '함축 의미', ('의미하는', '함축', 'imply')),
    ((22,), 'main_idea', '요지', ('요지', 'main idea')),
    ((23,), 'topic', '주제', ('주제로', 'topic')),
    ((24,), 'title', '제목', ('제목', 'title')),
    ((25,), 'chart', '도표', ('도표', 'graph', 'chart')),
    ((26,), 'content_match', '내용 일치·불일치', ('내용과 일치', '내용과 일치하지')),
    ((27, 28), 'notice', '안내문', ('안내문', 'event', 'competition')),
    ((29,), 'grammar', '어법', ('어법', 'grammar')),
    ((30,), 'vocabulary', '어휘', ('낱말', '어휘', '문맥상 적절하지', 'word')),
    ((31, 32, 33, 34), 'blank', '빈칸', ('빈칸', 'blank')),
    ((35,), 'irrelevant', '무관 문장', ('관계 없는', '무관')),
    ((36, 37), 'order', '글의 순서', ('순서', 'order')),
    ((38, 39), 'insertion', '문장 삽입', ('들어가기에', '삽입')),
    ((40,), 'summary', '요약', ('요약', 'summary')),
    ((41, 42), 'long_complex', '장문 복합', ('41', '42', '윗글')),
    ((43, 44, 45), 'long_sequence_reference_content', '장문 순서·지칭·내용', ('43', '44', '45', '윗글')),
]
HALF_FEATURED_NUMBERS = {21, 24, 29, 30, *range(31, 35), *range(36, 43)}


def csat_type_for_number(question_no: int) -> dict:
    for numbers, type_id, name, signals in CSAT_TYPE_RANGES:
        if question_no in numbers:
            return {'id': type_id, 'name': name, 'signals': signals}
    return {'id': 'other', 'name': '기타', 'signals': ()}


def classify_csat_question(question_no: int, prompt: str, choices: list[str]) -> dict:
    expected = csat_type_for_number(question_no)
    haystack = f"{prompt} {' '.join(choices)}".lower()
    matched = any(signal.lower() in haystack for signal in expected['signals'])
    # 복합문항은 그룹 발문이 한 번만 나올 수 있어 보기 5개 또는 번호 표기도 구조 신호로 인정한다.
    if question_no >= 41 and len(choices) == 5:
        matched = True
    return {
        'type_id': expected['id'], 'type_name': expected['name'],
        'type_mismatch': not matched,
        'mismatch_reason': '' if matched else f"{question_no}번 표준 유형({expected['name']})과 발문/구조 신호가 일치하지 않습니다.",
    }


def classify_csat_prompt(prompt: str, choices: list[str] | None = None) -> dict:
    """발문만으로 분류하되 같은 유형의 여러 수능 번호 중 하나를 임의 선택하지 않는다."""
    # Choice prose routinely contains words such as ``event`` or ``purpose`` and
    # must not override the actual instruction. Classification is prompt-scoped.
    haystack = prompt.lower()
    matches = []
    for numbers, type_id, name, signals in CSAT_TYPE_RANGES:
        if any(signal.lower() in haystack for signal in signals):
            matches.append({'numbers': list(numbers), 'id': type_id, 'name': name})
    if len(matches) > 1:
        singleton = [item for item in matches if len(item['numbers']) == 1]
        if len(singleton) == 1:
            matches = singleton
    if len(matches) > 1:
        exact = [item for item in matches if item['id'] in ('grammar', 'vocabulary', 'blank', 'order', 'insertion', 'summary')]
        if len(exact) == 1:
            matches = exact
    if len(matches) != 1:
        return {'type_id': 'review', 'type_name': 'REVIEW', 'candidate_numbers': [],
                'classification_status': 'review', 'classification_reason': '발문 규칙으로 수능 번호 유형을 하나로 판별하지 못했습니다.'}
    match = matches[0]
    candidates = match['numbers']
    return {'type_id': match['id'], 'type_name': match['name'], 'candidate_numbers': candidates,
            'classification_status': 'classified' if len(candidates) == 1 else 'review',
            'classification_reason': '' if len(candidates) == 1 else f"{','.join(map(str, candidates))}번이 같은 발문 유형이므로 정확한 번호를 검수해 주세요."}


def _html_text(fragment: str) -> str:
    import html
    value = re.sub(r'(?is)<(?:script|style)\b.*?</(?:script|style)>', ' ', fragment)
    value = re.sub(r'(?i)<br\s*/?>', '\n', value)
    value = re.sub(r'(?i)</(?:p|div|li|h[1-6]|section|article)>', '\n', value)
    return re.sub(r'[ \t]+', ' ', html.unescape(re.sub(r'(?s)<[^>]+>', ' ', value))).strip()


def _first_html_part(fragment: str, class_names: tuple[str, ...], tags: str = r'[a-z][\w:-]*') -> str:
    names = '|'.join(re.escape(name) for name in class_names)
    pattern = rf'(?is)<(?P<tag>{tags})\b[^>]*class=["\'][^"\']*\b(?:{names})\b[^"\']*["\'][^>]*>(?P<body>.*?)</(?P=tag)>'
    found = re.search(pattern, fragment)
    return found.group('body').strip() if found else ''


def _question_html_fragments(html_text: str):
    """Yield balanced top-level question containers without truncating nested divs."""
    tag_pattern = re.compile(r'(?is)<(?P<closing>/)?(?P<tag>article|section|div)\b(?P<attrs>[^>]*)>')
    tokens = list(tag_pattern.finditer(html_text))
    for position, token in enumerate(tokens):
        if token.group('closing'):
            continue
        attrs = token.group('attrs')
        marked = re.search(r'(?i)data-question-no\s*=', attrs) or re.search(
            r'(?i)class\s*=\s*["\'][^"\']*\b(?:question|q)\b', attrs)
        if not marked:
            continue
        tag, depth = token.group('tag').lower(), 1
        for end in tokens[position + 1:]:
            if end.group('tag').lower() != tag:
                continue
            depth += -1 if end.group('closing') else 1
            if depth == 0:
                yield html_text[token.start():end.end()], attrs
                break


def split_csat_html(html_text: str) -> list[dict]:
    """Editable HTML을 보수적으로 분리하고 인식한 원문 fragment를 그대로 보존한다."""
    questions = []
    for index, (fragment, attrs) in enumerate(_question_html_fragments(html_text), start=1):
        number_match = re.search(r'(?i)data-(?:source-)?question-no\s*=\s*["\']?(\d{1,2})', attrs)
        visible = _html_text(fragment)
        if not number_match:
            number_match = re.search(r'^\s*(\d{1,2})\s*[.\)]', visible)
        source_no = int(number_match.group(1)) if number_match else index
        prompt_html = _first_html_part(fragment, ('prompt', 'stem', 'instruction'))
        if not prompt_html:
            heading = re.search(r'(?is)<(?:h[1-6]|p)\b[^>]*>(.*?)</(?:h[1-6]|p)>', fragment)
            prompt_html = heading.group(1) if heading else ''
        prompt = re.sub(r'^\s*\d{1,2}\s*[.\)]\s*', '', _html_text(prompt_html))
        passage_html = _first_html_part(fragment, ('passage', 'source-text', 'reading'))
        choices_html = _first_html_part(fragment, ('choices', 'options'))
        choices = [_html_text(item) for item in re.findall(r'(?is)<li\b[^>]*>(.*?)</li>', choices_html or fragment)]
        answer_html = _first_html_part(fragment, ('answer', 'correct-answer'))
        explanation_html = _first_html_part(fragment, ('explanation', 'solution', 'commentary'))
        classification = classify_csat_prompt(prompt, choices)
        questions.append({
            'source_question_no': source_no,
            'question_no': classification['candidate_numbers'][0] if len(classification['candidate_numbers']) == 1 else 0,
            'csat_type_id': classification['type_id'], 'csat_type_name': classification['type_name'],
            'stem': prompt, 'prompt': prompt, 'passage': _html_text(passage_html), 'choices': choices,
            'answer': _html_text(answer_html), 'explanation': _html_text(explanation_html),
            'question_html': fragment, 'explanation_html': explanation_html, **classification,
        })
    return questions


CSAT_JSON_TYPE_ALIASES = {
    '목적': ('purpose', '목적', [18]), '심경': ('emotion', '심경', [19]),
    '주장': ('claim', '필자의 주장', [20]), '함축 의미': ('implication', '함축 의미', [21]),
    '요지': ('main_idea', '요지', [22]), '주제': ('topic', '주제', [23]), '제목': ('title', '제목', [24]),
    '도표': ('chart', '도표', [25]), '내용 일치': ('content_match', '내용 일치·불일치', [26]),
    '내용 불일치': ('content_match', '내용 일치·불일치', [26]), '안내문': ('notice', '안내문', [27, 28]),
    '어법': ('grammar', '어법', [29]),
    '어휘': ('vocabulary', '어휘', [30]), '빈칸': ('blank', '빈칸 추론', [31, 32, 33, 34]),
    '무관한 문장': ('irrelevant', '무관한 문장', [35]), '순서': ('order', '글의 순서', [36, 37]),
    '문장 삽입': ('insertion', '문장 삽입', [38, 39]), '삽입': ('insertion', '문장 삽입', [38, 39]),
    '요약': ('summary', '요약문 완성', [40]), '장문': ('long_complex', '장문', [41, 42, 43, 44, 45]),
    '공통지문': ('long_complex', '장문', [41, 42, 43, 44, 45]),
}


def _normalise_csat_problem_choices(problem: dict, raw_choices: object) -> tuple[dict, list[str]]:
    """Normalize the production V1 layout where choices trail the passage.

    Only a final, sequential circled-number block is extracted. Digits elsewhere
    in the passage are never considered. The original problem remains available
    as ``source_problem`` in raw_json and the source file is retained verbatim.
    """
    if isinstance(raw_choices, list) and raw_choices:
        return {**problem, 'choices': raw_choices}, [str(choice) for choice in raw_choices]
    passage = str(problem.get('passage') or '')
    markers = '①②③④⑤'
    markup = problem.get('markup') if isinstance(problem.get('markup'), list) else []
    # In insertion and underlined-word questions, the source markup is itself the
    # explicit option model. Preserve it instead of inventing text choices.
    for item in markup:
        if not isinstance(item, dict):
            continue
        values = item.get('labels') if item.get('type') == 'insertion_positions' else \
            item.get('tokens') if item.get('type') == 'underline_targets' else None
        if isinstance(values, list) and len(values) >= 2 and all(str(value).strip() for value in values):
            choices = [str(value) for value in values]
            return {**problem, 'choices': choices}, choices

    occurrences = [(match.start(), markers.index(match.group()) + 1)
                   for match in re.finditer(r'[①②③④⑤]', passage)]
    selected = None
    for position, (offset, marker_no) in enumerate(occurrences):
        tail = occurrences[position:]
        numbers = [number for _, number in tail]
        if marker_no == 1 and numbers == list(range(1, len(numbers) + 1)) and len(numbers) >= 2:
            selected = (offset, tail)
    if selected is None:
        return dict(problem), []
    start_offset, tail = selected
    choices = []
    for position, (offset, _) in enumerate(tail):
        end = tail[position + 1][0] if position + 1 < len(tail) else len(passage)
        choices.append(passage[offset:end].strip())
    normalized = {**problem, 'passage': passage[:start_offset].rstrip(), 'choices': choices}
    return normalized, choices


def _parse_csat_number_spec(value: object) -> list[int]:
    """Parse only an explicit CSAT-number field; never mine numbers from prose."""
    if value is None or value == '' or isinstance(value, bool):
        return []
    if isinstance(value, int):
        numbers = [value]
    elif isinstance(value, list):
        if not value or any(isinstance(item, bool) or not isinstance(item, int) for item in value):
            return []
        numbers = value
    elif isinstance(value, str):
        source = value.strip()
        range_match = re.fullmatch(r'(\d{1,2})\s*[-~～〜–—]\s*(\d{1,2})', source)
        if range_match:
            start, end = map(int, range_match.groups())
            if start > end:
                return []
            numbers = list(range(start, end + 1))
        elif re.fullmatch(r'\d{1,2}(?:\s*[·.,]\s*\d{1,2})+', source):
            numbers = [int(item) for item in re.split(r'\s*[·.,]\s*', source)]
        elif re.fullmatch(r'\d{1,2}', source):
            numbers = [int(source)]
        else:
            return []
    else:
        return []
    return numbers if len(numbers) == len(set(numbers)) and all(18 <= number <= 45 for number in numbers) else []


def _normalise_csat_json_question(raw: dict, inherited: dict, index: int) -> dict:
    if not isinstance(raw, dict):
        raise ValueError('문항은 JSON object여야 합니다.')
    problem = raw.get('problem') or {}
    if not isinstance(problem, dict):
        raise ValueError('problem은 JSON object여야 합니다.')
    source_no = int(raw.get('source_question_no', raw.get('number', index)) or index)
    source_number_raw = next((raw[key] for key in ('csat_question_no', 'question_no', 'question_number')
                              if raw.get(key) not in (None, '')), '')
    declared_numbers = _parse_csat_number_spec(source_number_raw)
    explicit_no = declared_numbers[0] if len(declared_numbers) == 1 else 0
    supplied_type = str(raw.get('type') or raw.get('csat_question_type') or '').strip()
    alias = CSAT_JSON_TYPE_ALIASES.get(supplied_type)
    instruction = str(problem.get('instruction', raw.get('instruction', raw.get('prompt', raw.get('stem', '')))) or '')
    raw_choices = problem.get('choices', raw.get('choices', [])) or []
    if not isinstance(raw_choices, list):
        raise ValueError('problem.choices는 array여야 합니다.')
    normalized_problem, choices = _normalise_csat_problem_choices(problem, raw_choices)
    passage = str(normalized_problem.get('passage', raw.get('passage', '')) or '')
    if explicit_no:
        classified = classify_csat_question(explicit_no, instruction, choices)
        classification_status = 'review' if classified['type_mismatch'] else 'classified'
        candidates = [explicit_no]
        candidate_source = 'explicit'
    elif len(declared_numbers) > 1:
        standards = [csat_type_for_number(number) for number in declared_numbers]
        same_type = len({item['id'] for item in standards}) == 1
        classified = {
            'type_id': standards[0]['id'] if same_type else 'review',
            'type_name': standards[0]['name'] if same_type else 'REVIEW',
            'type_mismatch': not same_type,
            'mismatch_reason': '' if same_type else '복합 수능번호가 서로 다른 표준 유형에 걸쳐 있습니다.',
        }
        candidates = declared_numbers
        classification_status = 'review'
        candidate_source = 'declared'
    elif alias:
        type_id, type_name, candidates = alias
        explicit_no = candidates[0] if len(candidates) == 1 else 0
        classified = {'type_id': type_id, 'type_name': type_name, 'type_mismatch': False,
                      'mismatch_reason': '' if explicit_no else f'{supplied_type} 유형의 표준 번호 {candidates} 중 하나를 확인해 주세요.'}
        classification_status = 'classified' if explicit_no else 'review'
        candidate_source = 'mapping'
    else:
        prompt_classification = classify_csat_prompt(instruction, choices)
        candidates = prompt_classification['candidate_numbers']
        explicit_no = candidates[0] if len(candidates) == 1 else 0
        classified = {'type_id': prompt_classification['type_id'], 'type_name': prompt_classification['type_name'],
                      'type_mismatch': prompt_classification['classification_status'] == 'review',
                      'mismatch_reason': prompt_classification['classification_reason']}
        classification_status = prompt_classification['classification_status']
        candidate_source = 'mapping'
    markup = normalized_problem.get('markup', raw.get('markup', [])) or []
    if not isinstance(markup, list) or any(not isinstance(item, dict) for item in markup):
        raise ValueError('problem.markup은 object array여야 합니다.')
    return {
        'source_question_no': source_no, 'source_question_number_raw': source_number_raw,
        'question_no': explicit_no,
        'csat_type_id': classified['type_id'], 'csat_type_name': classified['type_name'],
        'stem': instruction, 'prompt': instruction, 'passage': passage, 'choices': choices,
        'answer': str(raw.get('answer', '') or ''), 'explanation': str(raw.get('explanation', '') or ''),
        'problem': normalized_problem, 'source_problem': problem, 'markup': markup,
        'group_label': str(raw.get('group_label') or raw.get('common_passage_id') or problem.get('shared_group') or ''),
        'common_passage_id': str(raw.get('common_passage_id') or problem.get('shared_group') or ''),
        'classification_status': classification_status, 'candidate_numbers': candidates,
        'candidate_source': candidate_source,
        'classification_reason': classified['mismatch_reason'], **classified,
        '_source_metadata': inherited,
    }


def _resolve_csat_candidate_runs(questions: list[dict]) -> None:
    """Assign an explicit standard candidate range to its complete source-order run.

    Examples are the four blank questions (31–34), two order questions (36–37),
    and two insertion questions (38–39). This uses classifier-declared standard
    mappings, never common-passage group size. Explicit individual numbers win.
    """
    index = 0
    while index < len(questions):
        question = questions[index]
        if question.get('question_no'):
            index += 1
            continue
        candidates = question.get('candidate_numbers') or []
        type_id = question.get('csat_type_id')
        end = index + 1
        while end < len(questions) and not questions[end].get('question_no') \
                and questions[end].get('csat_type_id') == type_id \
                and questions[end].get('candidate_numbers') == candidates:
            end += 1
        group = questions[index:end]
        if len(candidates) == len(group) and len(group) > 1:
            for item, number in zip(group, candidates):
                standard = csat_type_for_number(number)
                item.update({
                    'question_no': number, 'csat_type_id': standard['id'], 'csat_type_name': standard['name'],
                    'type_id': standard['id'], 'type_name': standard['name'], 'candidate_numbers': [number],
                    'classification_status': 'classified', 'classification_reason': '',
                    'type_mismatch': False, 'mismatch_reason': '',
                })
        index = end


def _csat_json_review_reason(question: dict) -> str:
    number = int(question.get('question_no') or 0)
    if not 18 <= number <= 45:
        return '유효한 수능 문항 번호가 필요합니다.'
    if question.get('classification_status') == 'review' or question.get('type_mismatch') or question.get('csat_type_id') == 'review':
        return question.get('classification_reason') or question.get('mismatch_reason') or '표준 유형 검수가 필요합니다.'
    problem = question.get('problem') if isinstance(question.get('problem'), dict) else {}
    if not str(problem.get('instruction') or question.get('prompt') or '').strip() and not str(problem.get('passage') or question.get('passage') or '').strip():
        return '문제 발문 또는 본문이 필요합니다.'
    choices = problem.get('choices', question.get('choices'))
    if not isinstance(choices, list) or not any(str(choice).strip() for choice in choices):
        return '객관식 선택지가 필요합니다.'
    if not str(question.get('answer') or '').strip():
        return '정답이 필요합니다.'
    return ''


def parse_csat_question_json(value: object, ui_source_kind: str = '', ui_metadata: dict | None = None) -> list[dict]:
    """Map CSAT Question JSON Standard V1 (single or multi-round) without rewriting content."""
    if not isinstance(value, dict):
        raise ValueError('최상위 JSON은 object여야 합니다.')
    schema = str(value.get('schema') or '')
    if schema and schema != 'buildit.csat.question.v1':
        raise ValueError('지원하지 않는 schema입니다. buildit.csat.question.v1을 사용해 주세요.')
    base = {
        'source_kind': str(ui_source_kind or value.get('source_kind') or '').lower(),
        'source_title': str((ui_metadata or {}).get('title') or value.get('exam_name') or ''),
        'source_year': str((ui_metadata or {}).get('year') or value.get('year') or ''),
        'source_round': str((ui_metadata or {}).get('round') or value.get('round') or ''),
    }
    rounds = value.get('rounds')
    groups = rounds if isinstance(rounds, list) else [value]
    parsed = []
    for round_index, group in enumerate(groups, start=1):
        if not isinstance(group, dict):
            raise ValueError(f'{round_index}번째 회차는 object여야 합니다.')
        metadata = {
            'source_kind': str(group.get('source_kind') or base['source_kind']).lower(),
            'source_title': str(group.get('exam_name') or group.get('title') or base['source_title']),
            'source_year': str(group.get('year') or base['source_year']),
            'source_round': str(group.get('round') or base['source_round'] or round_index),
        }
        if metadata['source_kind'] not in ('half', 'full', 'wrong', 'daily'):
            raise ValueError('source_kind는 HALF/FULL/WRONG/DAILY 중 하나여야 합니다.')
        questions = group.get('questions')
        if not isinstance(questions, list) or not questions:
            raise ValueError(f'{metadata["source_round"]}회차에 questions가 없습니다.')
        round_questions = [_normalise_csat_json_question(raw, metadata, index)
                           for index, raw in enumerate(questions, start=1)]
        _resolve_csat_candidate_runs(round_questions)
        parsed.extend(round_questions)
    return parsed


def _extract_hwpx_text(data: bytes) -> str:
    try:
        from xml.etree import ElementTree
        with zipfile.ZipFile(__import__('io').BytesIO(data)) as archive:
            section_names = sorted(name for name in archive.namelist() if re.fullmatch(r'Contents/section\d+\.xml', name))
            if not section_names:
                raise ValueError('HWPX 본문 section XML이 없습니다.')
            chunks = []
            for name in section_names:
                root = ElementTree.fromstring(archive.read(name))
                chunks.extend(node.text for node in root.iter() if node.tag.endswith('}t') and node.text)
            return '\n'.join(chunks)
    except (zipfile.BadZipFile, ValueError, ElementTree.ParseError) as error:
        raise HTTPException(status_code=400, detail=f'HWPX 본문을 읽지 못했습니다: {error}') from error


def _hwp_record_payloads(stream: bytes):
    offset = 0
    while offset + 4 <= len(stream):
        header = struct.unpack_from('<I', stream, offset)[0]
        offset += 4
        tag, size = header & 0x3ff, (header >> 20) & 0xfff
        if size == 0xfff:
            if offset + 4 > len(stream):
                raise ValueError('HWP 레코드 길이가 잘렸습니다.')
            size = struct.unpack_from('<I', stream, offset)[0]
            offset += 4
        if offset + size > len(stream):
            raise ValueError('HWP 본문 레코드가 손상되었습니다.')
        yield tag, stream[offset:offset + size]
        offset += size


def _clean_hwp_para_text(payload: bytes) -> str:
    # HWP_TAG_PARA_TEXT payload is UTF-16LE. Inline controls occupy code points below
    # U+0020; tabs/newlines are retained and other binary control markers are removed.
    text = payload.decode('utf-16le', errors='strict')
    return ''.join('\n' if char in ('\r', '\n') else '\t' if char == '\t' else char
                   for char in text if ord(char) >= 32 or char in ('\r', '\n', '\t')).strip()


def _extract_hwp5_text(data: bytes) -> str:
    """Extract HWP 5.x OLE paragraph text without invoking desktop software.

    Encrypted/distributable documents are rejected. We never pretend that partially
    decoded binary is source text.
    """
    if not data.startswith(bytes.fromhex('D0CF11E0A1B11AE1')):
        raise HTTPException(status_code=400, detail='HWP 5.x OLE 문서가 아닙니다. 구형 HWP 3.x 또는 손상 파일은 지원하지 않습니다.')
    try:
        import olefile
        with olefile.OleFileIO(data) as ole:
            required = ('FileHeader',)
            if any(not ole.exists(name) for name in required):
                raise ValueError('HWP FileHeader가 없습니다.')
            header = ole.openstream('FileHeader').read()
            if not header.startswith(b'HWP Document File') or len(header) < 40:
                raise ValueError('HWP 5.x 헤더가 올바르지 않습니다.')
            flags = struct.unpack_from('<I', header, 36)[0]
            compressed, encrypted, distributable = bool(flags & 1), bool(flags & 2), bool(flags & 4)
            if encrypted:
                raise ValueError('암호화된 HWP는 지원하지 않습니다. 암호를 해제한 사본을 올려 주세요.')
            if distributable:
                raise ValueError('배포용 HWP는 본문 추출을 지원하지 않습니다. 편집 가능한 사본을 올려 주세요.')
            sections = sorted((entry for entry in ole.listdir() if len(entry) == 2 and entry[0] == 'BodyText' and re.fullmatch(r'Section\d+', entry[1])), key=lambda entry: int(entry[1][7:]))
            if not sections:
                raise ValueError('HWP BodyText section이 없습니다.')
            paragraphs = []
            for entry in sections:
                stream = ole.openstream(entry).read()
                if compressed:
                    try: stream = zlib.decompress(stream, -15)
                    except zlib.error as error: raise ValueError('압축된 HWP 본문이 손상되었습니다.') from error
                for tag, payload in _hwp_record_payloads(stream):
                    if tag == 67:  # HWP_TAG_PARA_TEXT
                        text = _clean_hwp_para_text(payload)
                        if text: paragraphs.append(text)
            if not paragraphs:
                raise ValueError('HWP에서 읽을 수 있는 문단 텍스트를 찾지 못했습니다.')
            return '\n'.join(paragraphs)
    except HTTPException:
        raise
    except (OSError, ValueError, UnicodeDecodeError) as error:
        raise HTTPException(status_code=400, detail=f'HWP 본문을 읽지 못했습니다: {error}') from error


def _normalise_uploaded_questions(raw_text: str) -> list[dict]:
    """HWPX/TXT/JSON 어댑터의 공통 경계. JSON은 완전 구조화, 문서는 번호 블록을 보수적으로 분리한다."""
    stripped = raw_text.strip()
    if stripped.startswith(('{', '[')):
        try:
            value = json.loads(stripped)
            return value.get('questions', []) if isinstance(value, dict) else value
        except json.JSONDecodeError:
            pass
    matches = list(re.finditer(r'(?m)^\s*(1[8-9]|[2-3]\d|4[0-5])\s*[\.\)]\s*', stripped))
    result = []
    for index, match in enumerate(matches):
        block = stripped[match.end():matches[index + 1].start() if index + 1 < len(matches) else len(stripped)].strip()
        lines = [line.strip() for line in block.splitlines() if line.strip()]
        prompt = lines[0] if lines else ''
        choices = [re.sub(r'^[①②③④⑤]\s*', '', line) for line in lines if re.match(r'^[①②③④⑤]', line)]
        passage = '\n'.join(line for line in lines[1:] if not re.match(r'^[①②③④⑤]', line))
        result.append({'question_no': int(match.group(1)), 'stem': prompt, 'prompt': prompt, 'passage': passage, 'choices': choices})
    return result


def _canonical_question(raw: dict, exam_kind: str, metadata: dict, filename: str, raw_text: str, now: str) -> dict:
    number = int(raw.get('question_number', raw.get('question_no', 0)) or 0)
    choices = [str(choice) for choice in raw.get('choices', [])]
    problem = raw.get('problem') if isinstance(raw.get('problem'), dict) else {}
    stem = str(raw.get('stem', raw.get('prompt', problem.get('instruction', ''))))
    classification = classify_csat_question(number, stem, choices)
    return {
        'source_exam': str(metadata.get('title', metadata.get('source_exam', ''))),
        'year': str(metadata.get('year', '')), 'session': str(metadata.get('round', metadata.get('session', ''))),
        'exam_kind': exam_kind, 'question_number': number, 'question_no': number,
        'csat_type': classification['type_id'], 'csat_type_id': classification['type_id'],
        'csat_type_name': classification['type_name'], 'stem': stem, 'prompt': stem,
        'passage': str(raw.get('passage', problem.get('passage', ''))), 'choices': choices,
        'answer': str(raw.get('answer', '')), 'explanation': str(raw.get('explanation', raw.get('solution', ''))),
        'solution': str(raw.get('explanation', raw.get('solution', ''))), 'footnote': str(raw.get('footnote', '')),
        'group_label': str(raw.get('group_label', '')), 'source_file': filename,
        'source_reference': f'{filename}#question-{number}', 'raw_text': raw_text,
        'source_question_no': raw.get('source_question_no'), 'question_html': str(raw.get('question_html', '')),
        'explanation_html': str(raw.get('explanation_html', '')),
        # Structured JSON V1 content stays in raw_json when the canonical object is committed.
        # This is the lossless master used by renderers; DB columns remain search indexes.
        'problem': problem, 'markup': raw.get('markup', problem.get('markup', [])),
        'common_passage_id': str(raw.get('common_passage_id', '')),
        'classification_status': raw.get('classification_status', 'review' if classification.get('type_mismatch') else 'classified'),
        'created_at': now, **classification,
    }


CsatSourceKind = Literal['half', 'full', 'wrong', 'daily']


def _source_storage_kind(source_kind: str) -> Literal['half', 'full']:
    """Keep the applied half/full schema while source_kind records ingestion intent in metadata."""
    return 'full' if source_kind == 'full' else 'half'


class CsatImportCommit(BaseModel):
    exam_kind: Literal['half', 'full']
    source_kind: Optional[CsatSourceKind] = None
    source_filename: str
    source_hash: str = Field(pattern=r'^[0-9a-f]{64}$')
    parser: str
    metadata: dict = {}
    raw_text: str
    questions: list[dict]


def _csat_question_response(row: sqlite3.Row) -> dict:
    item = dict(row)
    item['choices'] = json.loads(item.pop('choices_json'))
    item['exam_metadata'] = json.loads(item.pop('exam_metadata_json', '{}')) if 'exam_metadata_json' in item else {}
    raw = json.loads(item.get('raw_json') or '{}')
    item['source_kind'] = item['exam_metadata'].get('source_kind') or item.get('exam_kind')
    item['source_title'] = (item['exam_metadata'].get('source_title') or item['exam_metadata'].get('title')
                            or item['exam_metadata'].get('source_name') or item['exam_metadata'].get('exam_name') or '')
    item['source_year'] = str(item['exam_metadata'].get('source_year') or item['exam_metadata'].get('year') or '')
    item['source_round'] = str(item['exam_metadata'].get('source_round') or item['exam_metadata'].get('round') or '')
    item['question_html'] = raw.get('question_html') or ''
    item['explanation_html'] = raw.get('explanation_html') or ''
    item['has_answer'] = bool(str(item.get('answer') or '').strip())
    item['has_explanation'] = bool(str(item.get('solution') or '').strip() or str(raw.get('explanation_html') or '').strip())
    item['review_status'] = 'review' if item.get('type_mismatch') or raw.get('classification_status') == 'review' else 'classified'
    item['type_mismatch'] = bool(item['type_mismatch'])
    return item


@app.get('/api/csat/types')
def list_csat_types():
    return [{'numbers': list(numbers), 'id': type_id, 'name': name} for numbers, type_id, name, _ in CSAT_TYPE_RANGES]


@app.post('/api/csat/uploads', status_code=201)
async def upload_csat_exam(exam_kind: Literal['half', 'full'] = Form(...), metadata_json: str = Form('{}'), file: UploadFile = File(...)):
    filename = file.filename or 'upload'
    suffix = Path(filename).suffix.lower()
    data = await file.read()
    if suffix not in ('.hwp', '.hwpx', '.json', '.txt'):
        raise HTTPException(status_code=415, detail='현재 HWP, HWPX, JSON, TXT를 지원합니다.')
    if not data:
        raise HTTPException(status_code=400, detail='빈 파일은 업로드할 수 없습니다.')
    if len(data) > 30 * 1024 * 1024:
        raise HTTPException(status_code=413, detail='수능 자료 파일은 30MB 이하여야 합니다.')
    parser = {'hwp': 'hwp5-ole-v1', 'hwpx': 'hwpx-xml-v1', 'json': 'json', 'txt': 'txt'}[suffix[1:]]
    raw_text = _extract_hwp5_text(data) if suffix == '.hwp' else _extract_hwpx_text(data) if suffix == '.hwpx' else data.decode('utf-8-sig', errors='strict')
    questions = _normalise_uploaded_questions(raw_text)
    if not questions:
        raise HTTPException(status_code=400, detail='18~45번 문항 블록을 찾지 못했습니다. 구조화 JSON 또는 번호가 포함된 HWPX를 확인해 주세요.')
    try:
        metadata = json.loads(metadata_json or '{}')
    except json.JSONDecodeError as error:
        raise HTTPException(status_code=400, detail='시험 메타데이터 JSON 형식이 올바르지 않습니다.') from error
    now = datetime.now().isoformat()
    canonical = [_canonical_question(raw, exam_kind, metadata, filename, raw_text, now) for raw in questions]
    canonical = [item for item in canonical if 1 <= item['question_number'] <= 99]
    failures = [{'index': index + 1, 'reason': '문항번호를 인식하지 못했습니다.'} for index, raw in enumerate(questions) if not int(raw.get('question_number', raw.get('question_no', 0)) or 0)]
    if not canonical:
        raise HTTPException(status_code=400, detail='저장 가능한 문항을 찾지 못했습니다.')
    return {'source_id': None, 'source_hash': hashlib.sha256(data).hexdigest(), 'parser': parser, 'raw_text': raw_text,
            'metadata': metadata, 'source_filename': filename, 'original_stored': False, 'owner_scope': 'shared',
            'questions': canonical, 'failures': failures, 'requires_review': True}


@app.post('/api/csat/json-imports/preview', status_code=201)
async def preview_csat_json_imports(source_kind: Optional[CsatSourceKind] = Form(None),
                                    metadata_json: str = Form('{}'), files: list[UploadFile] = File(...)):
    if not files or len(files) > 50:
        raise HTTPException(status_code=400, detail='JSON 파일은 한 번에 1~50개까지 업로드할 수 있습니다.')
    try:
        ui_metadata = json.loads(metadata_json or '{}')
    except json.JSONDecodeError as error:
        raise HTTPException(status_code=400, detail='시험 메타데이터 JSON 형식이 올바르지 않습니다.') from error
    with get_db() as conn:
        existing_rows = conn.execute('''SELECT q.raw_json, s.exam_metadata_json FROM csat_questions q
            JOIN csat_question_sources s ON s.id=q.source_id''').fetchall()
    existing = set()
    for row in existing_rows:
        raw = json.loads(row['raw_json'] or '{}')
        metadata = json.loads(row['exam_metadata_json'] or '{}')
        existing.add((str(metadata.get('source_kind') or '').lower(),
                      str(metadata.get('source_title') or metadata.get('title') or ''),
                      str(metadata.get('source_year') or metadata.get('year') or ''),
                      str(metadata.get('source_round') or metadata.get('round') or ''),
                      int(raw.get('source_question_no') or raw.get('question_number') or 0)))
    previews, total_bytes = [], 0
    for file in files:
        filename = file.filename or 'questions.json'
        data = await file.read()
        total_bytes += len(data)
        failure, raw_text = '', ''
        if Path(filename).suffix.lower() != '.json': failure = 'JSON 파일만 업로드할 수 있습니다.'
        elif not data: failure = '빈 JSON 파일입니다.'
        elif len(data) > 10 * 1024 * 1024: failure = 'JSON은 파일당 10MB 이하여야 합니다.'
        elif total_bytes > 30 * 1024 * 1024: failure = '전체 JSON 크기는 30MB 이하여야 합니다.'
        try:
            raw_text = data.decode('utf-8-sig', errors='strict') if not failure else ''
            value = json.loads(raw_text) if raw_text else None
            parsed = parse_csat_question_json(value, source_kind or '', ui_metadata) if value is not None else []
        except (UnicodeDecodeError, json.JSONDecodeError, ValueError, TypeError) as error:
            failure, parsed = f'JSON 검증 실패: {error}', []
        if failure:
            previews.append({'source_id': None, 'source_hash': hashlib.sha256(data).hexdigest(), 'parser': 'csat-json-v1',
                'raw_text': raw_text, 'metadata': {**ui_metadata, 'source_kind': (source_kind or '')},
                'source_filename': filename, 'display_source_name': filename, 'original_stored': False,
                'owner_scope': 'shared', 'questions': [], 'failures': [{'index': 1, 'reason': failure}],
                'requires_review': False, 'status': 'FAILED', 'duplicate_source_id': None})
            continue
        groups = {}
        for question in parsed:
            metadata = question.pop('_source_metadata')
            key = tuple(metadata[field] for field in ('source_kind', 'source_title', 'source_year', 'source_round'))
            groups.setdefault(key, {'metadata': metadata, 'questions': []})['questions'].append(question)
        for key, group in groups.items():
            metadata, questions = group['metadata'], group['questions']
            for question in questions:
                provenance = (*key, int(question['source_question_no']))
                question['source_kind'] = metadata['source_kind']
                question['exam_kind'] = _source_storage_kind(metadata['source_kind'])
                review_reason = _csat_json_review_reason(question)
                question['review_reason'] = review_reason
                question['import_status'] = ('DUPLICATE' if provenance in existing else
                                             'REVIEW' if review_reason else 'READY')
            identity = json.dumps(metadata, sort_keys=True, ensure_ascii=False).encode()
            source_hash = hashlib.sha256(data + b'\0' + identity).hexdigest()
            storage_kind = _source_storage_kind(metadata['source_kind'])
            with get_db() as conn:
                duplicate_source = conn.execute('SELECT id FROM csat_question_sources WHERE source_hash=? AND exam_kind=?',
                                                (source_hash, storage_kind)).fetchone()
            states = [question['import_status'] for question in questions]
            status = 'DUPLICATE' if states and all(state == 'DUPLICATE' for state in states) else 'REVIEW' if 'REVIEW' in states else 'READY'
            display_name = f'{filename} · {metadata["source_round"]}회'
            previews.append({'source_id': None, 'source_hash': source_hash, 'parser': 'csat-json-v1',
                'raw_text': raw_text, 'metadata': metadata, 'source_filename': filename,
                'display_source_name': display_name, 'original_stored': False, 'owner_scope': 'shared',
                'questions': questions, 'failures': [], 'requires_review': 'REVIEW' in states,
                'status': 'DUPLICATE' if duplicate_source else status,
                'duplicate_source_id': duplicate_source['id'] if duplicate_source else None})
    return {'files': previews, 'file_count': len(previews),
            'question_count': sum(len(item['questions']) for item in previews),
            'review_count': sum(state == 'REVIEW' for item in previews for state in [q.get('import_status') for q in item['questions']]),
            'failed_count': sum(item['status'] == 'FAILED' for item in previews),
            'duplicate_count': sum(q.get('import_status') == 'DUPLICATE' for item in previews for q in item['questions'])}


@app.post('/api/csat/html-imports/preview', status_code=201)
async def preview_csat_html_imports(exam_kind: Optional[Literal['half', 'full']] = Form(None),
                                    source_kind: Optional[CsatSourceKind] = Form(None),
                                    metadata_json: str = Form('{}'), files: list[UploadFile] = File(...)):
    if not files or len(files) > 50:
        raise HTTPException(status_code=400, detail='HTML 파일은 한 번에 1~50개까지 업로드할 수 있습니다.')
    try:
        metadata = json.loads(metadata_json or '{}')
    except json.JSONDecodeError as error:
        raise HTTPException(status_code=400, detail='시험 메타데이터 JSON 형식이 올바르지 않습니다.') from error
    selected_source_kind = source_kind or exam_kind
    if not selected_source_kind:
        raise HTTPException(status_code=400, detail='자료 종류(HALF/FULL/WRONG/DAILY)를 선택해 주세요.')
    storage_kind = _source_storage_kind(selected_source_kind)
    metadata = {**metadata, 'source_kind': selected_source_kind}
    previews, total_bytes = [], 0
    for file in files:
        filename = file.filename or 'upload.html'
        data = await file.read()
        total_bytes += len(data)
        source_hash = hashlib.sha256(data).hexdigest()
        failure = ''
        if Path(filename).suffix.lower() not in ('.html', '.htm'):
            failure = 'HTML 파일만 Batch Import할 수 있습니다.'
        elif not data:
            failure = '빈 HTML 파일입니다.'
        elif len(data) > 10 * 1024 * 1024:
            failure = 'HTML은 파일당 10MB 이하여야 합니다.'
        elif total_bytes > 30 * 1024 * 1024:
            failure = 'Batch HTML 전체 크기는 30MB 이하여야 합니다.'
        if failure:
            previews.append({'source_id': None, 'source_hash': source_hash, 'parser': 'editable-html-v1',
                'raw_text': '', 'metadata': metadata, 'source_filename': filename, 'original_stored': False,
                'owner_scope': 'shared', 'questions': [], 'failures': [{'index': 1, 'reason': failure}],
                'requires_review': False, 'status': 'FAILED', 'duplicate_source_id': None})
            continue
        try:
            raw_html = data.decode('utf-8-sig', errors='strict')
        except UnicodeDecodeError:
            previews.append({'source_id': None, 'source_hash': source_hash, 'parser': 'editable-html-v1',
                'raw_text': '', 'metadata': metadata, 'source_filename': filename, 'original_stored': False,
                'owner_scope': 'shared', 'questions': [], 'failures': [{'index': 1, 'reason': 'UTF-8 HTML이 아닙니다.'}],
                'requires_review': False, 'status': 'FAILED', 'duplicate_source_id': None})
            continue
        questions = split_csat_html(raw_html)
        with get_db() as conn:
            duplicate = conn.execute('SELECT id FROM csat_question_sources WHERE source_hash=? AND exam_kind=?',
                                     (source_hash, storage_kind)).fetchone()
        review = any(question['classification_status'] == 'review' for question in questions)
        status = 'DUPLICATE' if duplicate else 'FAILED' if not questions else 'REVIEW' if review else 'READY'
        preview_questions = [{**question, 'exam_kind': storage_kind, 'source_kind': selected_source_kind} for question in questions]
        previews.append({
            'source_id': None, 'source_hash': source_hash, 'parser': 'editable-html-v1',
            'raw_text': raw_html, 'metadata': metadata, 'source_filename': filename, 'original_stored': False,
            'owner_scope': 'shared', 'questions': preview_questions,
            'failures': [] if questions else [{'index': 1, 'reason': '문항 container를 찾지 못했습니다.'}],
            'requires_review': review, 'status': status,
            'duplicate_source_id': duplicate['id'] if duplicate else None,
        })
    return {'files': previews, 'file_count': len(previews),
            'question_count': sum(len(item['questions']) for item in previews),
            'review_count': sum(1 for item in previews for question in item['questions'] if question['classification_status'] == 'review'),
            'failed_count': sum(1 for item in previews if item['status'] == 'FAILED'),
            'duplicate_count': sum(1 for item in previews if item['status'] == 'DUPLICATE')}


@app.post('/api/csat/imports', status_code=201)
def commit_csat_import(payload: CsatImportCommit):
    now = datetime.now().isoformat()
    source_kind = payload.source_kind or payload.metadata.get('source_kind') or payload.exam_kind
    storage_kind = _source_storage_kind(source_kind)
    metadata = {**payload.metadata, 'source_kind': source_kind}
    if payload.parser == 'csat-json-v1':
        invalid = [{'source_question_no': raw.get('source_question_no'), 'reason': _csat_json_review_reason(raw)}
                   for raw in payload.questions if _csat_json_review_reason(raw)]
        if invalid:
            raise HTTPException(status_code=422, detail={'message': 'READY 문항만 저장할 수 있습니다.', 'questions': invalid})
    canonical = [_canonical_question(raw, storage_kind, metadata, payload.source_filename, payload.raw_text, now) for raw in payload.questions]
    canonical = [item for item in canonical if 1 <= item['question_number'] <= 99]
    if not canonical:
        raise HTTPException(status_code=400, detail='검수 후 저장할 문항이 없습니다.')
    saved = []
    with get_db() as conn:
        duplicate = conn.execute('SELECT id FROM csat_question_sources WHERE source_hash=? AND exam_kind=?', (payload.source_hash, storage_kind)).fetchone()
        if duplicate:
            raise HTTPException(status_code=409, detail=f'같은 원본 파일이 이미 저장되어 있습니다. (source_id={duplicate["id"]})')
        source = conn.execute(
            'INSERT INTO csat_question_sources (exam_kind, source_filename, exam_metadata_json, parser, raw_text, created_at, source_hash, original_stored, owner_scope) VALUES (?,?,?,?,?,?,?,?,?)',
            (storage_kind, payload.source_filename, json.dumps(metadata, ensure_ascii=False), payload.parser, payload.raw_text, now, payload.source_hash, 0, 'shared'),
        )
        source_id = source.lastrowid
        for raw in canonical:
            number, choices = raw['question_number'], raw['choices']
            cursor = conn.execute('''INSERT INTO csat_questions
                (source_id, exam_kind, question_no, csat_type_id, csat_type_name, passage, prompt, choices_json,
                 answer, solution, footnote, group_label, source_reference, type_mismatch, mismatch_reason, raw_json, created_at)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',
                (source_id, storage_kind, number, raw['type_id'], raw['type_name'], raw['passage'], raw['stem'],
                 json.dumps(choices, ensure_ascii=False), raw['answer'], raw['explanation'], raw['footnote'], raw['group_label'],
                 raw['source_reference'], int(raw['type_mismatch']), raw['mismatch_reason'], json.dumps(raw, ensure_ascii=False), now))
            saved.append(cursor.lastrowid)
        if not saved:
            raise HTTPException(status_code=400, detail='18~45번 범위의 저장 가능한 문항이 없습니다.')
        rows = conn.execute(f'''SELECT q.*, s.source_filename, s.exam_metadata_json FROM csat_questions q
            JOIN csat_question_sources s ON s.id=q.source_id WHERE q.id IN ({','.join('?' * len(saved))}) ORDER BY q.question_no''', saved).fetchall()
    return {'source_id': source_id, 'parser': payload.parser, 'original_stored': False, 'owner_scope': 'shared', 'questions': [_csat_question_response(row) for row in rows]}


@app.get('/api/csat/questions')
def list_csat_questions(exam_kind: Optional[Literal['half', 'full']] = None, question_no: Optional[int] = None,
                        type_id: Optional[str] = None, mismatch: Optional[bool] = None,
                        exclude_used: bool = False, used_in: Optional[Literal['half', 'full', 'wrong', 'daily']] = None,
                        source_id: Optional[int] = None, source_filename: Optional[str] = None,
                        source_round: Optional[str] = None, has_answer: Optional[bool] = None,
                        has_explanation: Optional[bool] = None,
                        review_status: Optional[Literal['classified', 'review']] = None,
                        source_kind: Optional[CsatSourceKind] = None, source_title: Optional[str] = None,
                        source_year: Optional[str] = None, used: Optional[bool] = None):
    filters, values = [], []
    for clause, value in (('q.exam_kind=?', exam_kind), ('q.question_no=?', question_no), ('q.csat_type_id=?', type_id)):
        if value is not None: filters.append(clause); values.append(value)
    if source_id is not None: filters.append('q.source_id=?'); values.append(source_id)
    if source_filename: filters.append('s.source_filename=?'); values.append(source_filename)
    if mismatch is not None: filters.append('q.type_mismatch=?'); values.append(int(mismatch))
    if exclude_used:
        if used_in:
            filters.append('NOT EXISTS (SELECT 1 FROM csat_question_usage u WHERE u.question_id=q.id AND u.assessment_type=?)'); values.append(used_in)
        else:
            filters.append('NOT EXISTS (SELECT 1 FROM csat_question_usage u WHERE u.question_id=q.id)')
    if used is not None:
        filters.append(('EXISTS' if used else 'NOT EXISTS') + ' (SELECT 1 FROM csat_question_usage u WHERE u.question_id=q.id)')
    where = f"WHERE {' AND '.join(filters)}" if filters else ''
    with get_db() as conn:
        rows = conn.execute(f'''SELECT q.*, s.source_filename, s.exam_metadata_json,
            (SELECT COUNT(*) FROM csat_question_usage u WHERE u.question_id=q.id) usage_count
            FROM csat_questions q
            JOIN csat_question_sources s ON s.id=q.source_id {where} ORDER BY q.created_at DESC, q.question_no''', values).fetchall()
    result = [_csat_question_response(row) for row in rows]
    if source_round:
        result = [item for item in result if item['source_round'] == source_round]
    if source_kind:
        result = [item for item in result if item['source_kind'] == source_kind]
    if source_title:
        result = [item for item in result if item['source_title'] == source_title]
    if source_year:
        result = [item for item in result if item['source_year'] == source_year]
    if has_answer is not None:
        result = [item for item in result if item['has_answer'] is has_answer]
    if has_explanation is not None:
        result = [item for item in result if item['has_explanation'] is has_explanation]
    if review_status:
        result = [item for item in result if item['review_status'] == review_status]
    return result


AssessmentType = Literal['half', 'full', 'wrong', 'daily']


def _csat_delete_protection(conn: sqlite3.Connection, question_ids: list[int]) -> list[str]:
    if not question_ids:
        return []
    placeholders = ','.join('?' * len(question_ids))
    reasons = []
    if conn.execute(f'SELECT 1 FROM csat_assessment_items WHERE question_id IN ({placeholders}) LIMIT 1', question_ids).fetchone():
        reasons.append('Assessment 사용 이력')
    elif conn.execute(f'SELECT 1 FROM csat_question_usage WHERE question_id IN ({placeholders}) LIMIT 1', question_ids).fetchone():
        reasons.append('Assessment 사용 이력')
    if conn.execute(f'SELECT 1 FROM csat_wrong_answers WHERE question_id IN ({placeholders}) LIMIT 1', question_ids).fetchone():
        reasons.append('학생 오답 연결')
    return reasons


def _raise_csat_delete_conflict(reasons: list[str]):
    if reasons:
        raise HTTPException(status_code=409, detail=f"사용 중인 문항은 삭제할 수 없습니다. ({', '.join(reasons)})")


@app.delete('/api/csat/questions/{question_id}')
def delete_csat_question(question_id: int):
    with get_db() as conn:
        question = conn.execute('SELECT id, source_id FROM csat_questions WHERE id=?', (question_id,)).fetchone()
        if not question:
            raise HTTPException(status_code=404, detail='수능 문항을 찾을 수 없습니다.')
        _raise_csat_delete_conflict(_csat_delete_protection(conn, [question_id]))
        conn.execute('DELETE FROM csat_questions WHERE id=?', (question_id,))
    return {'deleted_question_id': question_id, 'source_id': question['source_id']}


@app.delete('/api/csat/question-sources/{source_id}')
def delete_csat_question_source(source_id: int):
    with get_db() as conn:
        source = conn.execute('SELECT id FROM csat_question_sources WHERE id=?', (source_id,)).fetchone()
        if not source:
            raise HTTPException(status_code=404, detail='수능 회차 원본을 찾을 수 없습니다.')
        question_ids = [row['id'] for row in conn.execute(
            'SELECT id FROM csat_questions WHERE source_id=? ORDER BY id', (source_id,)
        ).fetchall()]
        _raise_csat_delete_conflict(_csat_delete_protection(conn, question_ids))
        conn.execute('DELETE FROM csat_questions WHERE source_id=?', (source_id,))
        conn.execute('DELETE FROM csat_question_sources WHERE id=?', (source_id,))
    return {'deleted_source_id': source_id, 'deleted_question_ids': question_ids, 'deleted_count': len(question_ids)}


class CsatAssessmentCreate(BaseModel):
    assessment_type: AssessmentType
    title: str = ''
    template_id: str = ''
    question_ids: list[int] = Field(default_factory=list)
    question_no: Optional[int] = Field(default=None, ge=18, le=45)
    type_id: str = ''
    count: int = Field(default=4, ge=1, le=45)
    days: int = Field(default=1, ge=1, le=30)
    exclude_used: bool = False


def _assessment_detail(conn: sqlite3.Connection, assessment_id: int) -> dict:
    assessment = conn.execute('SELECT * FROM csat_assessments WHERE id=?', (assessment_id,)).fetchone()
    if not assessment:
        raise HTTPException(status_code=404, detail='수능 Assessment를 찾을 수 없습니다.')
    items = conn.execute('''SELECT i.position, i.day_no, q.*, s.source_filename, s.exam_metadata_json
        FROM csat_assessment_items i JOIN csat_questions q ON q.id=i.question_id
        JOIN csat_question_sources s ON s.id=q.source_id
        WHERE i.assessment_id=? ORDER BY i.position''', (assessment_id,)).fetchall()
    result = dict(assessment)
    result['request'] = json.loads(result.pop('request_json'))
    result['items'] = [{**_csat_question_response(row), 'position': row['position'], 'day_no': row['day_no']} for row in items]
    result['question_ids'] = [item['id'] for item in result['items']]
    result['problem_html_url'] = f'/problem-generator/wiki/api/csat/assessments/{assessment_id}/problem.html'
    result['explanation_html_url'] = f'/problem-generator/wiki/api/csat/assessments/{assessment_id}/explanation.html'
    return result


def compose_csat_assessment(payload: CsatAssessmentCreate) -> dict:
    """Select and persist atomically; only this saved Assessment path creates usage rows."""
    target_count = 28 if payload.assessment_type == 'full' else payload.days * 4 if payload.assessment_type == 'daily' else payload.count
    requested_ids = list(dict.fromkeys(payload.question_ids))
    if requested_ids and len(requested_ids) != target_count:
        raise HTTPException(status_code=400, detail=f'저장할 Assessment에는 정확히 {target_count}개 Question ID가 필요합니다.')
    if payload.assessment_type == 'full' and (payload.question_no is not None or payload.type_id):
        raise HTTPException(status_code=400, detail='FULL은 특정 번호/유형 필터가 아니라 18~45번을 각각 한 문항씩 구성합니다.')
    with get_db() as conn:
        if payload.assessment_type == 'full':
            if requested_ids:
                rows = conn.execute(f'''SELECT q.id, q.question_no FROM csat_questions q
                    WHERE q.id IN ({','.join('?' * len(requested_ids))})
                    {'AND NOT EXISTS (SELECT 1 FROM csat_question_usage u WHERE u.question_id=q.id)' if payload.exclude_used else ''}
                    ORDER BY q.question_no''', requested_ids).fetchall()
                if [row['question_no'] for row in rows] != list(range(18, 46)):
                    raise HTTPException(status_code=409, detail='FULL은 18~45번 Question이 번호별로 정확히 하나씩 필요합니다.')
                selected = [row['id'] for row in rows]
            else:
                selected = []
                for question_no in range(18, 46):
                    row = conn.execute(f'''SELECT q.id FROM csat_questions q WHERE q.question_no=?
                        {'AND NOT EXISTS (SELECT 1 FROM csat_question_usage u WHERE u.question_id=q.id)' if payload.exclude_used else ''}
                        ORDER BY q.created_at, q.id LIMIT 1''', (question_no,)).fetchone()
                    if not row:
                        raise HTTPException(status_code=409, detail=f'FULL 구성에 필요한 {question_no}번 미사용 문항이 없습니다.')
                    selected.append(row['id'])
        else:
            filters, values = [], []
            if requested_ids:
                filters.append(f"q.id IN ({','.join('?' * len(requested_ids))})"); values.extend(requested_ids)
            if payload.question_no is not None:
                filters.append('q.question_no=?'); values.append(payload.question_no)
            if payload.type_id:
                filters.append('q.csat_type_id=?'); values.append(payload.type_id)
            if payload.exclude_used:
                filters.append('NOT EXISTS (SELECT 1 FROM csat_question_usage u WHERE u.question_id=q.id)')
            where = f"WHERE {' AND '.join(filters)}" if filters else ''
            rows = conn.execute(f'''SELECT q.id FROM csat_questions q {where}
                ORDER BY q.created_at, q.id LIMIT ?''', (*values, target_count)).fetchall()
            selected = [row['id'] for row in rows]
            if requested_ids:
                # SQL order is deterministic, while explicit editing order must be preserved.
                selected_set = set(selected)
                selected = [question_id for question_id in requested_ids if question_id in selected_set]
        if len(selected) < target_count:
            raise HTTPException(status_code=409, detail=f'조건에 맞는 문항이 {len(selected)}개뿐입니다. {target_count}개가 필요합니다.')
        now = datetime.now().isoformat()
        title = payload.title.strip() or f'BuildIt {payload.assessment_type.upper()}'
        template_id = payload.template_id.strip() or f'csat-{payload.assessment_type}-v1'
        cursor = conn.execute('''INSERT INTO csat_assessments
            (assessment_type,title,template_id,request_json,created_at,updated_at) VALUES(?,?,?,?,?,?)''',
            (payload.assessment_type, title, template_id, payload.model_dump_json(), now, now))
        assessment_id = cursor.lastrowid
        for position, question_id in enumerate(selected, start=1):
            day_no = ((position - 1) // 4) + 1 if payload.assessment_type == 'daily' else None
            conn.execute('''INSERT INTO csat_assessment_items
                (assessment_id,question_id,position,day_no,created_at) VALUES(?,?,?,?,?)''',
                (assessment_id, question_id, position, day_no, now))
            conn.execute('''INSERT INTO csat_question_usage
                (question_id,assessment_id,assessment_type,used_at) VALUES(?,?,?,?)''',
                (question_id, assessment_id, payload.assessment_type, now))
        return _assessment_detail(conn, assessment_id)


@app.post('/api/csat/assessments', status_code=201)
def create_csat_assessment(payload: CsatAssessmentCreate):
    return compose_csat_assessment(payload)


@app.get('/api/csat/assessments')
def list_csat_assessments(assessment_type: Optional[AssessmentType] = None):
    with get_db() as conn:
        where, values = ('WHERE a.assessment_type=?', [assessment_type]) if assessment_type else ('', [])
        return [dict(row) for row in conn.execute(f'''SELECT a.*,
            (SELECT COUNT(*) FROM csat_assessment_items i WHERE i.assessment_id=a.id) question_count
            FROM csat_assessments a {where} ORDER BY a.id DESC''', values).fetchall()]


@app.get('/api/csat/assessments/{assessment_id}')
def get_csat_assessment(assessment_id: int):
    with get_db() as conn:
        return _assessment_detail(conn, assessment_id)


class CsatAssessmentItemReplace(BaseModel):
    question_id: int


@app.patch('/api/csat/assessments/{assessment_id}/items/{position}')
def replace_csat_assessment_item(assessment_id: int, position: int, payload: CsatAssessmentItemReplace):
    """Replace one saved item; both problem and explanation renderers follow this single row."""
    with get_db() as conn:
        assessment = conn.execute('SELECT * FROM csat_assessments WHERE id=?', (assessment_id,)).fetchone()
        current = conn.execute('SELECT * FROM csat_assessment_items WHERE assessment_id=? AND position=?', (assessment_id, position)).fetchone()
        question = conn.execute('SELECT id, question_no FROM csat_questions WHERE id=?', (payload.question_id,)).fetchone()
        if not assessment or not current or not question:
            raise HTTPException(status_code=404, detail='Assessment item 또는 Question을 찾을 수 없습니다.')
        if assessment['assessment_type'] == 'full':
            current_question = conn.execute('SELECT question_no FROM csat_questions WHERE id=?', (current['question_id'],)).fetchone()
            if question['question_no'] != current_question['question_no']:
                raise HTTPException(status_code=400, detail='FULL 문항 교체는 같은 수능 번호의 Question으로만 가능합니다.')
        duplicate = conn.execute('SELECT 1 FROM csat_assessment_items WHERE assessment_id=? AND question_id=? AND position<>?',
                                 (assessment_id, payload.question_id, position)).fetchone()
        if duplicate:
            raise HTTPException(status_code=409, detail='같은 Assessment에 동일 Question을 중복 배치할 수 없습니다.')
        now = datetime.now().isoformat()
        conn.execute('UPDATE csat_assessment_items SET question_id=?, created_at=? WHERE id=?', (payload.question_id, now, current['id']))
        conn.execute('DELETE FROM csat_question_usage WHERE assessment_id=? AND question_id=?', (assessment_id, current['question_id']))
        conn.execute('''INSERT INTO csat_question_usage(question_id,assessment_id,assessment_type,used_at)
            VALUES(?,?,?,?)''', (payload.question_id, assessment_id, assessment['assessment_type'], now))
        conn.execute('UPDATE csat_assessments SET updated_at=? WHERE id=?', (now, assessment_id))
        return _assessment_detail(conn, assessment_id)


@app.get('/api/csat/question-usage')
def list_csat_question_usage(question_id: Optional[int] = None, assessment_type: Optional[AssessmentType] = None):
    filters, values = [], []
    if question_id is not None: filters.append('u.question_id=?'); values.append(question_id)
    if assessment_type: filters.append('u.assessment_type=?'); values.append(assessment_type)
    where = f"WHERE {' AND '.join(filters)}" if filters else ''
    with get_db() as conn:
        return [dict(row) for row in conn.execute(f'''SELECT u.*, a.title FROM csat_question_usage u
            JOIN csat_assessments a ON a.id=u.assessment_id {where} ORDER BY u.id DESC''', values).fetchall()]


class CsatWrongAnswerCreate(BaseModel):
    question_id: Optional[int] = None
    student_name: str
    exam_name: str = ''
    question_no: int = Field(ge=18, le=45)
    wrong_answer: str = ''
    correct_answer: str = ''
    memo: str = ''
    source_filename: str = ''


@app.post('/api/csat/wrong-answers', status_code=201)
def create_csat_wrong_answer(payload: CsatWrongAnswerCreate):
    if not payload.student_name.strip():
        raise HTTPException(status_code=400, detail='학생 이름을 입력해 주세요.')
    csat_type = csat_type_for_number(payload.question_no)
    now = datetime.now().isoformat()
    with get_db() as conn:
        if payload.question_id is not None:
            question = conn.execute('SELECT id, question_no FROM csat_questions WHERE id=?', (payload.question_id,)).fetchone()
            if not question:
                raise HTTPException(status_code=400, detail='연결할 Question ID를 찾을 수 없습니다.')
            if question['question_no'] != payload.question_no:
                raise HTTPException(status_code=400, detail='오답 번호와 연결할 Question의 수능 번호가 다릅니다.')
        cursor = conn.execute('''INSERT INTO csat_wrong_answers
            (question_id, student_name, exam_name, question_no, csat_type_id, csat_type_name, wrong_answer, correct_answer, memo, source_filename, created_at)
            VALUES (?,?,?,?,?,?,?,?,?,?,?)''', (payload.question_id, payload.student_name.strip(), payload.exam_name.strip(), payload.question_no, csat_type['id'], csat_type['name'],
            payload.wrong_answer.strip(), payload.correct_answer.strip(), payload.memo.strip(), payload.source_filename.strip(), now))
        return dict(conn.execute('SELECT * FROM csat_wrong_answers WHERE id=?', (cursor.lastrowid,)).fetchone())


@app.get('/api/csat/wrong-answers')
def list_csat_wrong_answers(student: Optional[str] = None, exam: Optional[str] = None, type_id: Optional[str] = None,
                            question_no: Optional[int] = None, linked: Optional[bool] = None):
    filters, values = [], []
    for clause, value in (('w.student_name=?', student), ('w.exam_name=?', exam), ('w.csat_type_id=?', type_id), ('w.question_no=?', question_no)):
        if value: filters.append(clause); values.append(value)
    if linked is not None: filters.append('w.question_id IS NOT NULL' if linked else 'w.question_id IS NULL')
    where = f"WHERE {' AND '.join(filters)}" if filters else ''
    with get_db() as conn:
        rows = conn.execute(f'''SELECT w.*, w.id wrong_answer_id,
            q.source_id question_source_id, s.source_filename question_source_filename,
            q.answer question_answer, q.solution question_solution, q.raw_json question_raw_json
            FROM csat_wrong_answers w
            LEFT JOIN csat_questions q ON q.id=w.question_id
            LEFT JOIN csat_question_sources s ON s.id=q.source_id
            {where} ORDER BY w.id DESC''', values).fetchall()
    result = []
    for row in rows:
        item = dict(row)
        raw = json.loads(item.pop('question_raw_json') or '{}')
        item['question_linked'] = item['question_id'] is not None
        item['has_answer'] = bool(str(item.pop('question_answer') or '').strip())
        item['has_explanation'] = bool(str(item.pop('question_solution') or '').strip() or str(raw.get('explanation_html') or '').strip())
        result.append(item)
    return result


class CsatWrongLink(BaseModel):
    question_id: int


@app.patch('/api/csat/wrong-answers/{wrong_answer_id}/question')
def link_csat_wrong_question(wrong_answer_id: int, payload: CsatWrongLink):
    with get_db() as conn:
        wrong = conn.execute('SELECT * FROM csat_wrong_answers WHERE id=?', (wrong_answer_id,)).fetchone()
        question = conn.execute('SELECT * FROM csat_questions WHERE id=?', (payload.question_id,)).fetchone()
        if not wrong or not question:
            raise HTTPException(status_code=404, detail='오답 또는 Question을 찾을 수 없습니다.')
        if wrong['question_no'] != question['question_no']:
            raise HTTPException(status_code=400, detail='오답 번호와 Question의 수능 번호가 다릅니다.')
        conn.execute('UPDATE csat_wrong_answers SET question_id=? WHERE id=?', (payload.question_id, wrong_answer_id))
        return dict(conn.execute('SELECT * FROM csat_wrong_answers WHERE id=?', (wrong_answer_id,)).fetchone())


class CsatWrongAssessmentCreate(BaseModel):
    title: str = ''
    wrong_answer_ids: list[int] = Field(min_length=1, max_length=45)
    exclude_used: bool = False


@app.post('/api/csat/wrong-assessments', status_code=201)
def create_csat_wrong_assessment(payload: CsatWrongAssessmentCreate):
    unique_wrong_ids = list(dict.fromkeys(payload.wrong_answer_ids))
    with get_db() as conn:
        rows = conn.execute(f'''SELECT id, question_id FROM csat_wrong_answers
            WHERE id IN ({','.join('?' * len(unique_wrong_ids))})''', unique_wrong_ids).fetchall()
    by_id = {row['id']: row['question_id'] for row in rows}
    if len(by_id) != len(unique_wrong_ids):
        raise HTTPException(status_code=404, detail='선택한 오답 중 존재하지 않는 항목이 있습니다.')
    if any(by_id[wrong_id] is None for wrong_id in unique_wrong_ids):
        raise HTTPException(status_code=409, detail='Question ID가 연결되지 않은 오답이 있습니다.')
    question_ids = [by_id[wrong_id] for wrong_id in unique_wrong_ids]
    if len(set(question_ids)) != len(question_ids):
        raise HTTPException(status_code=409, detail='같은 Question이 중복 연결된 오답은 하나의 WRONG에 함께 넣을 수 없습니다.')
    return compose_csat_assessment(CsatAssessmentCreate(
        assessment_type='wrong', title=payload.title or 'BuildIt WRONG', question_ids=question_ids,
        count=len(question_ids), exclude_used=payload.exclude_used,
    ))


def render_csat_html(questions: list[dict], title: str = 'BuildIt 수능 영어', preserve_order: bool = False) -> str:
    def esc(value):
        import html
        return html.escape(str(value or ''))
    blocks = []
    ordered = questions if preserve_order else sorted(questions, key=lambda item: item['question_no'])
    for q in ordered:
        choices = ''.join(f'<li>{esc(choice)}</li>' for choice in q.get('choices', []))
        flag = '<span class="flag">유형 확인 필요</span>' if q.get('type_mismatch') else ''
        blocks.append(f'''<article class="question" data-question-id="{q.get('id', '')}" data-day-no="{q.get('day_no') or ''}"><nav class="question-tools"><button onclick="moveQuestion(this,-1)">↑</button><button onclick="moveQuestion(this,1)">↓</button><button onclick="this.closest('.question').remove()">삭제</button></nav><h2>{q['question_no']}. <span contenteditable="true">{esc(q.get('prompt'))}</span> {flag}</h2>
            <div class="passage" contenteditable="true">{esc(q.get('passage')).replace(chr(10), '<br>')}</div>
            <ol class="choices" contenteditable="true">{choices}</ol><small contenteditable="true">{esc(q.get('footnote'))}</small></article>''')
    return f'''<!doctype html><html lang="ko"><meta charset="utf-8"><title>{esc(title)}</title><style>
        @page{{size:A4;margin:12mm}}*{{box-sizing:border-box}}body{{font-family:"Noto Serif KR","Malgun Gothic",serif;margin:0;color:#111}}
        header{{display:flex;justify-content:space-between;border-bottom:2px solid #111;margin-bottom:12px}}header b{{font-size:22px}}
        .toolbar{{position:sticky;top:0;background:#fff4cf;padding:8px;text-align:center;z-index:2}}.paper{{column-count:2;column-gap:28px;column-rule:1px solid #222}}
        .question{{break-inside:avoid;page-break-inside:avoid;margin:0 0 18px;font-size:11.5px;line-height:1.55;position:relative}}.question-tools{{display:none;position:absolute;right:0;top:-5px}}.question:hover .question-tools{{display:block}}h2{{font-size:12px;margin:0 0 7px}}.passage{{text-align:justify;margin-left:14px}}
        .choices{{list-style:none;padding:0 0 0 14px;margin:8px 0;counter-reset:item}}.choices li{{counter-increment:item}}.choices li:before{{content:counter(item, circle) " ";}}
        .flag{{font:10px sans-serif;color:#a33;border:1px solid #a33;padding:1px 4px}}[contenteditable=true]:focus{{outline:2px solid #e5b94f;background:#fff9e8}}
        @media print{{.toolbar,.question-tools{{display:none!important}}.paper{{column-fill:auto}}.question:nth-of-type(4n+1):not(:first-child){{break-before:page}}}}@media(max-width:800px){{.paper{{column-count:1;column-rule:0}}}}
    </style><body><div class="toolbar">문구를 클릭해 수정할 수 있습니다. <button onclick="print()">인쇄</button> <button onclick="saveEdited()">편집 HTML 저장</button></div>
    <header><b contenteditable="true">{esc(title)}</b><span>BuildIt 자체 제작 시험지 · A4 2단 · 페이지당 4문항 목표</span></header><main class="paper">{''.join(blocks)}</main><script>function moveQuestion(button,direction){{const item=button.closest('.question'),parent=item.parentElement,sibling=direction<0?item.previousElementSibling:item.nextElementSibling;if(!sibling)return;direction<0?parent.insertBefore(item,sibling):parent.insertBefore(sibling,item)}}function saveEdited(){{const blob=new Blob(['<!doctype html>'+document.documentElement.outerHTML],{{type:'text/html;charset=utf-8'}});const a=document.createElement('a');a.href=URL.createObjectURL(blob);a.download='buildit-csat-edited.html';a.click();setTimeout(()=>URL.revokeObjectURL(a.href),500)}}</script></body></html>'''


def render_csat_explanation_html(questions: list[dict], title: str) -> str:
    import html
    blocks = []
    for item in questions:
        raw = json.loads(item.get('raw_json') or '{}')
        explanation = raw.get('explanation_html') or html.escape(item.get('solution') or '해설 없음').replace('\n', '<br>')
        blocks.append(f'''<article class="explanation" data-question-id="{item['id']}" data-day-no="{item.get('day_no') or ''}">
            <h2>{item['question_no']}번 · {html.escape(item.get('csat_type_name') or '')}</h2>
            <p><b>정답</b> {html.escape(item.get('answer') or '')}</p><div contenteditable="true">{explanation}</div></article>''')
    return f'''<!doctype html><html lang="ko"><meta charset="utf-8"><title>{html.escape(title)} 해설</title><style>
        @page{{size:A4;margin:14mm}}body{{font-family:"Noto Serif KR","Malgun Gothic",serif;color:#211b13;background:#fffaf0}}
        main{{columns:2;column-gap:28px}}article{{break-inside:avoid;border-top:2px solid #d58a1f;padding:8px 4px 18px;line-height:1.6}}
        h1{{color:#c66f12}}h2{{font-size:14px}}[contenteditable=true]:focus{{outline:2px solid #e5b94f;background:#fff}}
        @media(max-width:800px){{main{{columns:1}}}}</style><body><h1 contenteditable="true">{html.escape(title)} 해설</h1><main>{''.join(blocks)}</main></body></html>'''


def _safe_csat_fragment(value: str) -> str:
    fragment = re.sub(r'(?is)<(?:script|iframe|object|embed|form)\b.*?</(?:script|iframe|object|embed|form)>', '', value or '')
    fragment = re.sub(r'''(?is)\s+on[a-z]+\s*=\s*(?:"[^"]*"|'[^']*')''', '', fragment)
    return re.sub(r'(?i)javascript\s*:', '', fragment)


def _daily_question_content(item: dict) -> str:
    import html
    raw = json.loads(item.get('raw_json') or '{}')
    if raw.get('question_html'):
        return _safe_csat_fragment(raw['question_html'])
    choices = ''.join(f'<li>{html.escape(str(choice))}</li>' for choice in item.get('choices', []))
    return f'''<h2>{item['question_no']}. {html.escape(item.get('prompt') or '')}</h2>
        <div class="passage">{html.escape(item.get('passage') or '').replace(chr(10), '<br>')}</div><ol>{choices}</ol>'''


def render_daily_problem_html(questions: list[dict], title: str) -> str:
    import html
    pages = []
    days = sorted({int(item['day_no']) for item in questions})
    for day_no in days:
        day_items = [item for item in questions if item['day_no'] == day_no]
        for page_index in range(0, len(day_items), 2):
            blocks = ''.join(f'''<article class="daily-question" data-question-id="{item['id']}" data-day-no="{day_no}">
                <div class="daily-question-content" contenteditable="true">{_daily_question_content(item)}</div></article>'''
                for item in day_items[page_index:page_index + 2])
            pages.append(f'''<section class="daily-page" data-day-no="{day_no}"><header><strong>Day {day_no}</strong>
                <span>{html.escape(title)} · 수능 번호 유형 파악 훈련</span></header><main>{blocks}</main><footer>BuildIt DAILY · {day_no}</footer></section>''')
    return f'''<!doctype html><html lang="ko"><meta charset="utf-8"><title>{html.escape(title)}</title><style>
        @page{{size:A4;margin:0}}*{{box-sizing:border-box}}body{{margin:0;background:#ddd;color:#17120d;font-family:"Noto Serif KR","Malgun Gothic",serif}}
        .daily-page{{width:210mm;min-height:297mm;margin:0 auto 8mm;background:white;padding:20mm 20mm 14mm;break-after:page;position:relative}}
        header{{display:flex;align-items:flex-end;border-bottom:3px double #ad6708;margin-bottom:10mm;color:#a95f04}}
        header strong{{display:grid;place-items:end start;width:35mm;height:28mm;padding:0 3mm 2mm;background:#b66b08;color:white;font:900 27px Arial}}
        header span{{margin-left:auto;padding:0 0 3mm;font:800 14px "Malgun Gothic"}}main{{display:grid;gap:15mm}}
        .daily-question{{break-inside:avoid;min-height:105mm;font-size:11.5px;line-height:1.55}}.daily-question :is(article,section){{margin:0}}
        .daily-question h2{{font-size:12.5px}}.daily-question .passage{{text-align:justify}}.daily-question ol{{padding-left:22px}}
        footer{{position:absolute;left:20mm;bottom:10mm;color:#b66b08;font:700 10px Arial}}[contenteditable=true]:focus{{outline:2px solid #df9b34;background:#fff8e9}}
        @media(max-width:800px){{.daily-page{{width:100%;min-height:auto;margin:0;padding:24px}}}}@media print{{body{{background:white}}.daily-page{{margin:0}}}}</style>
        <body>{''.join(pages)}</body></html>'''


def render_daily_explanation_html(questions: list[dict], title: str) -> str:
    import html
    pages = []
    for day_no in sorted({int(item['day_no']) for item in questions}):
        blocks = []
        for item in [question for question in questions if question['day_no'] == day_no]:
            raw = json.loads(item.get('raw_json') or '{}')
            explanation = _safe_csat_fragment(raw.get('explanation_html') or '')
            if not explanation:
                explanation = html.escape(item.get('solution') or '해설 없음').replace('\n', '<br>')
            blocks.append(f'''<article class="daily-explanation" data-question-id="{item['id']}" data-day-no="{day_no}">
                <h2>{item['question_no']}번 <span>정답 {html.escape(item.get('answer') or '')}</span></h2>
                <div contenteditable="true">{explanation}</div></article>''')
        pages.append(f'''<section class="daily-answer-page" data-day-no="{day_no}"><header><b>Day {day_no}</b><strong>정답과 해설</strong></header>
            <main>{''.join(blocks)}</main><footer>{html.escape(title)} · {day_no}</footer></section>''')
    return f'''<!doctype html><html lang="ko"><meta charset="utf-8"><title>{html.escape(title)} 정답과 해설</title><style>
        @page{{size:A4;margin:0}}*{{box-sizing:border-box}}body{{margin:0;background:#d9d2c4;color:#271f15;font-family:"Noto Sans KR","Malgun Gothic",sans-serif}}
        .daily-answer-page{{width:210mm;min-height:297mm;margin:0 auto 8mm;padding:18mm;background:#fff9e9;break-after:page;position:relative}}
        header{{border-top:2px solid #b76b08;padding:7mm 0 5mm;color:#b76b08;display:flex;gap:8px;align-items:baseline}}header b{{font-size:18px}}header strong{{font-size:13px}}
        main{{column-count:2;column-gap:10mm;column-rule:1px solid #ead5ad}}.daily-explanation{{break-inside:avoid;border-top:1px solid #d18a24;padding:5mm 0 7mm;font-size:10.5px;line-height:1.62}}
        h2{{font-size:13px;margin:0 0 3mm;color:#2d251b}}h2 span{{color:#b76b08;margin-left:8px}}footer{{position:absolute;right:18mm;bottom:9mm;color:#9b6b2d;font-size:9px}}
        [contenteditable=true]:focus{{outline:2px solid #df9b34;background:white}}@media(max-width:800px){{.daily-answer-page{{width:100%;min-height:auto;margin:0;padding:24px}}main{{column-count:1}}}}
        @media print{{body{{background:white}}.daily-answer-page{{margin:0}}}}</style><body>{''.join(pages)}</body></html>'''


def _structured_csat_text(value: str, markup: list, field: str) -> str:
    """Render JSON V1 text without treating model-provided strings as HTML."""
    import html
    rendered = html.escape(str(value or '')).replace('\n', '<br>')
    for mark in markup if isinstance(markup, list) else []:
        if not isinstance(mark, dict) or str(mark.get('field') or 'passage') != field:
            continue
        target = html.escape(str(mark.get('text') or mark.get('target') or ''))
        if not target or target not in rendered:
            continue
        kind = str(mark.get('type') or '').lower()
        replacement = (f'<u>{target}</u>' if kind == 'underline' else
                       f'<strong>{target}</strong>' if kind in ('bold', 'emphasis') else
                       f'<span class="blank">{target}</span>' if kind == 'blank' else target)
        if kind in ('box', 'given_sentence'):
            replacement = f'<span class="given-sentence">{target}</span>'
        rendered = rendered.replace(target, replacement, 1)
    return rendered


def _structured_csat_question(item: dict, include_passage: bool = True) -> str:
    import html
    raw = json.loads(item.get('raw_json') or '{}')
    if raw.get('question_html'):
        return _safe_csat_fragment(raw['question_html'])
    problem = raw.get('problem') if isinstance(raw.get('problem'), dict) else {}
    markup = raw.get('markup') or problem.get('markup') or []
    instruction = problem.get('instruction', item.get('prompt') or '')
    passage = problem.get('passage', item.get('passage') or '')
    choices = problem.get('choices') if isinstance(problem.get('choices'), list) else item.get('choices', [])
    choice_html = ''.join(f'<li>{_structured_csat_text(str(choice), markup, "choices")}</li>' for choice in choices)
    passage_html = f'<div class="passage" contenteditable="true">{_structured_csat_text(passage, markup, "passage")}</div>' if include_passage else ''
    return f'''<h2><span contenteditable="true">{_structured_csat_text(instruction, markup, 'instruction')}</span></h2>
        {passage_html}<ol class="choices" contenteditable="true">{choice_html}</ol>'''


def render_csat_problem_document(questions: list[dict], title: str, renderer: Literal['compact', 'full']) -> str:
    """Editable A4 two-column output. CSS flows content; it never hard-codes questions per page."""
    import html
    ordered = questions if renderer == 'compact' else sorted(questions, key=lambda item: item['question_no'])
    blocks, seen_common = [], set()
    for item in ordered:
        raw = json.loads(item.get('raw_json') or '{}')
        common_id = str(raw.get('common_passage_id') or '')
        include_passage = not common_id or common_id not in seen_common
        if common_id:
            seen_common.add(common_id)
        group = f' data-common-passage-id="{html.escape(common_id)}"' if common_id else ''
        day = f' data-day-no="{item.get("day_no") or ""}"'
        blocks.append(f'''<article class="question" data-question-id="{item.get('id', '')}"{day}{group}>
            <nav class="question-tools"><button onclick="moveQuestion(this,-1)">↑</button><button onclick="moveQuestion(this,1)">↓</button><button onclick="this.closest('.question').remove()">삭제</button></nav>
            <div class="question-number">{item['question_no']}.</div><div class="question-content">{_structured_csat_question(item, include_passage)}</div></article>''')
    density = 'full-reading' if renderer == 'full' else 'compact'
    subtitle = '18–45 수능 독해형 · 문항 길이에 따른 자동 흐름' if renderer == 'full' else 'HALF · WRONG · DAILY Compact'
    return f'''<!doctype html><html lang="ko"><meta charset="utf-8"><title>{html.escape(title)}</title><style>
        @page{{size:A4;margin:11mm}}*{{box-sizing:border-box}}body{{font-family:"Noto Serif KR","Malgun Gothic",serif;margin:0;color:#111}}
        .toolbar{{position:sticky;top:0;background:#fff4cf;padding:7px;text-align:center;z-index:3}}header{{display:flex;justify-content:space-between;align-items:end;border-bottom:2px solid #111;margin-bottom:8px}}header b{{font-size:19px}}
        .paper{{column-count:2;column-gap:8mm;column-rule:1px solid #333;column-fill:auto}}.question{{display:grid;grid-template-columns:20px 1fr;break-inside:avoid;page-break-inside:avoid;margin:0 0 11px;font-size:10.5px;line-height:1.44;position:relative}}
        .full-reading .question{{font-size:9.5pt;line-height:1.38;margin-bottom:9px}}.question-number{{font-weight:800}}.question-content h2{{font:inherit;font-weight:600;margin:0 0 4px}}.passage{{text-align:justify;white-space:normal}}.choices{{list-style:none;padding:0;margin:5px 0;counter-reset:item}}.choices li{{counter-increment:item}}.choices li:before{{content:counter(item, circle) " ";}}
        .blank{{display:inline-block;min-width:42px;border-bottom:1px solid #111;color:transparent}}.given-sentence{{display:block;border:1px solid #333;padding:5px 7px;margin:5px 0}}.question-tools{{display:none;position:absolute;right:0;top:-6px;background:white}}.question:hover .question-tools{{display:block}}[contenteditable=true]:focus{{outline:2px solid #e5b94f;background:#fff9e8}}
        @media(max-width:800px){{.paper{{column-count:1;column-rule:0}}}}@media print{{.toolbar,.question-tools{{display:none!important}}}}
    </style><body class="{density}"><div class="toolbar">문구를 클릭해 수정할 수 있습니다. <button onclick="print()">인쇄</button> <button onclick="saveEdited()">편집 HTML 저장</button></div>
    <header><b contenteditable="true">{html.escape(title)}</b><span>{subtitle}</span></header><main class="paper">{''.join(blocks)}</main>
    <script>function moveQuestion(button,direction){{const item=button.closest('.question'),parent=item.parentElement,sibling=direction<0?item.previousElementSibling:item.nextElementSibling;if(!sibling)return;direction<0?parent.insertBefore(item,sibling):parent.insertBefore(sibling,item)}}function saveEdited(){{const blob=new Blob(['<!doctype html>'+document.documentElement.outerHTML],{{type:'text/html;charset=utf-8'}});const a=document.createElement('a');a.href=URL.createObjectURL(blob);a.download='buildit-csat-edited.html';a.click();setTimeout(()=>URL.revokeObjectURL(a.href),500)}}</script></body></html>'''


def render_assessment_problem_html(assessment: dict) -> str:
    if assessment['assessment_type'] == 'daily':
        return render_daily_problem_html(assessment['items'], assessment['title']).replace(
            '<body>', f'<body data-template-id="{assessment["template_id"]}" data-assessment-type="daily">', 1)
    renderer = 'full' if assessment['assessment_type'] == 'full' else 'compact'
    document = render_csat_problem_document(assessment['items'], assessment['title'], renderer)
    return document.replace('<body ', f'<body data-template-id="{assessment["template_id"]}" data-assessment-type="{assessment["assessment_type"]}" ', 1)


def render_assessment_explanation_html(assessment: dict) -> str:
    if assessment['assessment_type'] == 'daily':
        return render_daily_explanation_html(assessment['items'], assessment['title'])
    document = render_csat_explanation_html(assessment['items'], assessment['title'])
    return document.replace('<body>', f'<body data-template-id="{assessment["template_id"]}" data-assessment-type="{assessment["assessment_type"]}">', 1)


@app.get('/api/csat/assessments/{assessment_id}/problem.html', response_class=HTMLResponse)
def render_saved_assessment_problem(assessment_id: int):
    with get_db() as conn:
        assessment = _assessment_detail(conn, assessment_id)
    return render_assessment_problem_html(assessment)


@app.get('/api/csat/assessments/{assessment_id}/explanation.html', response_class=HTMLResponse)
def render_saved_assessment_explanation(assessment_id: int):
    with get_db() as conn:
        assessment = _assessment_detail(conn, assessment_id)
    return render_assessment_explanation_html(assessment)


@app.get('/api/csat/exam.html', response_class=HTMLResponse)
def export_csat_exam_html(ids: str = '', exam_kind: Optional[Literal['half', 'full']] = None, title: str = 'BuildIt 수능 영어'):
    selected = [int(value) for value in ids.split(',') if value.strip().isdigit()]
    filters, values = [], []
    if selected: filters.append(f"q.id IN ({','.join('?' * len(selected))})"); values.extend(selected)
    if exam_kind: filters.append('q.exam_kind=?'); values.append(exam_kind)
    where = f"WHERE {' AND '.join(filters)}" if filters else ''
    with get_db() as conn:
        rows = conn.execute(f'SELECT q.* FROM csat_questions q {where} ORDER BY q.question_no', values).fetchall()
    items = []
    for row in rows:
        item = dict(row); item['choices'] = json.loads(item['choices_json']); item['type_mismatch'] = bool(item['type_mismatch']); items.append(item)
    return render_csat_problem_document(items, title, 'full' if exam_kind == 'full' else 'compact')


PASSAGE_MEMORY_SCHEMA = {
    'type': 'object', 'properties': {
        'title': {'type': 'string'},
        'sentences': {'type': 'array', 'items': {'type': 'object', 'properties': {
            'english': {'type': 'string'}, 'korean': {'type': 'string'},
        }, 'required': ['english', 'korean']}},
    }, 'required': ['title', 'sentences'],
}

PASSAGE_ANALYSIS_SCHEMA = {
    'type': 'object', 'properties': {
        'title': {'type': 'string'},
        'topic': {'type': 'string'}, 'mainIdea': {'type': 'string'}, 'summary': {'type': 'string'},
        'sentences': {'type': 'array', 'items': {'type': 'object', 'properties': {
            'slashed': {'type': 'string'}, 'directTranslation': {'type': 'string'},
            'translation': {'type': 'string'}, 'grammarGuide': {'type': 'string'},
        }, 'required': ['slashed', 'directTranslation', 'translation', 'grammarGuide']}},
        'tfQuestions': {'type': 'array', 'items': {'type': 'object', 'properties': {
            'language': {'type': 'string', 'enum': ['KO', 'EN']}, 'statement': {'type': 'string'},
            'answer': {'type': 'string', 'enum': ['T', 'F']},
        }, 'required': ['language', 'statement', 'answer']}},
    }, 'required': ['title', 'topic', 'mainIdea', 'summary', 'sentences', 'tfQuestions'],
}


def _passage_source(payload: PassageRequest) -> str:
    source = payload.source_text.strip()
    if not source:
        raise HTTPException(status_code=400, detail='분석할 영어 지문이 없습니다.')
    return source[:30000]


@app.post('/api/passages/memory')
def create_passage_memory(payload: PassageRequest):
    source = _passage_source(payload)
    rules = ('영어 학습용 영한 대조 자료를 만든다. 원문을 빠짐없이 자연스러운 문장 단위로 나누고, '
             '각 문장에 정확하고 자연스러운 한국어 번역을 붙인다. 영어 원문의 의미나 순서를 바꾸지 않는다.')
    try:
        ai = ai_router.generate('auto', rules, source, schema=PASSAGE_MEMORY_SCHEMA, temperature=0.15)
        result = json.loads(re.sub(r'```json|```', '', ai.text).strip())
        return {**result, 'provider': ai.provider, 'model': ai.model}
    except Exception as exc:
        _handle_gemini_error(exc)


@app.post('/api/passages/analysis')
def create_passage_analysis(payload: PassageRequest):
    source = _passage_source(payload)
    rules = ('외부 영어 지문을 밀도 높은 학습 자료로 정리한다. 영어 원문 전체를 별도 반복하지 않는다. '
             '모든 문장을 원래 순서대로 정확히 한 번만 sentences에 넣는다. slashed는 영어 문장 자체에 의미 단위 / 를 넣은 문자열이고, '
             '바로 이어 directTranslation은 슬래시 단위의 직독직해, translation은 자연스러운 정식 해석, grammarGuide는 핵심 문법과 구문 설명이다. '
             '내용일치 T/F는 정확히 10문항으로 앞 5개 한국어(KO), 뒤 5개 영어(EN), T/F를 모두 포함하고 지문만으로 판정 가능하게 한다.')
    try:
        ai = ai_router.generate('auto', rules, source, schema=PASSAGE_ANALYSIS_SCHEMA, temperature=0.2)
        result = json.loads(re.sub(r'```json|```', '', ai.text).strip())
        if len(result.get('tfQuestions', [])) != 10:
            raise HTTPException(status_code=502, detail='T/F 문항이 10개로 생성되지 않았습니다. 다시 시도해 주세요.')
        return {**result, 'provider': ai.provider, 'model': ai.model}
    except HTTPException:
        raise
    except Exception as exc:
        _handle_gemini_error(exc)


@app.post('/api/ocr/clean-text')
def clean_ocr_text(payload: OcrCleanRequest):
    raw_text = payload.raw_text.strip()
    if not raw_text:
        raise HTTPException(status_code=400, detail='정리할 추출 텍스트가 없습니다.')
    rules = pl.build_ocr_clean_rules()
    try:
        ai = ai_router.generate('auto', rules, raw_text[:20000], temperature=0.1, allow_noncritical_fallback=True)
        text = ai.text
    except Exception as exc:
        _handle_gemini_error(exc)
        return
    return {'cleanedText': text.strip(), 'provider': ai.provider, 'model': ai.model}


# ==================================================================================
# 학교 프린트 사진 → 구조화 JSON + 편집 가능 HTML
# ==================================================================================
class SchoolPrintPage(BaseModel):
    page: int
    name: str
    ocr_text: str = ''
    image_data_url: str = ''

class SchoolPrintReconstructRequest(BaseModel):
    model: str = ai_router.QWEN_MODEL
    metadata: dict = Field(default_factory=dict)
    pages: list[SchoolPrintPage]

class SchoolPrintPatchRequest(BaseModel):
    model: str = ai_router.QWEN_MODEL
    instruction: str
    html: str
    canonical_json: Optional[dict] = None

def _school_print_json(text: str) -> dict:
    try:return json.loads(re.sub(r'^\s*```(?:json)?|```\s*$', '', text.strip(), flags=re.I))
    except (json.JSONDecodeError, TypeError):raise HTTPException(status_code=502, detail='AI 응답을 구조화 JSON으로 해석하지 못했습니다.')

@app.post('/api/school-prints/reconstruct')
def reconstruct_school_print(payload: SchoolPrintReconstructRequest):
    if not payload.pages:raise HTTPException(status_code=400, detail='복원할 페이지가 없습니다.')
    system='''학교 영어 프린트 2단계 복원기의 2차 HTML 복원 단계다. 각 PAGE의 OCR TEXT는 이미 브라우저 Tesseract/PDF 텍스트 레이어로 확보한 원문 내용의 기준(source of truth)이고, 함께 제공된 원본 이미지는 배치와 시각 구조의 기준이다. OCR TEXT의 문제·문장·보기·정답을 창작·요약·변형·누락하지 않는다. 이미지에서 제목, A/B/C 섹션, 문항번호, 국가/국적 4열 표, 영영풀이 박스, 밑줄, 빈칸, 답안선, 보기와 강조를 확인해 실제 편집 가능한 semantic HTML(table, thead, tbody, section, u, span.blank, div.answer-line, div.definition-box)로 재구성한다. background-image, 원본 이미지 삽입, 캔버스 기반 가짜 복원은 금지한다. 판독 불가능한 곳은 원문을 유지하고 class="review-required"로 표시한다. {"html":string,"canonicalJson":object} JSON만 반환한다. canonicalJson은 source/pages/sections/blocks(type,order,content,page,reviewRequired)/sourceLinkage/restorationStatus/restorationVersion/version/editedHtml/model을 포함한다.'''
    content=[{'type':'text','text':'metadata: '+json.dumps(payload.metadata,ensure_ascii=False)}]
    for page in payload.pages:
        content.append({'type':'text','text':f'PAGE {page.page} / {page.name}\nOCR TEXT (CONTENT SOURCE OF TRUTH):\n{page.ocr_text}'})
        if page.image_data_url.startswith('data:image/'):content.append({'type':'image_url','image_url':{'url':page.image_data_url}})
    try:ai=ai_router.generate_multimodal(payload.model,system,content)
    except Exception as exc:_handle_gemini_error(exc)
    result=_school_print_json(ai.text);html=str(result.get('html') or '').strip()
    if not html:raise HTTPException(status_code=502,detail='선택한 모델이 HTML을 반환하지 않았습니다.')
    canonical=result.get('canonicalJson') if isinstance(result.get('canonicalJson'),dict) else {}
    extracted_text=[{'page':p.page,'sourceName':p.name,'text':p.ocr_text} for p in payload.pages]
    source={**payload.metadata,'fileNames':[p.name for p in payload.pages]}
    source.update(canonical.get('source') if isinstance(canonical.get('source'),dict) else {})
    canonical_pages=canonical.get('pages') if isinstance(canonical.get('pages'),list) else []
    by_page={item.get('page'):item for item in canonical_pages if isinstance(item,dict)}
    canonical_pages=[{**by_page.get(p.page,{}),'page':p.page,'sourceName':p.name,'extractedText':p.ocr_text,'sections':by_page.get(p.page,{}).get('sections',[])} for p in payload.pages]
    canonical.update({'source':source,'pages':canonical_pages,'extractedText':extracted_text,'sourceLinkage':canonical.get('sourceLinkage') or {'pageCount':len(payload.pages),'uploadedAt':datetime.now().isoformat()},'restorationStatus':'structured','restorationVersion':2,'version':int(canonical.get('version') or 1),'editedHtml':html,'model':ai.model})
    return {'html':html,'canonicalJson':canonical,'provider':ai.provider,'model':ai.model}

@app.post('/api/school-prints/patch-html')
def patch_school_print_html(payload: SchoolPrintPatchRequest):
    if not payload.instruction.strip() or not payload.html.strip():raise HTTPException(status_code=400,detail='현재 HTML과 수정 지시가 필요합니다.')
    system='''학교 프린트 HTML 부분 수정기다. 사용자가 지적한 부분만 고치고 다른 문구·순서·표·스타일은 바꾸지 않는다. 전체 재작성과 내용 창작, background-image를 금지한다. {"html":string,"canonicalJson":object} JSON만 반환하고 blocks/version/editedHtml을 동기화한다.'''
    request=json.dumps({'instruction':payload.instruction,'currentHtml':payload.html,'canonicalJson':payload.canonical_json or {}},ensure_ascii=False)
    try:ai=ai_router.generate('qwen',system,request,temperature=.1)
    except Exception as exc:_handle_gemini_error(exc)
    result=_school_print_json(ai.text);html=str(result.get('html') or '').strip()
    if not html:raise HTTPException(status_code=502,detail='수정된 HTML이 비어 있습니다.')
    canonical=result.get('canonicalJson') if isinstance(result.get('canonicalJson'),dict) else (payload.canonical_json or {})
    canonical.update({'version':int(canonical.get('version') or 1)+1,'editedHtml':html,'model':ai.model})
    return {'html':html,'canonicalJson':canonical,'provider':ai.provider,'model':ai.model}


# ==================================================================================
# 학생 오답 관리 — 사진/PDF에서 브라우저 OCR + Gemini로 정리한 "최종 추출 텍스트"를
# 사용자가 화면에서 확인·수정한 뒤 저장한다. 이 저장 단계에서는 더 이상 AI를 호출하지
# 않는다(문제2, 6차): 예전에는 저장 버튼을 누르면 서버가 다시 한 번 Gemini를 호출해
# "틀린 문제만" 요약 추출했는데, 이 재처리 과정에서 사용자가 이미 확인한 내용의 일부가
# 잘려나가는 문제가 있었다. 이제 저장 버튼은 사용자가 화면에서 편집한 텍스트를 그대로,
# 손실 없이 DB에 넣는다. 각 저장은 독립된 record이며(덮어쓰지 않고 계속 누적),
# student_id로 학생과 연결된다 — 학교/학년 등은 note에 복사해 두지 않고, 조회 시
# students 테이블을 다시 조인해 항상 최신 정보를 보여준다(문제7).
# ==================================================================================

class WrongNoteCreate(BaseModel):
    exam_label: str = ''
    title: str = ''
    question_type: str = ''
    question_text: str
    answer_text: str = ''
    source_text: str = ''


def _wrong_note_response(row: dict, student: dict) -> dict:
    return {
        'id': row['id'], 'student_id': row['student_id'],
        'student_name': student['name'], 'student_school': student['school'], 'student_grade': student['grade'],
        'exam_label': row['exam_label'], 'title': row['title'], 'question_type': row['question_type'],
        'wrong_text': row['wrong_text'], 'answer_text': row['answer_text'], 'source_text': row['source_text'],
        'created_at': row['created_at'],
    }


@app.post('/api/students/{student_id}/wrong-notes')
def create_wrong_note(student_id: int, payload: WrongNoteCreate):
    if not payload.question_text.strip():
        raise HTTPException(status_code=400, detail='저장할 오답 문제 텍스트가 없습니다. 오답 사진을 업로드해 추출한 뒤 저장해 주세요.')
    with get_db() as conn:
        student = conn.execute('SELECT * FROM students WHERE id = ?', (student_id,)).fetchone()
        if not student:
            raise HTTPException(status_code=404, detail='학생 프로필을 찾지 못했습니다. 먼저 학생을 등록해 주세요.')
        now = datetime.now().isoformat()
        cursor = conn.execute(
            '''INSERT INTO student_wrong_notes
               (student_id, exam_label, title, question_type, wrong_text, answer_text, source_text, created_at)
               VALUES (?,?,?,?,?,?,?,?)''',
            (
                student_id, payload.exam_label.strip(), payload.title.strip(), payload.question_type.strip(),
                payload.question_text.strip()[:20000], payload.answer_text.strip()[:2000],
                payload.source_text.strip()[:20000], now,
            ),
        )
        row = conn.execute('SELECT * FROM student_wrong_notes WHERE id = ?', (cursor.lastrowid,)).fetchone()
        return _wrong_note_response(dict(row), dict(student))


@app.get('/api/students/{student_id}/wrong-notes')
def list_wrong_notes(student_id: int):
    with get_db() as conn:
        student = conn.execute('SELECT * FROM students WHERE id = ?', (student_id,)).fetchone()
        if not student:
            raise HTTPException(status_code=404, detail='학생 프로필을 찾지 못했습니다.')
        rows = conn.execute('SELECT * FROM student_wrong_notes WHERE student_id = ? ORDER BY id DESC', (student_id,)).fetchall()
        return [_wrong_note_response(dict(row), dict(student)) for row in rows]


class WrongNoteUpdate(BaseModel):
    exam_label: Optional[str] = None
    title: Optional[str] = None
    question_type: Optional[str] = None
    question_text: Optional[str] = None
    answer_text: Optional[str] = None


@app.patch('/api/students/{student_id}/wrong-notes/{note_id}')
def update_wrong_note(student_id: int, note_id: int, payload: WrongNoteUpdate):
    """문제6(6차): "전체보기"에서 수정한 뒤 다시 저장할 수 있도록 부분 수정을 지원한다."""
    with get_db() as conn:
        student = conn.execute('SELECT * FROM students WHERE id = ?', (student_id,)).fetchone()
        existing = conn.execute('SELECT * FROM student_wrong_notes WHERE id = ? AND student_id = ?', (note_id, student_id)).fetchone()
        if not student or not existing:
            raise HTTPException(status_code=404, detail='수정할 오답 자료를 찾지 못했습니다.')
        fields = {
            'exam_label': payload.exam_label, 'title': payload.title, 'question_type': payload.question_type,
            'wrong_text': payload.question_text, 'answer_text': payload.answer_text,
        }
        updates = {k: v.strip() for k, v in fields.items() if v is not None}
        if updates:
            set_clause = ', '.join(f'{k} = ?' for k in updates)
            conn.execute(f'UPDATE student_wrong_notes SET {set_clause} WHERE id = ?', (*updates.values(), note_id))
        row = conn.execute('SELECT * FROM student_wrong_notes WHERE id = ?', (note_id,)).fetchone()
        return _wrong_note_response(dict(row), dict(student))


@app.delete('/api/students/{student_id}/wrong-notes/{note_id}')
def delete_wrong_note(student_id: int, note_id: int):
    with get_db() as conn:
        conn.execute('DELETE FROM student_wrong_notes WHERE id = ? AND student_id = ?', (note_id, student_id))
    return {'ok': True}


class WrongNoteExportRequest(BaseModel):
    note_ids: list[int] = []
    title: str = ''


@app.post('/api/students/{student_id}/wrong-notes/export-docx')
def export_wrong_notes_docx(student_id: int, payload: WrongNoteExportRequest):
    """문제9~11(6차): 선택한 오답 문제들을 순서대로 하나의 통합 문제지(.docx)로 내보낸다.
    내신 실전 제작과 동일한 A4/2단/맑은 고딕 10pt/페이지 나누기 스타일을 재사용한다."""
    if not payload.note_ids:
        raise HTTPException(status_code=400, detail='통합할 오답 문항을 먼저 선택해 주세요.')
    with get_db() as conn:
        student = conn.execute('SELECT * FROM students WHERE id = ?', (student_id,)).fetchone()
        if not student:
            raise HTTPException(status_code=404, detail='학생 프로필을 찾지 못했습니다.')
        rows = conn.execute(
            f"SELECT * FROM student_wrong_notes WHERE student_id = ? AND id IN ({','.join('?' * len(payload.note_ids))})",
            (student_id, *payload.note_ids),
        ).fetchall()
    by_id = {row['id']: dict(row) for row in rows}
    ordered = [by_id[nid] for nid in payload.note_ids if nid in by_id]
    if not ordered:
        raise HTTPException(status_code=400, detail='선택한 오답 문항을 찾지 못했습니다.')

    title = payload.title.strip() or f"{student['name']} 오답 복습 문제"
    subtitle = ' · '.join(filter(None, [student['school'], student['grade'], student['name']]))
    doc = _build_wrong_note_combined_document(title, subtitle, ordered)

    export_dir = WORKSPACE / 'data' / 'exports'
    export_dir.mkdir(parents=True, exist_ok=True)
    safe_title = re.sub(r'[\\/:*?"<>|]', '_', title) or '오답_복습_문제'
    doc_path = export_dir / 'wrong-notes-combined.docx'
    doc.save(doc_path)
    return FileResponse(
        doc_path,
        media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        filename=f'{safe_title}.docx',
    )


class ExamMaterialWrongNoteExportRequest(BaseModel):
    material_ids: list[int] = []
    title: str = ''


def _exam_material_json_to_text(json_text: str, original_name: str) -> str:
    """exam_material_sets.json_text는 사용자가 자유롭게 올린 JSON이라 고정된 스키마가 없다
    (BUILDIT_CLAUDE_HANDOFF.md 5번: "세부 정보는 JSON 내부에서 확장 가능"). 어떤 형태든
    한 문단짜리 요약 텍스트로 풀어내며, docx 헬퍼는 줄바꿈을 별도 처리하지 않으므로
    개행 없이 ' / '로만 이어 붙인다."""
    try:
        data = json.loads(json_text) if json_text else None
    except (ValueError, TypeError):
        data = None
    if isinstance(data, dict):
        parts = [f'{k}: {v}' for k, v in data.items() if v not in (None, '', [], {})]
        text = ' / '.join(parts) if parts else json.dumps(data, ensure_ascii=False)
    elif isinstance(data, list):
        parts = []
        for i, item in enumerate(data, start=1):
            item_text = ' / '.join(f'{k}: {v}' for k, v in item.items()) if isinstance(item, dict) else str(item)
            parts.append(f'{i}) {item_text}')
        text = '  '.join(parts) if parts else '[]'
    else:
        text = (json_text or '').strip()
    return f'{text} (원본 파일: {original_name})' if original_name else text


@app.post('/api/exam-materials/wrong-notes/export-docx')
def export_exam_material_wrong_notes_docx(payload: ExamMaterialWrongNoteExportRequest):
    """NaesinWikiV1(원본+JSON 통합 저장)에서 모은 학생 오답 exam_material_sets 행을 선택해
    통합 복습 문제지로 내보낸다. student_wrong_notes 테이블의 export_wrong_notes_docx와 동일한
    _build_wrong_note_combined_document를 그대로 재사용한다 — 새 스타일을 만들지 않는다."""
    if not payload.material_ids:
        raise HTTPException(status_code=400, detail='내보낼 학생 오답 자료를 먼저 선택해 주세요.')
    with get_db() as conn:
        rows = conn.execute(
            f"""SELECT m.*, s.name student_name FROM exam_material_sets m
                LEFT JOIN students s ON s.id = m.student_id
                WHERE m.core_type='student_wrong_note' AND m.id IN ({','.join('?' * len(payload.material_ids))})""",
            payload.material_ids,
        ).fetchall()
    by_id = {row['id']: dict(row) for row in rows}
    ordered = [by_id[mid] for mid in payload.material_ids if mid in by_id]
    if not ordered:
        raise HTTPException(status_code=400, detail='선택한 학생 오답 자료를 찾지 못했습니다.')

    notes = [{
        'title': ' · '.join(filter(None, [row.get('student_name'), row.get('title') or row.get('original_name')])),
        'wrong_text': _exam_material_json_to_text(row.get('json_text') or '', row.get('original_name') or ''),
        'answer_text': '',
        'question_type': '',
    } for row in ordered]
    first = ordered[0]
    title = payload.title.strip() or f"{first.get('school', '')} {first.get('exam', '')} 학생 오답 복습 문제".strip()
    subtitle = ' · '.join(filter(None, [first.get('school'), first.get('grade'), f"{first.get('year')}년 {first.get('semester')}" if first.get('year') else '']))
    doc = _build_wrong_note_combined_document(title, subtitle, notes, '학생 오답 복습 문제')

    export_dir = WORKSPACE / 'data' / 'exports'
    export_dir.mkdir(parents=True, exist_ok=True)
    safe_title = re.sub(r'[\\/:*?"<>|]', '_', title) or '학생_오답_복습_문제'
    doc_path = export_dir / 'exam-material-wrong-notes-combined.docx'
    doc.save(doc_path)
    return FileResponse(
        doc_path,
        media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        filename=f'{safe_title}.docx',
    )


@app.get('/api/export/project')
def export_project():
    files = ['backend/main.py', 'backend/prompt_logic.py', 'src', 'package.json', 'package-lock.json', 'vite.config.ts', 'tsconfig.json', 'index.html', 'requirements.txt', 'data/app.db']
    try:
        if EXPORT_PATH.exists():
            EXPORT_PATH.unlink()
        with zipfile.ZipFile(EXPORT_PATH, 'w', zipfile.ZIP_DEFLATED) as archive:
            for relative in files:
                target = WORKSPACE / relative
                if target.is_file():
                    archive.write(target, target.relative_to(WORKSPACE))
                elif target.is_dir():
                    for item in target.rglob('*'):
                        if item.is_file() and '.DS_Store' not in item.name and not item.name.endswith('.map'):
                            archive.write(item, item.relative_to(WORKSPACE))
    except OSError as error_message:
        raise HTTPException(status_code=500, detail='압축파일을 준비하지 못했습니다.') from error_message
    return FileResponse(EXPORT_PATH, media_type='application/zip', filename='feel-to-real-project-export.zip')

class SchoolPrintExportItem(BaseModel):
    title: str = ''
    text: str = ''
    answer: str = ''

class SchoolPrintExportRequest(BaseModel):
    title: str = ''
    subtitle: str = ''
    items: list[SchoolPrintExportItem] = []

@app.post('/api/school-prints/export-docx')
def export_school_prints_docx(payload: SchoolPrintExportRequest):
    if not payload.items:
        raise HTTPException(status_code=400, detail='내보낼 학교 프린트를 선택해 주세요.')
    notes = [{'title': x.title, 'wrong_text': x.text, 'answer_text': x.answer, 'question_type': ''} for x in payload.items]
    title = payload.title.strip() or '학교 프린트'
    doc = _build_wrong_note_combined_document(title, payload.subtitle.strip(), notes, '학교 프린트')
    export_dir = WORKSPACE / 'data' / 'exports'; export_dir.mkdir(parents=True, exist_ok=True)
    path = export_dir / 'school-prints-combined.docx'; doc.save(path)
    safe = re.sub(r'[\\/:*?"<>|]', '_', title) or '학교_프린트'
    return FileResponse(path, media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document', filename=f'{safe}.docx')


# ==================================================================================
# MCP Server — 위 Workspace Context API(get_final_workspace_context /
# get_jikbo_workspace_context / create_final_deliverable) 위의 얇은 adapter를
# /api/mcp 와 /api/mcp/에 redirect 없이 연결한다. 실제 Context 조립·저장 로직은
# 전부 이 파일에 그대로 있고,
# mcp_server.py는 그 함수들을 호출만 한다(자세한 설계는 mcp_server.py 모듈 docstring).
# ==================================================================================
from mcp_server import McpEndpointMiddleware, build_mcp_asgi_app  # noqa: E402
app.add_middleware(McpEndpointMiddleware, mcp_app=build_mcp_asgi_app())
